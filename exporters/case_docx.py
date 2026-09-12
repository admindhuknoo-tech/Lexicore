"""Professional DOCX renderer for Case Analysis."""
from io import BytesIO
from identity_profile import get_identity_profile
import re
from docx import Document
from .common import (
    _case_export_sections,
    _export_scalar,
    prepare_case_export_for_reader,
    _presentation_source_corpus,
    LexiCoreCivilPresentationSanitizer,
    LexiCoreLowLevelRenderGovernor,
)


def _roman(number: int) -> str:
    values = (
        (1000, "M"), (900, "CM"), (500, "D"), (400, "CD"),
        (100, "C"), (90, "XC"), (50, "L"), (40, "XL"),
        (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I"),
    )
    n = int(number)
    out = []
    for value, symbol in values:
        while n >= value:
            out.append(symbol)
            n -= value
    return "".join(out)


def _docx_set_cell_shading(cell, fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def _docx_set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def _docx_add_section_title(doc, number, heading):
    from docx.shared import Pt, Mm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Mm(0)
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f'{_roman(number)}. {str(heading or "").strip().rstrip(".").upper()}')
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12.5)
    return p


def _apply_legal_body_format(paragraph, *, indent=True):
    from docx.shared import Pt, Mm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.first_line_indent = Mm(0)
    paragraph.paragraph_format.space_after = Pt(3)
    return paragraph


def _set_run_body_font(run, *, bold=False, size=10.5):
    from docx.shared import Pt
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.bold = bold
    return run


def _split_pipe(value):
    return [part.strip() for part in str(value or '').split(' | ')]


def _docx_add_reference_table(doc, headers, rows, widths_mm):
    from docx.shared import Pt, Mm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = 'Table Grid'

    for idx, (header, width) in enumerate(zip(headers, widths_mm)):
        cell = table.rows[0].cells[idx]
        cell.width = Mm(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        _docx_set_cell_shading(cell, '10263D')
        _docx_set_cell_margins(cell, 70, 80, 70, 80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Mm(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(header))
        _set_run_body_font(r, bold=True, size=8.3)
        r.font.color.rgb = RGBColor(255, 255, 255)

    for row in rows:
        cells = table.add_row().cells
        for idx, (value, width) in enumerate(zip(row, widths_mm)):
            cell = cells[idx]
            cell.width = Mm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            _docx_set_cell_margins(cell, 70, 80, 70, 80)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Mm(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value or '-'))
            _set_run_body_font(r, size=8.1)
    return table


def export_case_docx(x):
    profile = get_identity_profile()
    from docx.shared import Pt, Mm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def is_structured_line(text: str) -> bool:
        prefixes = (
            'ISU:', 'FAKTA/DALIL:', 'SUMBER:', 'ATURAN:', 'ANALISIS:', 'STATUS:',
            'HUKUM YANG BERLAKU:', 'UNSUR HUKUM:', 'PERBUATAN YANG DIDALILKAN:',
            'BUKTI:', 'BUKTI/BANTAHAN LAWAN:', 'UJI UNSUR:', 'KAUSALITAS:',
            'RISIKO:', 'KLASIFIKASI PROSEDURAL/POKOK PERKARA:',
            'TINDAKAN YANG DIREKOMENDASIKAN:', 'TINDAKAN:', 'RANTAI ',
            'Langkah ', 'Ranah hukum:',
        )
        return text.startswith(prefixes) or (' | ' in text)

    def summary_matrix_projection(lines):
        before, rows, after = [], [], []
        in_matrix = False
        for line0 in lines:
            value = governor.intercept(line0)
            if value is None:
                continue
            text = str(value).strip()
            if text.startswith('Matriks temuan utama:'):
                in_matrix = True
                continue
            parts = _split_pipe(text)
            if in_matrix and len(parts) == 4 and not text.startswith('- '):
                rows.append(parts)
                continue
            if in_matrix:
                in_matrix = False
            (after if rows else before).append(text)
        return before, rows, after

    def evidence_table_projection(lines):
        before, rows, after = [], [], []
        marker_seen = False
        for line0 in lines:
            value = governor.intercept(line0)
            if value is None:
                continue
            text = str(value).strip()
            if text.startswith('TABEL:'):
                marker_seen = True
                continue
            parts = _split_pipe(text)
            if marker_seen and len(parts) == 4 and text.startswith('['):
                rows.append(parts)
                continue
            (after if marker_seen else before).append(text)
        return before, rows, after, marker_seen

    def action_plan_projection(lines):
        rows, remainder = [], []
        for line0 in lines:
            value = governor.intercept(line0)
            if value is None:
                continue
            text = str(value).strip()
            parts = _split_pipe(text)
            if text.startswith('Langkah ') and len(parts) >= 6:
                step, window, priority, action, objective, condition = parts[:6]
                step = step.split(' ', 1)[1].strip() if step.startswith('Langkah ') else step
                condition = condition.split(':', 1)[1].strip() if condition.startswith('Kondisi:') else condition
                rows.append((step, f'{window}\n{priority}', f'{action}\n{objective}', condition))
            else:
                remainder.append(text)
        return rows, remainder

    def candidate_law_projection(lines):
        rows, remainder = [], []
        for line0 in lines:
            value = governor.intercept(line0)
            if value is None:
                continue
            text = str(value).strip()
            parts = _split_pipe(text)
            if len(parts) == 3 and parts[0].startswith('Ranah hukum:') and parts[1].startswith('Sumber:'):
                rows.append((
                    parts[0].split(':', 1)[1].strip(),
                    parts[1].split(':', 1)[1].strip(),
                    parts[2].split(':', 1)[1].strip() if ':' in parts[2] else parts[2],
                ))
            else:
                remainder.append(text)
        return rows, remainder

    meta, sections = _case_export_sections(x)
    meta, sections = prepare_case_export_for_reader(meta, sections)
    posture_label = LexiCoreCivilPresentationSanitizer.resolve_civil_posture_label(
        _presentation_source_corpus(x)
    )
    governor = LexiCoreLowLevelRenderGovernor(is_civil=bool(posture_label))

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Mm(215)
    sec.page_height = Mm(330)
    sec.top_margin = Mm(24)
    sec.left_margin = Mm(28)
    sec.bottom_margin = Mm(22)
    sec.right_margin = Mm(22)

    styles = doc.styles
    styles['Normal'].font.name = 'Times New Roman'
    styles['Normal'].font.size = Pt(10.5)
    styles['Normal'].font.color.rgb = RGBColor(31, 41, 55)
    styles['Normal'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    styles['Normal'].paragraph_format.line_spacing = 1.15
    styles['Normal'].paragraph_format.first_line_indent = Mm(0)
    styles['Normal'].paragraph_format.space_after = Pt(3)

    # Preserve LexiCore identity while using the wider, reference-inspired F4 working-paper text area.
    mast = doc.add_table(rows=1, cols=2)
    mast.alignment = WD_TABLE_ALIGNMENT.CENTER
    mast.autofit = False
    mast.columns[0].width = Mm(20)
    mast.columns[1].width = Mm(145)
    left, right = mast.rows[0].cells
    left.width = Mm(20)
    right.width = Mm(145)
    _docx_set_cell_shading(left, 'D8B65C')
    _docx_set_cell_shading(right, '10263D')
    for c in (left, right):
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _docx_set_cell_margins(c, 120, 120, 120, 120)
    lp = left.paragraphs[0]
    lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lp.paragraph_format.first_line_indent = Mm(0)
    lr = lp.add_run('LC\n' + (profile.get('watermark_text') or profile['display_name'])[:10])
    lr.bold = True
    lr.font.name = 'Times New Roman'
    lr.font.size = Pt(12)
    lr.font.color.rgb = RGBColor(16, 38, 61)
    rp = right.paragraphs[0]
    rp.paragraph_format.first_line_indent = Mm(0)
    rp.paragraph_format.space_after = Pt(0)
    rr = rp.add_run('LexiCore')
    rr.bold = True
    rr.font.name = 'Times New Roman'
    rr.font.size = Pt(17)
    rr.font.color.rgb = RGBColor(255, 255, 255)
    br = rp.add_run(f"\nEvidence-to-Action Case Analysis | {profile['display_name']}")
    br.font.name = 'Times New Roman'
    br.font.size = Pt(7.2)
    br.font.color.rgb = RGBColor(223, 229, 235)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.first_line_indent = Mm(0)
    spacer.paragraph_format.space_after = Pt(0)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.first_line_indent = Mm(0)
    t.paragraph_format.space_after = Pt(2)
    tr = t.add_run(governor.sanitize(_export_scalar(x.get('title') or 'Case Analysis')))
    tr.bold = True
    tr.font.name = 'Times New Roman'
    tr.font.size = Pt(16)
    tr.font.color.rgb = RGBColor(18, 39, 63)

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.paragraph_format.first_line_indent = Mm(0)
    s.paragraph_format.space_after = Pt(6)
    sr = s.add_run('KERTAS KERJA - MENUNGGU VERIFIKASI PROFESIONAL')
    sr.bold = True
    sr.font.name = 'Times New Roman'
    sr.font.size = Pt(8.5)
    sr.font.color.rgb = RGBColor(151, 112, 24)

    mt = doc.add_table(rows=0, cols=2)
    mt.alignment = WD_TABLE_ALIGNMENT.CENTER
    mt.autofit = False
    mt.columns[0].width = Mm(38)
    mt.columns[1].width = Mm(127)
    for k, v in meta:
        cells = mt.add_row().cells
        cells[0].width = Mm(38)
        cells[1].width = Mm(127)
        _docx_set_cell_shading(cells[0], 'EEF2F6')
        _docx_set_cell_shading(cells[1], 'FFFFFF')
        for c in cells:
            _docx_set_cell_margins(c, 70, 100, 70, 100)
        p0 = cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p0.paragraph_format.first_line_indent = Mm(0)
        p0.paragraph_format.line_spacing = 1.0
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(governor.sanitize(str(k)))
        _set_run_body_font(r0, bold=True, size=9)
        r0.font.color.rgb = RGBColor(73, 89, 105)
        p1 = cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p1.paragraph_format.first_line_indent = Mm(0)
        p1.paragraph_format.line_spacing = 1.0
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(governor.sanitize(_export_scalar(v)))
        _set_run_body_font(r1, size=9.2)

    for n, (heading, lines) in enumerate(sections, 1):
        governor.enter_section(heading)
        rendered_heading = governor.sanitize(heading)
        _docx_add_section_title(doc, n, rendered_heading)

        if heading in {'Case Readiness Review', 'Kesiapan Analisis'}:
            vals = {}
            disclaimer = None
            for line in lines:
                text = governor.intercept(line)
                if text is None:
                    continue
                text = str(text)
                if ':' in text and '%' in text:
                    k, v = text.split(':', 1)
                    vals[k.strip()] = v.strip()
                elif text:
                    disclaimer = text
            rt = doc.add_table(rows=1, cols=4)
            rt.alignment = WD_TABLE_ALIGNMENT.CENTER
            rt.autofit = True
            mapping = [
                ('Kesiapan keseluruhan', 'Keseluruhan'),
                ('Pemetaan bukti', 'Bukti'),
                ('Analisis Hukum', 'Hukum'),
                ('Rencana tindakan', 'Tindakan'),
            ]
            for idx, (key, label) in enumerate(mapping):
                c = rt.rows[0].cells[idx]
                _docx_set_cell_shading(c, 'F7F9FB' if idx else 'EAF6F0')
                _docx_set_cell_margins(c, 90, 70, 90, 70)
                p = c.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Mm(0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(1)
                r = p.add_run(governor.sanitize(vals.get(key, '0%')))
                _set_run_body_font(r, bold=True, size=10.5)
                r.font.color.rgb = RGBColor(21, 116, 82) if idx == 0 else RGBColor(18, 39, 63)
                q = c.add_paragraph()
                q.alignment = WD_ALIGN_PARAGRAPH.CENTER
                q.paragraph_format.first_line_indent = Mm(0)
                q.paragraph_format.line_spacing = 1.0
                q.paragraph_format.space_after = Pt(0)
                qr = q.add_run(governor.sanitize(label))
                _set_run_body_font(qr, size=8)
                qr.font.color.rgb = RGBColor(91, 100, 113)
            if disclaimer:
                dp = doc.add_paragraph()
                _apply_legal_body_format(dp, indent=False)
                dr = dp.add_run(governor.sanitize(disclaimer))
                _set_run_body_font(dr, size=9.5)
                dr.italic = True
                dr.font.color.rgb = RGBColor(100, 108, 118)
            governor.enter_section('')
            continue

        normalized_heading = str(rendered_heading or '').strip().upper()

        if normalized_heading == 'RINGKASAN ANALISIS HUKUM':
            before, rows, after = summary_matrix_projection(lines)
            for text in before:
                if text:
                    p = doc.add_paragraph()
                    _apply_legal_body_format(p, indent=False)
                    _set_run_body_font(p.add_run(text), size=10.5)
            if rows:
                p = doc.add_paragraph()
                _apply_legal_body_format(p, indent=False)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(2)
                _set_run_body_font(p.add_run('Matriks temuan utama'), bold=True, size=10.5)
                _docx_add_reference_table(
                    doc,
                    ['Aspek', 'Posisi/teks dokumen', 'Temuan', 'Rekomendasi'],
                    rows,
                    [32, 34, 49, 50],
                )
            lines = after

        elif normalized_heading == 'PEMETAAN BUKTI':
            before, rows, after, marker_seen = evidence_table_projection(lines)
            for text in before:
                if text:
                    p = doc.add_paragraph()
                    _apply_legal_body_format(p, indent=False)
                    _set_run_body_font(p.add_run(text), size=10.5)
            if rows:
                _docx_add_reference_table(
                    doc,
                    ['Sumber/Bukti Potensial', 'Proposisi Faktual yang Perlu Diuji', 'Status Pembuktian', 'Uji Lanjut'],
                    rows,
                    [38, 62, 30, 35],
                )
            elif marker_seen:
                p = doc.add_paragraph()
                _apply_legal_body_format(p, indent=False)
                _set_run_body_font(p.add_run('TABEL: [Sumber/Bukti Potensial] | [Proposisi Faktual yang Perlu Diuji] | [Status Pembuktian] | [Uji Lanjut]'), size=8.5)
            lines = after

        elif normalized_heading == 'RENCANA TINDAKAN':
            rows, remainder = action_plan_projection(lines)
            if rows:
                _docx_add_reference_table(
                    doc,
                    ['Langkah', 'Waktu/Prioritas', 'Tindakan dan Tujuan', 'Kondisi'],
                    rows,
                    [18, 28, 92, 27],
                )
            lines = remainder

        elif normalized_heading in {'KANDIDAT DASAR HUKUM YANG PERLU DIVERIFIKASI', 'DASAR HUKUM TERVERIFIKASI'}:
            rows, remainder = candidate_law_projection(lines)
            if rows:
                _docx_add_reference_table(
                    doc,
                    ['Ranah Hukum', 'Instrumen / Sumber', 'Status'],
                    rows,
                    [38, 72, 55],
                )
            lines = remainder

        for line in lines:
            intercepted = governor.intercept(line)
            if intercepted is None:
                continue
            text = str(intercepted).strip()
            if not text:
                continue

            if text.endswith(':') and len(text) < 80:
                p = doc.add_paragraph()
                _apply_legal_body_format(p, indent=False)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(text[:-1])
                _set_run_body_font(r, bold=True, size=10.5)
                r.font.color.rgb = RGBColor(43, 58, 73)
            elif text.startswith('- '):
                p = doc.add_paragraph(style='List Bullet')
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.first_line_indent = Mm(0)
                p.paragraph_format.left_indent = Mm(6)
                p.paragraph_format.line_spacing = 1.1
                p.paragraph_format.space_after = Pt(2)
                _set_run_body_font(p.add_run(text[2:]), size=10.5)
            elif re.match(r'^\[(P\d|SOURCE|ADMISSION|DENIAL|ALLEGATION)', text, re.I) or is_structured_line(text):
                p = doc.add_paragraph()
                _apply_legal_body_format(p, indent=False)
                _set_run_body_font(p.add_run(text), size=10.5)
            else:
                p = doc.add_paragraph()
                _apply_legal_body_format(p, indent=True)
                _set_run_body_font(p.add_run(text), size=10.5)
        governor.enter_section('')

    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.first_line_indent = Mm(0)
    hr = hp.add_run(f"LEXICORE | {profile['display_name'].upper()}")
    _set_run_body_font(hr, bold=True, size=8)
    hr.font.color.rgb = RGBColor(93, 105, 118)

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.first_line_indent = Mm(0)
    fr = fp.add_run('Kertas Kerja Rahasia | Menunggu Verifikasi Profesional | Halaman ')
    _set_run_body_font(fr, size=8)
    fr.font.color.rgb = RGBColor(93, 105, 118)
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    fp._p.append(fld)

    doc.core_properties.title = 'LexiCore Case Analysis - ' + governor.sanitize(str(x.get('title') or 'Case Analysis'))
    doc.core_properties.author = profile["signatory_name"]
    doc.core_properties.subject = 'Evidence-to-Action Legal Working Paper'

    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio
