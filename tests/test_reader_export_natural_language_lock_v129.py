from exporters.common import LexiCoreLowLevelRenderGovernor


def test_v129_branding_is_immutable_at_low_level_governor():
    g = LexiCoreLowLevelRenderGovernor()
    brand = "Evidence-to-Action Case Analysis | Initiated by ELF - Erfan's Law Firm"
    assert g.sanitize(brand) == brand
    assert g.sanitize("Evidence-to-Action Legal Working Paper") == "Evidence-to-Action Legal Working Paper"


def test_v129_natural_language_normalization_removes_mechanical_artifacts():
    g = LexiCoreLowLevelRenderGovernor()
    raw = (
        "kerugian nyata (actual loss); kontribusi pihak lain/intervening acts; "
        "rantai keputusan-akibat/intervening act belum terbukti; recovery; "
        "berdasarkan record yang tersedia; melalui tempus engine; "
        "tidak di-hard-code sebagai dasar final; sumber non-pleading"
    )
    out = g.sanitize(raw)
    assert "kerugian nyata (kerugian nyata)" not in out
    assert "kerugian nyata" in out
    assert "kontribusi atau tindakan pihak lain yang memengaruhi hubungan kausal" in out
    assert "rantai keputusan-akibat serta faktor atau tindakan perantara" in out
    assert "recovery" not in out.lower()
    assert "pemulihan kerugian/aset" in out
    assert "berdasarkan data yang tersedia" in out
    assert "mekanisme verifikasi waktu berlaku" in out
    assert "tidak ditetapkan secara tetap sebagai dasar final" in out
    assert "sumber di luar dokumen argumentasi para pihak" in out


def test_v129_pdf_masthead_preserves_branding(monkeypatch):
    import exporters.case_pdf as mod
    from pypdf import PdfReader

    monkeypatch.setattr(mod, "_case_export_sections", lambda x: ([(("Judul"), "Case Analysis")], []))
    monkeypatch.setattr(mod, "prepare_case_export_for_reader", lambda m, s: (m, s))
    bio = mod.export_case_pdf({"title": "Case Analysis", "source_text": ""})
    text = "\n".join(page.extract_text() or "" for page in PdfReader(bio).pages)
    assert "Evidence-to-Action Case Analysis" in text
    assert "analisis bukti ke tindakan Case Analysis" not in text


def test_v129_docx_masthead_preserves_branding(monkeypatch):
    import exporters.case_docx as mod
    from docx import Document

    monkeypatch.setattr(mod, "_case_export_sections", lambda x: ([(("Judul"), "Case Analysis")], []))
    monkeypatch.setattr(mod, "prepare_case_export_for_reader", lambda m, s: (m, s))
    bio = mod.export_case_docx({"title": "Case Analysis", "source_text": ""})
    doc = Document(bio)
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.extend(p.text for p in cell.paragraphs)
    text = "\n".join(chunks)
    assert "Evidence-to-Action Case Analysis" in text
    assert "analisis bukti ke tindakan Case Analysis" not in text
