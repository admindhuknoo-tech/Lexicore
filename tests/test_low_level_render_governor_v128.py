from io import BytesIO

from exporters.common import LexiCoreLowLevelRenderGovernor


def test_governor_sanitizes_final_residue_at_boundary():
    g = LexiCoreLowLevelRenderGovernor(is_civil=False)
    raw = (
        "case keterkaitan | candidate result | UNKNOWN_DATE | actual loss | "
        "intervening acts | impact | existing_4 | outstanding principal | "
        "CREDIT, FIDUCIARY, LOSS, RESPONSIBILITY | LEGAL_VERIFICATION_REQUIRED"
    )
    out = g.sanitize(raw)
    for residue in (
        "case keterkaitan", "candidate result", "UNKNOWN_DATE", "actual loss",
        "intervening acts", "existing_4", "outstanding principal",
        "LEGAL_VERIFICATION_REQUIRED",
    ):
        assert residue not in out
    assert "keterkaitan materi perkara" in out
    assert "Hasil Kandidat Hukum" in out
    assert "Tanggal Belum Terverifikasi" in out
    assert "kerugian nyata" in out
    assert "faktor atau tindakan intervensi pihak lain" in out
    assert "elemen keterkaitan aktif" in out
    assert "sisa pokok kewajiban" in out


def test_governor_deduplicates_only_immediately_repeated_actions_inside_action_section():
    g = LexiCoreLowLevelRenderGovernor()
    g.enter_section("Rantai Penalaran Hukum Kanonik")
    first = g.intercept("TINDAKAN: Dapatkan dan uji bukti primer yang sama.")
    second = g.intercept("TINDAKAN:   Dapatkan   dan uji bukti primer yang sama.")
    assert first is not None
    assert second is None

    # A non-action line resets the single-line memory ledger.
    assert g.intercept("BUKTI: belum tersedia") is not None
    assert g.intercept("TINDAKAN: Dapatkan dan uji bukti primer yang sama.") is not None

    # Leaving the action section also resets it.
    g.enter_section("Audit Dokumen Terperinci")
    assert g.intercept("TINDAKAN: Dapatkan dan uji bukti primer yang sama.") is not None


def test_governor_civil_swaps_are_gated():
    criminal = LexiCoreLowLevelRenderGovernor(is_civil=False)
    civil = LexiCoreLowLevelRenderGovernor(is_civil=True)
    raw = "surat dakwaan | teks dakwaan | didakwakan | Pengadilan Tipikor"
    assert "surat dakwaan" in criminal.sanitize(raw)
    out = civil.sanitize(raw)
    assert "surat dakwaan" not in out
    assert "surat gugatan" in out
    assert "posita gugatan" in out
    assert "digugat atau disengketakan" in out
    assert "Pengadilan Negeri" in out


def _fake_export_payload():
    meta = [("Acuan waktu", "UNKNOWN_DATE")]
    sections = [
        ("Rantai Penalaran Hukum Kanonik", [
            "RANTAI E3 | Status: HOLD",
            "TINDAKAN: Dapatkan dan uji bukti primer yang secara langsung menjawab kekosongan pembuktian ini.",
            "TINDAKAN:   Dapatkan   dan uji bukti primer yang secara langsung menjawab gap ini.",
            "Acuan waktu: UNKNOWN_DATE | case keterkaitan | candidate result | actual loss | intervening acts | existing_4",
        ])
    ]
    return meta, sections


def test_pdf_actual_renderer_boundary_intercepts_and_deduplicates(monkeypatch):
    import exporters.case_pdf as mod
    from pypdf import PdfReader

    payload = _fake_export_payload()
    monkeypatch.setattr(mod, "_case_export_sections", lambda x: payload)
    # Deliberately bypass the earlier presentation pass: this proves the actual
    # ReportLab write loop is protected independently.
    monkeypatch.setattr(mod, "prepare_case_export_for_reader", lambda m, s: (m, s))

    bio = mod.export_case_pdf({"title": "Case Analysis", "source_text": ""})
    text = "\n".join(page.extract_text() or "" for page in PdfReader(bio).pages)

    normalized = " ".join(text.split())
    assert normalized.count("Dapatkan dan uji bukti primer yang secara langsung menjawab kekosongan pembuktian ini.") == 1
    for residue in ("UNKNOWN_DATE", "case keterkaitan", "candidate result", "actual loss", "intervening acts", "existing_4"):
        assert residue not in text
    assert "Tanggal Belum Terverifikasi" in text


def test_docx_actual_renderer_boundary_intercepts_and_deduplicates(monkeypatch):
    import exporters.case_docx as mod
    from docx import Document

    payload = _fake_export_payload()
    monkeypatch.setattr(mod, "_case_export_sections", lambda x: payload)
    monkeypatch.setattr(mod, "prepare_case_export_for_reader", lambda m, s: (m, s))

    bio = mod.export_case_docx({"title": "Case Analysis", "source_text": ""})
    doc = Document(bio)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    text = "\n".join(parts)

    assert text.count("Dapatkan dan uji bukti primer yang secara langsung menjawab kekosongan pembuktian ini.") == 1
    for residue in ("UNKNOWN_DATE", "case keterkaitan", "candidate result", "actual loss", "intervening acts", "existing_4"):
        assert residue not in text
    assert "Tanggal Belum Terverifikasi" in text
