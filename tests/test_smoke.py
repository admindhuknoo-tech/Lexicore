"""LexiCore API smoke tests.

Hardened for the v1.3.4 stability baseline. Every bug found across
the v1.3.2 and v1.3.3 audits (stale client_communications schema, contract
party-extraction regex, PDF export silently dropping lines, redundant
Evidence-to-Action calls) was only caught by manually exercising every
endpoint by hand after the fact. This file automates that sweep so the same
class of regression is caught before a new zip ships, not after.

IMPORTANT — isolation from real data:
This suite points LEXICORE_DB_PATH at a throwaway SQLite file in a temp
directory. It never touches the firm's real lexicore.db, and it
deliberately exercises the destructive endpoints (/api/history DELETE,
purge-all) — which would be dangerous to run against real client data.

The env vars are actually set in tests/conftest.py, not here (v1.3.5.4 fix):
database.py reads LEXICORE_DB_PATH exactly once, at first module import, so
setting it here only protected the suite if this file happened to be the
first test module pytest imported. tests/test_migrations.py imports
`database` too and sorts before this file alphabetically, so pytest's
default collection order silently defeated that guard and let real
migrations run against the real lexicore.db. conftest.py runs before any
test module in this directory is collected, so it's the only placement that
is actually guaranteed safe. The lines below are kept as a harmless,
redundant safety net (setdefault, so they never override conftest.py) in
case this file is ever executed directly rather than via pytest.

Run with:
    py -m pip install pytest --break-system-packages   # if not installed
    py -m pytest -q
"""
import os
import io
import json
import tempfile

# Redundant safety net only — see the module docstring above. The real
# guarantee comes from tests/conftest.py, which pytest always loads first.
os.environ.setdefault("LEXICORE_DB_PATH", os.path.join(tempfile.mkdtemp(prefix="lexicore_test_"), "test_lexicore.db"))
os.environ.setdefault("LEXICORE_DISABLE_AUTO_BACKUP", "1")
os.environ.setdefault("LEXICORE_BACKUP_DIR", os.path.join(tempfile.mkdtemp(prefix="lexicore_test_backups_"), "backups"))
os.environ.setdefault("LEXICORE_AI_PROVIDER", "local")
os.environ.setdefault("LEXICORE_AI_MODE", "local")
os.environ.setdefault("GEMINI_API_KEY", "")
os.environ.setdefault("OPENAI_API_KEY", "")

import pytest
import app as appmod  # noqa: E402  (import must follow the env var above)
from contract_review import RiskDetector  # noqa: E402
from database import EXPECTED_COLUMNS, get_db_connection, SCHEMA_VERSION  # noqa: E402
from regulatory_db import get_all_regulations, search_regulations  # noqa: E402


@pytest.fixture()
def client():
    appmod.app.testing = True
    return appmod.app.test_client()


def post_json(client, url, payload):
    return client.post(url, data=json.dumps(payload), content_type="application/json")


def _sample_docx_bytes(paragraphs):
    """Build a minimal real .docx in memory for upload-based endpoints."""
    from docx import Document
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio


# --------------------------------------------------------------------------
# Read-only / status endpoints
# --------------------------------------------------------------------------

def test_health(client):
    r = client.get("/api/health")
    assert r.status_code == 200
    assert r.get_json()["status"] == "online"


def test_ai_status(client):
    assert client.get("/api/ai/status").status_code == 200


def test_dashboard_metrics(client):
    r = client.get("/api/dashboard/metrics")
    assert r.status_code == 200
    assert "metrics" in r.get_json()


def test_regulations_catalog_and_search(client):
    assert client.get("/api/regulations/catalog").status_code == 200
    assert client.get("/api/regulations/search?q=korupsi").status_code == 200


def test_regulatory_intelligence_catalog_graph_timeline(client):
    assert client.get("/api/regulatory-intelligence/catalog").status_code == 200
    assert client.get("/api/regulatory-intelligence/graph?q=korupsi").status_code == 200
    assert client.get("/api/regulatory-intelligence/timeline?q=korupsi").status_code == 200


def test_legal_sources(client, monkeypatch):
    assert client.get("/api/legal-sources").status_code == 200
    # Release gate must be deterministic/offline: stub the external health
    # probe while still testing the Flask endpoint shape.
    monkeypatch.setattr(appmod, "source_health", lambda: [])
    r = client.get("/api/legal-sources/health")
    assert r.status_code == 200
    assert r.get_json()["count"] == 0


def test_audit_logs_empty_ok(client):
    assert client.get("/api/audit/logs").status_code == 200


# --------------------------------------------------------------------------
# Legal Drafting lifecycle: generate -> save -> get -> update -> export -> history
# --------------------------------------------------------------------------

def test_draft_lifecycle(client):
    r = post_json(client, "/api/generate/draft", {
        "doc_type": "Somasi", "party1": "PT A", "party2": "PT B",
        "effective_date": "2026-09-03", "duration": "", "prompt": "wanprestasi pembayaran",
    })
    assert r.status_code == 200
    content = r.get_json()["content"]
    assert len(content) > 50

    r = post_json(client, "/api/drafts", {
        "title": "Test Draft", "doc_type": "Somasi", "party1": "PT A", "party2": "PT B",
        "content": content, "word_count": len(content.split()), "clause_count": 1, "status": "saved",
    })
    assert r.status_code == 200
    draft_id = r.get_json()["draft_id"]
    assert draft_id

    assert client.get(f"/api/drafts/{draft_id}").status_code == 200

    r = client.put(f"/api/drafts/{draft_id}", data=json.dumps({"title": "Updated Title"}),
                    content_type="application/json")
    assert r.status_code == 200

    r = client.get(f"/api/drafts/{draft_id}/export/docx")
    assert r.status_code == 200
    assert r.headers["Content-Type"].startswith("application/vnd.openxmlformats")

    r = client.get(f"/api/history/drafts/{draft_id}")
    assert r.status_code == 200

    r = client.delete(f"/api/history/drafts/{draft_id}")
    assert r.status_code == 200
    assert client.get(f"/api/history/drafts/{draft_id}").status_code == 404


# --------------------------------------------------------------------------
# Contract Review — regression guard for the party-extraction bug
# --------------------------------------------------------------------------

def test_contract_review_party_extraction_does_not_pick_up_garbage(client):
    bio = _sample_docx_bytes([
        "PERJANJIAN KERJASAMA",
        "Pihak Pertama: PT Sumber Makmur, Pihak Kedua: Budi Santoso.",
        "Pihak Kedua wajib membayar denda sebesar Rp 50.000.000 apabila terjadi "
        "keterlambatan pembayaran lebih dari 30 hari.",
        "Pihak Pertama dapat mengakhiri perjanjian ini secara sepihak tanpa "
        "pemberitahuan sebelumnya.",
    ])
    r = client.post("/api/review", data={"file": (bio, "sample.docx")},
                     content_type="multipart/form-data")
    assert r.status_code == 200
    parties = r.get_json()["data"]["parties"]
    # Regression guard: previously this returned a truncated sentence fragment
    # ("dapat mengakhiri perjanjian ini secara sepihak tan...") as a "party".
    assert "PT Sumber Makmur" in parties
    assert "Budi Santoso" in parties
    for p in parties:
        assert "mengakhiri" not in p.lower(), f"garbage party name leaked back in: {p!r}"

    assert client.get("/api/analyses").status_code == 200


# --------------------------------------------------------------------------
# Legal Research
# --------------------------------------------------------------------------

def test_research_summarize_and_list(client):
    r = post_json(client, "/api/research/summarize", {
        "source_text": "Wanprestasi dalam perjanjian jual beli terjadi ketika salah satu "
                        "pihak tidak memenuhi kewajibannya sesuai perjanjian yang disepakati.",
    })
    assert r.status_code == 200
    assert client.get("/api/research").status_code == 200


# --------------------------------------------------------------------------
# Compliance & Risk
# --------------------------------------------------------------------------

def test_compliance_assess_and_list(client):
    r = post_json(client, "/api/compliance/assess", {
        "business_type": "PT", "sector": "fintech", "activities": "pinjaman online",
    })
    assert r.status_code == 200
    assert client.get("/api/compliance").status_code == 200


# --------------------------------------------------------------------------
# Case Analysis — including export, and regression guard for the
# PDF KeepTogether content-loss bug
# --------------------------------------------------------------------------

def test_case_analysis_and_exports(client):
    narrative = (
        "Karyawan diputus hubungan kerja tanpa pesangon oleh PT Maju Jaya pada "
        "1 Januari 2026 setelah bekerja selama lima tahun tanpa surat peringatan "
        "sebelumnya. Perusahaan mendalilkan efisiensi namun tidak ada dokumen "
        "pendukung. Karyawan menuntut pesangon, uang penghargaan masa kerja, dan "
        "uang penggantian hak sesuai UU Ketenagakerjaan."
    )
    r = post_json(client, "/api/case-analysis", {"title": "Kasus PHK Uji", "narrative": narrative, "verify_online": False})
    assert r.status_code == 200
    x = r.get_json()["data"]
    assert "case_readiness" in x
    assert "case_working_paper" in x
    assert x["case_working_paper"].get("working_paper_percentage")
    assert x["case_working_paper"].get("evidence_map")
    assert x["case_working_paper"].get("legal_construction")
    assert x["case_working_paper"].get("action_plan")
    assert x.get("case_analysis_id")
    assert x.get("analysis_provenance", {}).get("mode") == "LOCAL_DETERMINISTIC"
    assert x.get("analysis_provenance", {}).get("ai_available_at_run") is False

    assert client.get("/api/case-analysis").status_code == 200

    r = post_json(client, "/api/case-analysis/export/pdf", x)
    assert r.status_code == 200
    assert r.headers["Content-Type"] == "application/pdf"
    pdf_bytes = r.get_data()
    assert len(pdf_bytes) > 2000
    # Regression guard: the KeepTogether bug used to silently drop the 2nd and
    # 3rd line of every section longer than 3 lines. Extract the text back out
    # and confirm known distinctive phrases from the narrative/analysis
    # actually made it into the PDF, not just page 1 boilerplate.
    try:
        from pypdf import PdfReader
        text = "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(pdf_bytes)).pages)
        assert "PT Maju Jaya" in text or "pesangon" in text
    except ImportError:
        pass  # pypdf not installed in this environment; byte-length check above still applies

    r = post_json(client, "/api/case-analysis/export/docx", x)
    assert r.status_code == 200
    assert r.headers["Content-Type"].startswith("application/vnd.openxmlformats")

    r = post_json(client, "/api/case-analysis/export/xyz", x)
    assert r.status_code == 400

    r = post_json(client, "/api/case-analysis/export/pdf", {})
    assert r.status_code == 400


# --------------------------------------------------------------------------
# Norm Conflicts
# --------------------------------------------------------------------------

def test_norm_conflicts(client):
    r = post_json(client, "/api/norm-conflicts", {
        "provisions": ["UU Tipikor Pasal 3", "KUHP Pasal 604"],
        "facts_context": "Dugaan korupsi dengan unsur perdata/wanprestasi",
    })
    assert r.status_code == 200
    data = r.get_json()["data"]
    assert "conflicts_detected" in data
    assert "relationships_only" in data


# --------------------------------------------------------------------------
# Client Communication — regression guard for the stale-schema bug
# --------------------------------------------------------------------------

def test_client_communication_generate_save_list(client):
    r = post_json(client, "/api/communication/generate", {
        "client_name": "Budi", "document_type": "client_update", "matter": "sengketa waris",
    })
    assert r.status_code == 200

    r = post_json(client, "/api/communications", {
        "client_name": "Budi", "client_email": "", "subject": "Update Perkara",
        "message": "Pesan pembaruan untuk klien", "document_type": "client_update", "status": "draft",
    })
    assert r.status_code == 200
    comm_id = r.get_json()["communication_id"]

    r = client.get("/api/communications")
    assert r.status_code == 200
    rows = r.get_json()["data"]
    assert any(row["id"] == comm_id for row in rows)


# --------------------------------------------------------------------------
# History cleanup endpoints (safe here — isolated test DB, not the real one)
# --------------------------------------------------------------------------

def test_history_clear_kind_and_purge_all(client):
    post_json(client, "/api/drafts", {
        "title": "Throwaway", "doc_type": "Somasi", "content": "isi " * 20,
        "word_count": 5, "clause_count": 0, "status": "saved",
    })
    r = client.delete("/api/history/drafts")
    assert r.status_code == 200
    assert r.get_json()["deleted"] >= 1

    r = client.delete("/api/history", data=json.dumps({"confirm": "salah"}),
                       content_type="application/json")
    assert r.status_code == 400

    r = client.delete("/api/history", data=json.dumps({"confirm": "HAPUS SEMUA RIWAYAT"}),
                       content_type="application/json")
    assert r.status_code == 200

    assert client.get("/api/drafts").get_json()["data"] == []


if __name__ == "__main__":
    import sys
    sys.exit(pytest.main([__file__, "-v"]))


# --------------------------------------------------------------------------
# Stability baseline guards
# --------------------------------------------------------------------------

def test_schema_matches_expected_columns():
    conn = get_db_connection()
    try:
        version_row = conn.execute("SELECT value FROM schema_meta WHERE key='schema_version'").fetchone()
        assert version_row and int(version_row[0]) == SCHEMA_VERSION
        for table, expected in EXPECTED_COLUMNS.items():
            existing = {row[1] for row in conn.execute(f"PRAGMA table_info({table})").fetchall()}
            missing = [name for name, _ in expected if name not in existing]
            assert not missing, f"schema drift in {table}: {missing}"
    finally:
        conn.close()


def test_contract_risk_detects_semantic_unilateral_termination():
    text = "Pihak Pertama dapat mengakhiri perjanjian ini secara sepihak tanpa pemberitahuan sebelumnya."
    risks = RiskDetector.analyze(text)
    assert any(r.category == 'Terminasi Sepihak' and r.risk_level == 'HIGH' for r in risks)


def test_regulatory_corpus_has_priority_domains():
    regs = get_all_regulations()
    assert len(regs) >= 10
    for query in ('ketenagakerjaan phk pesangon', 'perseroan direksi', 'bpr tata kelola kredit'):
        hits = search_regulations(query, limit=5)
        assert hits, f"no regulatory hit for {query}"

def test_health_version_is_centralized(client):
    from version import LEXICORE_VERSION
    assert client.get('/api/health').get_json()['version'] == LEXICORE_VERSION

# --------------------------------------------------------------------------
# v1.3.5 Regulatory Corpus Expansion / relevance regression tests
# --------------------------------------------------------------------------

def test_regulatory_corpus_expanded_and_structured():
    from regulatory_db import corpus_stats
    stats = corpus_stats()
    assert stats["regulations"] >= 39
    assert stats["articles"] >= 90
    assert stats["domains"] >= 20


def test_labor_query_does_not_rank_tipikor_as_primary_result():
    hits = search_regulations("PHK pesangon PKWT ketenagakerjaan", limit=5)
    assert hits
    ids = [h["regulation"]["id"] for h in hits]
    assert ids[0] in {"uu_ketenagakerjaan_13_2003_6_2023", "pp_35_2021", "uu_pphi_2_2004"}
    assert "uu_tipikor_31_1999" not in ids[:3]


def test_bpr_temporal_corpus_contains_pre_and_post_event_rules():
    from regulatory_db import get_regulation_by_id
    pre = get_regulation_by_id("pojk_4_2015_tata_kelola_bpr")
    post = get_regulation_by_id("pojk_23_2022_bmpk_bpr")
    current = get_regulation_by_id("pojk_9_2024_tata_kelola_bpr")
    assert pre and pre["effective_date"] == "2015-04-01"
    assert post and post["effective_date"] == "2022-11-23"
    assert current and current["effective_date"] == "2024-07-01"
    assert "September 2022" in post.get("temporal_note", "")


def test_regulatory_intelligence_reports_expanded_counts(client):
    r = client.get("/api/regulatory-intelligence/catalog")
    assert r.status_code == 200
    data = r.get_json()
    # Endpoint shape may wrap the catalog. Search recursively without coupling
    # the test to presentation-only response nesting.
    def find_counts(obj):
        if isinstance(obj, dict):
            if "counts" in obj and isinstance(obj["counts"], dict):
                return obj["counts"]
            for v in obj.values():
                found = find_counts(v)
                if found:
                    return found
        return None
    counts = find_counts(data)
    assert counts and counts.get("regulations", 0) >= 39 and counts.get("articles", 0) >= 90

# --------------------------------------------------------------------------
# v1.3.5.3 BAP benchmark / full-document pipeline corrective
# --------------------------------------------------------------------------

def test_full_document_chunk_plan_never_reports_zero_for_large_text():
    from ai_engine import chunk_diagnostics
    diag = chunk_diagnostics(("Dokumen pemeriksaan kredit BPR tahun 2022. " * 3500))
    assert diag["characters"] > 100000
    assert diag["segments_total"] > 0
    assert all(n > 0 for n in diag["segment_lengths"])


def test_case_applicable_law_prunes_orphan_and_irrelevant_domains():
    result = {
        "case_regulatory_snapshot": {"domains": [{"id":"financial_services"},{"id":"corruption"}]},
        "applicable_law": [
            {"domain":"Norma disebut dalam dokumen","source":"Pasal 603"},
            {"domain":"Norma disebut dalam dokumen","source":"Pasal 3J"},
            {"domain":"Ketenagakerjaan","source":"Peraturan ketenagakerjaan"},
            {"domain":"Pidana Materiil","source":"UU No. 1 Tahun 2023 tentang KUHP"},
        ],
    }
    cleaned = appmod._clean_applicable_law(result)
    sources = [x["source"] for x in cleaned]
    assert "Pasal 603" not in sources
    assert "Pasal 3J" not in sources
    assert all(x["domain"] != "Ketenagakerjaan" for x in cleaned)
    assert "UU No. 1 Tahun 2023 tentang KUHP" in sources


def test_norm_conflict_does_not_call_cross_domain_coexistence_a_conflict():
    from norm_conflict import analyze_conflicts
    out = analyze_conflicts(["UU Tipikor", "KUHPerdata"], "perkara kredit BPR perjanjian dan dugaan korupsi tahun 2022", [])
    assert not any(c.get("type") == "LEX_SPECIALIS_CANDIDATE" for c in out.get("conflicts_detected", []))
    assert any(r.get("classification") == "RELATIONSHIP_ONLY" for r in out.get("relationships_only", []))


def test_export_uses_compact_regulatory_section_and_executive_summary():
    from exporters.common import _case_export_sections
    sample = {
        "title":"Benchmark BAP", "analytical_method":"DETERMINISTIC_FALLBACK",
        "executive_summary":["Fakta kunci: fasilitas kredit BPR pada 2022.", "Isu utama: actual loss dan mens rea."],
        "case_regulatory_snapshot":{"event_year_candidate":2022,"event_date_candidate":"2022-09-27"},
        "regulatory_matches":[
            {"regulation":{"nomor":"POJK No. 23 Tahun 2022","effective_date":"2022-11-23"},"matched_articles":[]},
        ],
        "action_plan":[],
    }
    _, sections = _case_export_sections(sample)
    titles=[x[0] for x in sections]
    assert "Ringkasan Eksekutif" in titles
    assert "Regulasi & Tempus" in titles
    assert "Korpus Regulasi & Regulatory Intelligence" not in titles
    reg=dict(sections)["Regulasi & Tempus"]
    assert any("Terbit setelah peristiwa — tidak dipakai sebagai dasar materiil" in line for line in reg)


def test_export_same_year_without_exact_event_date_stays_uncertain():
    from exporters.common import _case_export_sections
    sample = {
        "title":"Tempus year-only",
        "case_regulatory_snapshot":{"event_year_candidate":2022},
        "regulatory_matches":[
            {"regulation":{"nomor":"POJK No. 23 Tahun 2022","effective_date":"2022-11-23"},"matched_articles":[]},
        ],
    }
    _, sections = _case_export_sections(sample)
    reg=dict(sections)["Regulasi & Tempus"]
    assert any("Tanggal perbuatan harus dipastikan" in line for line in reg)
    assert not any("Terbit setelah peristiwa — tidak dipakai sebagai dasar materiil" in line for line in reg)


def test_tempus_engine_keeps_machine_status_separate_from_lawyer_facing_label():
    from exporters.common import _tempus_status, _human_release_status
    post, _ = _tempus_status(
        {"nomor":"POJK No. 23 Tahun 2022","effective_date":"2022-11-23"},
        {"event_date_candidate":"2022-09-27"},
    )
    same_year, _ = _tempus_status(
        {"nomor":"POJK No. 23 Tahun 2022","effective_date":"2022-11-23"},
        {"event_year_candidate":2022},
    )
    assert post == "POST_TEMPUS_EXCLUDED"
    assert same_year == "TEMPUS_REQUIRES_EXACT_DATE"
    assert _human_release_status(post) == "Terbit setelah peristiwa — tidak dipakai sebagai dasar materiil"
    assert _human_release_status(same_year) == "Tanggal perbuatan harus dipastikan"


def test_bap_material_date_prefers_credit_date_over_procedural_and_regulation_dates():
    from services.regulatory_retrieval import detect_material_date
    sample = (
        "Berita Acara Pemeriksaan tanggal 21 Mei 2026. "
        "Diperlihatkan Perjanjian Kredit Nomor 10130003097 tanggal 27 September 2022. "
        "POJK No. 23 Tahun 2022 berlaku 23 November 2022."
    )
    assert detect_material_date(sample) == "2022-09-27"

# --------------------------------------------------------------------------
# v1.3.5.4 Regulatory & Full-Document Stabilization
# --------------------------------------------------------------------------

def test_ocr_split_day_is_normalized_without_becoming_material_event():
    from services.regulatory_retrieval import detect_case_dates, detect_material_date
    sample=(
        "Jawaban Tergugat atas Gugatan Penggugat tertanggal 2 6 Februari 2026. "
        "Eksepsi kompetensi absolut Pengadilan Agama. "
        "Sertipikat Hak Milik berdasarkan data fisik dan yuridis tanggal 03/11/2017."
    )
    dates=detect_case_dates(sample)
    assert any(x["date"] == "2026-02-26" and x["role"] == "procedural_or_filing" for x in dates)
    assert detect_material_date(sample) is None


def test_land_civil_pleading_domains_do_not_route_banking_or_employment():
    from services.regulatory_retrieval import detect_domains
    sample=(
        "Jawaban Tergugat atas Gugatan Penggugat. Eksepsi Obscuur Libel dan Plurium Litis Consortium. "
        "Kewenangan absolut Pengadilan Agama Pasal 49. Sertipikat Hak Milik, Surat Ukur, "
        "data fisik dan data yuridis serta Kantor Pertanahan."
    )
    ids={x["id"] for x in detect_domains(sample)}
    assert {"civil_procedure","land_property","religious_court"}.issubset(ids)
    assert "financial_services" not in ids
    assert "employment" not in ids
    assert "criminal" not in ids


def test_regulatory_seed_hard_gate_removes_cross_domain_contamination():
    from services.regulatory_retrieval import filter_regulatory_matches_for_domains
    matches=[
        {"regulation":{"nomor":"POJK No. 23 Tahun 2022","tentang":"BMPK BPR"},"matched_articles":[]},
        {"regulation":{"nomor":"UU No. 13 Tahun 2003","tentang":"Ketenagakerjaan"},"matched_articles":[]},
        {"regulation":{"nomor":"Staatsblad 1847 No. 23","tentang":"Kitab Undang-Undang Hukum Perdata"},"matched_articles":[]},
    ]
    domains=[{"id":"civil_procedure"},{"id":"land_property"},{"id":"civil_contract"}]
    out=filter_regulatory_matches_for_domains(matches,domains)
    titles=' '.join((x.get('regulation') or {}).get('tentang','') for x in out).lower()
    assert 'hukum perdata' in titles
    assert 'bpr' not in titles
    assert 'ketenagakerjaan' not in titles


def test_full_document_segment_trace_accepts_valid_payload(monkeypatch):
    import ai_engine
    monkeypatch.setenv('LEXICORE_AI_PROVIDER','gemini')
    monkeypatch.setenv('LEXICORE_AI_MODEL','gemini-test')
    monkeypatch.setenv('GEMINI_API_KEY','test-key')
    def fake(prompt):
        if 'Senior Legal Reasoning Engine' in prompt:
            return {"facts":["fakta"],"legal_issues":["isu"],"legal_analysis":"analisis","evidentiary_gaps":[],"action_plan":[]}
        return {"facts":[{"statement":"fakta","evidence":"kutipan","type":"DOCUMENT_FACT"}],"source_items":[{"label":"SOURCE FACT","statement":"fakta","evidence":"kutipan","segment":"1"}]}
    monkeypatch.setattr(ai_engine,'_gemini_json_retry',fake)
    out=ai_engine.analyze_full_document('Fakta perkara. '*1000,'Test',{},None)
    assert out["document_reading"]["segments_read"] > 0
    assert all(x["state"] == "ACCEPTED" for x in out["document_reading"]["segment_trace"])


def test_full_document_failure_exposes_auditable_reason(monkeypatch):
    import ai_engine
    monkeypatch.setenv('LEXICORE_AI_PROVIDER','gemini')
    monkeypatch.setenv('LEXICORE_AI_MODEL','gemini-test')
    monkeypatch.setenv('GEMINI_API_KEY','test-key')
    def fail(_prompt):
        raise RuntimeError('Gemini HTTP 429: quota')
    monkeypatch.setattr(ai_engine,'_gemini_json_retry',fail)
    with pytest.raises(ai_engine.FullDocumentFailure) as err:
        ai_engine.analyze_full_document('Dokumen perkara. '*900,'Test',{},None)
    diag=err.value.diagnostics
    assert diag["segments_total"] > 0
    assert diag["segments_read"] == 0
    assert diag["status"] == "ALL_SEGMENTS_REJECTED"
    assert all(x["reason"] == "HTTP_429_RATE_LIMIT" for x in diag["segment_trace"])


def test_orphan_article_with_ayat_is_removed_from_applicable_law():
    result={
        "case_regulatory_snapshot":{"domains":[{"id":"civil_procedure"},{"id":"land_property"}]},
        "applicable_law":[
            {"domain":"Norma disebut dalam dokumen","source":"Pasal 19 ayat (2)"},
            {"domain":"Pertanahan & Properti","source":"UUPA dan peraturan pertanahan yang berlaku"},
            {"domain":"Ketenagakerjaan","source":"Peraturan ketenagakerjaan"},
        ],
    }
    out=appmod._clean_applicable_law(result)
    sources=[x["source"] for x in out]
    assert "Pasal 19 ayat (2)" not in sources
    assert "UUPA dan peraturan pertanahan yang berlaku" in sources
    assert all(x.get("domain") != "Ketenagakerjaan" for x in out)


def test_ai_http_error_reasons_are_specific():
    import ai_engine
    assert ai_engine._http_reason(400) == 'HTTP_400_BAD_REQUEST'
    assert ai_engine._http_reason(403) == 'HTTP_403_AUTH_OR_PERMISSION'
    assert ai_engine._http_reason(429) == 'HTTP_429_RATE_LIMIT'
    assert ai_engine._http_reason(503) == 'HTTP_503_UPSTREAM'


def test_bpr_export_prunes_incidental_bw_seed():
    from exporters.common import _case_export_sections
    sample={
        'title':'BAP BPR',
        'case_regulatory_snapshot':{'domains':[{'id':'financial_services'},{'id':'corruption'},{'id':'criminal'},{'id':'regional_government'}], 'event_date_candidate':'2022-09-27'},
        'regulatory_matches':[
            {'regulation':{'nomor':'Staatsblad 1847 No. 23','tentang':'Kitab Undang-Undang Hukum Perdata','effective_date':'1848-05-01'},'matched_articles':[{'pasal':'Pasal 1233'}]},
            {'regulation':{'nomor':'POJK No. 13/POJK.03/2015','tentang':'Penerapan Manajemen Risiko bagi Bank Perkreditan Rakyat','effective_date':'2015-11-12'},'matched_articles':[{'topic':'Risiko Kredit'}]},
        ],
        'action_plan':[],
    }
    _, sections=_case_export_sections(sample)
    reg=' '.join(dict(sections)['Regulasi & Tempus']).lower()
    assert 'pojk' in reg
    assert 'staatsblad' not in reg


def test_action_plan_dedupes_same_substantive_issue():
    import app
    issue='Laporan hasil perhitungan kerugian keuangan negara/daerah lengkap beserta metode dan tanggal cut-off'
    result={
        'facts':['Dokumen membahas fasilitas kredit Perumda BPR tahun 2022.'],
        'legal_issues':['Apakah kerugian aktual telah terbukti?'],
        'legal_analysis':'Kerugian aktual perlu diverifikasi.',
        'evidentiary_gaps':[issue],
        'evidence_needed':[issue],
        'action_plan':[],
    }
    out=app._ensure_evidence_to_action(result)
    issues=[str(a.get('issue') or '') for a in out['action_plan']]
    assert sum(1 for x in issues if x == issue) == 1


def test_bpr_executive_summary_rejects_ocr_passthrough_noise():
    import app
    result={
        'source_text':'Berita Acara Pemeriksaan Tersangka dugaan Tindak Pidana Korupsi dalam Pemberian Fasilitas Kredit oleh Perumda BPR Kota Blitar kepada Debitur Dewi Mufarida Tahun 2022. Direktur Utama.',
        'facts':['Anak Dani (Alm) SOENARKO, sebagaimana diatur Jo.Pasal 603 Jo.Pasal 18', 'Mengertikah Tersangka mengapa Tersangka dimintai keterangan saat ini?'],
        'legal_issues':['Apakah kerugian aktual telah terbukti?'],
        'legal_analysis':'Pelanggaran prosedur tidak otomatis membuktikan tindak pidana.',
        'evidentiary_gaps':['Laporan audit kerugian aktual.'],
    }
    lines=app._compact_executive_summary(result)
    text=' '.join(lines)
    assert 'Anak Dani' not in text
    assert 'Mengertikah Tersangka' not in text
    assert 'Perumda BPR' in text


# --------------------------------------------------------------------------
# Regression tests (v1.3.5.4 audit)
# --------------------------------------------------------------------------

def test_cors_is_not_open_to_arbitrary_origins(client):
    """Security regression: CORS(app) used to reflect
    Access-Control-Allow-Origin for any requesting origin. Since this API has
    no authentication, that let any website the user had open in another tab
    silently read/write confidential client data from the locally running
    server. Cross-origin access must now be opt-in only (LEXICORE_CORS_ORIGINS)."""
    r = client.get("/api/health", headers={"Origin": "https://evil-example.com"})
    assert r.status_code == 200
    assert "Access-Control-Allow-Origin" not in r.headers


def _corrupt_pdf_bytes():
    return io.BytesIO(b"%PDF-1.4 this is not a real pdf structure at all")


def test_contract_review_rejects_corrupt_pdf_with_json_error(client):
    """Bug regression: corrupt/unreadable uploads must return a clean JSON
    error, never an unhandled 500 HTML page. Before this fix,
    PyPDF2.PdfReadError propagated straight out of the view function."""
    data = {"file": (_corrupt_pdf_bytes(), "corrupt.pdf")}
    r = client.post("/api/review", data=data, content_type="multipart/form-data")
    assert r.status_code == 422
    body = r.get_json()
    assert body is not None, "response must be JSON, not an HTML error page"
    assert body["success"] is False
    assert "tidak dapat dibaca" in body["error"]


def test_case_analysis_rejects_corrupt_pdf_with_json_error(client):
    data = {"title": "Test", "narrative": "", "file": (_corrupt_pdf_bytes(), "corrupt.pdf")}
    r = client.post("/api/case-analysis", data=data, content_type="multipart/form-data")
    assert r.status_code == 422
    body = r.get_json()
    assert body is not None, "response must be JSON, not an HTML error page"
    assert body["success"] is False
    assert "tidak dapat dibaca" in body["error"]


# --------------------------------------------------------------------------
# v1.3.6 comprehensive practice-expansion endpoint guards
# --------------------------------------------------------------------------

def test_legal_drafting_template_catalog_is_expanded(client):
    r=client.get('/api/drafting/templates')
    assert r.status_code==200
    body=r.get_json()
    assert body['count'] >= 24
    names={x['name'] for x in body['data']}
    for expected in ('Somasi','Jawaban Tergugat','Eksepsi Perdata','Nota Pembelaan / Pledoi','Permohonan Praperadilan','Legal Opinion'):
        assert expected in names


def test_offline_legal_research_uses_sqlite_regulatory_database(client):
    r=post_json(client,'/api/research/search',{'query':'BPR tata kelola kredit','mode':'offline','limit':10})
    assert r.status_code==200
    data=r.get_json()['data']
    assert data['mode']=='offline'
    assert data['counts']['local'] > 0
    assert data['counts']['online'] == 0
    assert all(x['source']=='LOCAL_DATABASE' for x in data['local_results'])


def test_regulatory_search_supports_offline_mode(client):
    r=client.get('/api/regulations/search?q=keuangan%20negara&mode=offline')
    assert r.status_code==200
    data=r.get_json()
    assert data['mode']=='offline'
    assert data['data']['counts']['local'] > 0


def test_case_payload_enforces_domain_contract_before_regulatory_intelligence():
    text=(
        'Jawaban Tergugat atas Gugatan Penggugat dengan eksepsi obscuur libel dan plurium litis consortium. '
        'Sengketa menyangkut Sertipikat Hak Milik, SKPT dan kompetensi absolut Pengadilan Agama.'
    )
    x=appmod._case_analysis_payload(text,'Sengketa Tanah','narrative','',None)
    active=set(x['domain_classification']['domain_contract'])
    assert 'financial_services' not in active and 'criminal' not in active
    regs=' '.join((m.get('regulation') or {}).get('tentang','') for m in x.get('regulatory_matches',[])).lower()
    assert 'bank perkreditan rakyat' not in regs and 'ketenagakerjaan' not in regs


def test_ai_synthesis_prompt_contains_immutable_domain_boundary():
    import ai_engine
    fallback={'case_posture':'PIDANA_KHUSUS_PENYIDIKAN','domain_classification':{'primary_domain':'corruption'},'domain_contract':{'primary_domain':'corruption'}}
    prompt=ai_engine._synthesis_prompt('Test',100,[],fallback,None)
    assert 'KONTRAK YURIDIS IMMUTABLE' in prompt
    assert 'Contoh A' in prompt and 'Contoh B' in prompt
    assert 'Perjanjian kredit' in prompt


def test_drafting_catalog_has_comprehensive_40_templates(client):
    r=client.get('/api/drafting/templates')
    assert r.status_code==200
    body=r.get_json()
    assert body['count'] >= 40
    names={x['name'] for x in body['data']}
    for expected in ('Gugatan PTUN','Jawaban PTUN','Gugatan PHI','Jawaban PHI','Permohonan Eksekusi','Permohonan Penangguhan Penahanan','Gugatan Cerai','Permohonan Cerai Talak','Permohonan Penetapan Ahli Waris'):
        assert expected in names


def test_ai_provider_configuration_is_dynamic(monkeypatch):
    import ai_engine
    monkeypatch.setenv('LEXICORE_AI_PROVIDER','local')
    monkeypatch.delenv('LEXICORE_AI_MODEL',raising=False)
    assert ai_engine.status()['provider']=='local'
    monkeypatch.setenv('LEXICORE_AI_PROVIDER','gemini')
    monkeypatch.setenv('LEXICORE_AI_MODEL','gemini-test')
    monkeypatch.setenv('GEMINI_API_KEY','test-key')
    st=ai_engine.status()
    assert st['provider']=='gemini' and st['model']=='gemini-test' and st['available'] is True

# --------------------------------------------------------------------------
# v1.3.8 scanned-document OCR ingestion
# --------------------------------------------------------------------------

def test_pdf_scan_uses_page_level_ocr_fallback(monkeypatch, tmp_path):
    from PIL import Image, ImageDraw
    import services.document_ocr as ocr

    image = Image.new('RGB', (1000, 1400), 'white')
    ImageDraw.Draw(image).text((80, 100), 'REPLIK PERKARA PERDATA', fill='black')
    pdf = tmp_path / 'scan.pdf'
    image.save(pdf, 'PDF', resolution=150.0)

    monkeypatch.setattr(ocr, '_ocr_pil_image', lambda image, diagnostics: 'REPLIK perkara perdata sertipikat hak milik dan eksepsi obscuur libel')
    text, diag = ocr.extract_pdf_text(str(pdf))
    assert 'REPLIK' in text
    assert diag['mode'] == 'OCR_FALLBACK'
    assert diag['pages_ocr'] == 1
    assert diag['pages_native'] == 0


def test_case_analysis_accepts_scanned_image_upload(client, monkeypatch):
    from contract_review import DocumentExtractor
    extracted = ('REPLIK perkara perdata. Penggugat menolak eksepsi obscuur libel dan plurium litis consortium. '
                 'Sengketa mengenai Sertipikat Hak Milik dan pembagian waris. ' * 4)
    diag = {'enabled': True, 'available': True, 'engine': 'embedded_rapidocr', 'language': 'ind+eng',
            'mode': 'OCR_IMAGE', 'pages_total': 1, 'pages_native': 0, 'pages_ocr': 1,
            'pages_failed': 0, 'characters_native': 0, 'characters_ocr': len(extracted),
            'tesseract_cmd': None, 'warnings': []}
    monkeypatch.setattr(DocumentExtractor, 'extract_with_diagnostics', staticmethod(lambda _path: (extracted, diag)))
    data = {'title': 'Replik scan', 'narrative': '', 'regulatory_mode': 'offline',
            'file': (io.BytesIO(b'fake-image'), 'replik.jpg')}
    r = client.post('/api/case-analysis', data=data, content_type='multipart/form-data')
    assert r.status_code == 200
    x = r.get_json()['data']
    assert x['document_ingestion']['mode'] == 'OCR_IMAGE'
    assert x['domain_classification']['primary_domain'] in {'civil_procedure', 'land_property', 'civil_contract'}


def test_ocr_status_endpoint_is_local_only(client):
    r = client.get('/api/ocr/status')
    assert r.status_code == 200
    data = r.get_json()['data']
    assert data['engine'] == 'embedded_rapidocr'
    assert data['portable'] is True
    assert data['os_executable_required'] is False
    assert data['privacy'] == 'LOCAL_HOST_PROCESSING'
    assert data['policy'] == 'NATIVE_TEXT_FIRST_PAGE_LEVEL_EMBEDDED_OCR_WITH_LOCAL_FALLBACK'


def test_embedded_ocr_result_order_and_confidence_filter():
    import services.document_ocr as ocr
    result = [
        ([[10,50],[110,50],[110,70],[10,70]], 'baris kedua', 0.97),
        ([[10,10],[110,10],[110,30],[10,30]], 'REPLIK', 0.99),
        ([[120,10],[220,10],[220,30],[120,30]], 'PERKARA', 0.98),
        ([[10,90],[110,90],[110,110],[10,110]], 'noise', 0.20),
    ]
    text, avg = ocr._rapidocr_result_to_text(result, 0.45)
    assert text.splitlines()[0] == 'REPLIK PERKARA'
    assert 'baris kedua' in text
    assert 'noise' not in text
    assert avg and avg > 0.9


def test_embedded_ocr_runtime_contract(monkeypatch):
    import services.document_ocr as ocr
    monkeypatch.setattr(ocr, 'embedded_ocr_available', lambda: True)
    st = ocr.ocr_runtime_status(load_engine=False)
    assert st['engine'] == 'embedded_rapidocr'
    assert st['portable'] is True
    assert st['os_executable_required'] is False
    assert st['privacy'] == 'LOCAL_HOST_PROCESSING'

# --------------------------------------------------------------------------
# v1.3.10 OCR evidence & jurisdiction stabilization
# --------------------------------------------------------------------------

def test_rapidocr_v3_numpy_boxes_do_not_trigger_fallback(monkeypatch):
    import types
    import numpy as np
    import services.document_ocr as ocr
    from PIL import Image

    class FakeEngine:
        def __call__(self, _arr):
            return types.SimpleNamespace(
                boxes=np.array([[[10,10],[110,10],[110,30],[10,30]]], dtype=float),
                txts=("REPLIK PERKARA PERDATA",),
                scores=(0.99,),
            )
    monkeypatch.setattr(ocr, '_get_rapidocr_engine', lambda: FakeEngine())
    d=ocr.OCRDiagnostics()
    text=ocr._ocr_embedded(Image.new('RGB',(300,100),'white'),d)
    assert 'REPLIK' in text
    assert d.engine=='embedded_rapidocr'
    assert d.portable is True
    assert not any('ambiguous' in str(x).lower() for x in d.warnings)




def test_embedded_ocr_timeout_returns_control_and_allows_fallback(monkeypatch):
    import time
    from PIL import Image
    import services.document_ocr as ocr

    # Keep this regression test fast while exercising the real timeout wrapper.
    monkeypatch.setenv('LEXICORE_EMBEDDED_OCR_TIMEOUT_SECONDS', '0.25')
    monkeypatch.setattr(ocr, '_RAPID_CIRCUIT_OPEN', False)
    monkeypatch.setattr(ocr, '_RAPID_CALL_INFLIGHT', False)
    monkeypatch.setattr(ocr, '_RAPID_TIMEOUT_COUNT', 0)

    def stuck_core(_image):
        time.sleep(0.8)
        return ('late result', 0.99)

    monkeypatch.setattr(ocr, '_ocr_embedded_core', stuck_core)
    monkeypatch.setattr(ocr, '_ocr_tesseract_optional', lambda _image, diagnostics: 'TESSERACT FALLBACK')
    d = ocr.OCRDiagnostics()
    started = time.monotonic()
    text = ocr._ocr_pil_image(Image.new('RGB', (100, 50), 'white'), d)
    elapsed = time.monotonic() - started

    assert text == 'TESSERACT FALLBACK'
    assert elapsed < 0.7
    assert ocr._RAPID_CIRCUIT_OPEN is True
    assert ocr._RAPID_TIMEOUT_COUNT == 1
    assert any('timeout' in str(w).lower() for w in d.warnings)

    # Do not leak the circuit-breaker state to following tests.
    ocr._RAPID_CIRCUIT_OPEN = False
    ocr._RAPID_CALL_INFLIGHT = False
    ocr._RAPID_TIMEOUT_COUNT = 0


def test_flask_dev_server_is_explicitly_threaded():
    from pathlib import Path
    source = (Path(__file__).resolve().parents[1] / 'app.py').read_text(encoding='utf-8')
    assert 'app.run(debug=debug_mode, host=host, port=port, threaded=True)' in source


def test_pdf_scan_total_ocr_budget_stops_remaining_pages(monkeypatch, tmp_path):
    import time
    from PIL import Image
    import services.document_ocr as ocr

    pages = [Image.new('RGB', (400, 500), 'white') for _ in range(3)]
    pdf = tmp_path / 'multi-scan.pdf'
    pages[0].save(pdf, 'PDF', save_all=True, append_images=pages[1:], resolution=120.0)

    monkeypatch.setattr(ocr, '_ocr_total_timeout_seconds', lambda: 0.15)

    def slow_page(_image, _diagnostics):
        time.sleep(0.20)
        return 'OCR PAGE'

    monkeypatch.setattr(ocr, '_ocr_pil_image', slow_page)
    started = time.monotonic()
    text, diag = ocr.extract_pdf_text(str(pdf))
    elapsed = time.monotonic() - started

    assert 'OCR PAGE' in text
    assert diag['total_timeout_exceeded'] is True
    assert diag['pages_ocr'] == 1
    assert diag['pages_failed'] >= 2
    assert elapsed < 1.0
    assert any('anggaran waktu total ocr' in str(w).lower() for w in diag['warnings'])


def test_ocr_runtime_status_accepts_local_fallback_when_embedded_unavailable(monkeypatch):
    import services.document_ocr as ocr
    monkeypatch.setattr(ocr,'embedded_ocr_available',lambda: False)
    monkeypatch.setattr(ocr,'find_tesseract',lambda: r'C:\\Program Files\\Tesseract-OCR\\tesseract.exe')
    # Avoid actually touching the host binary in this unit contract.
    class FakePT:
        class pytesseract:
            tesseract_cmd=''
        @staticmethod
        def get_languages(config=''):
            return ['eng','ind','osd']
    import sys
    monkeypatch.setitem(sys.modules,'pytesseract',FakePT)
    st=ocr.ocr_runtime_status(load_engine=True)
    assert st['available'] is True
    assert st['engine']=='tesseract_fallback'
    assert st['embedded_available'] is False
    assert st['tesseract_available'] is True


def test_replik_inheritance_forum_screen_is_secondary_not_assumed_primary():
    from services.case_domain_classifier import classify_case
    text=(
        'REPLIK DALAM EKSEPSI. Exceptio declinatoir mengenai pembagian waris dan para ahli waris. '
        'Diperdebatkan kewenangan absolut Pengadilan Agama. Objek sengketa Sertipikat Hak Milik dan SKPT. '
        'Terdapat eksepsi obscuur libel dan plurium litis consortium.'
    )
    c=classify_case(text)
    active=set(c['domain_contract'])
    assert c['posture']=='PERDATA_LITIGASI'
    assert 'religious_court' in active
    assert c['forum_screen']['inheritance_detected'] is True
    assert c['forum_screen']['absolute_competence_disputed'] is True
    assert c['forum_screen']['note']


def test_deterministic_source_ledger_extracts_material_ocr_facts():
    import app
    source='''
    REPLIK PERKARA NOMOR 44/Pdt.G/2026/PN.Kpn.
    DALAM EKSEPSI
    Tentang Exceptio Declinatoir. Pembagian waris telah dilakukan kepada para ahli waris.
    Tentang Eksepsi Plurium Litis Consortium.
    Tentang Eksepsi Obscuur Libel.
    DALAM KONVENSI
    Sertipikat Hak Milik Nomor 05392 seluas 73M2 dipersoalkan karena data tanah dan rumah dinilai tidak benar.
    Para pihak menyebut proses PTSL dan Surat Keterangan Pendaftaran Tanah (SKPT).
    DALAM REKONVENSI
    Para Penggugat meminta gugatan rekonvensi ditolak.
    KESIMPULAN
    Penggugat menawarkan penyelesaian pembagian objek 50% dan 50% apabila syarat hukum terpenuhi.
    PERMOHONAN
    Menolak Eksepsi Para Tergugat dan mengabulkan gugatan Penggugat.
    '''
    rows=app._source_ledger_candidates(source)
    text=' '.join(x['statement'] for x in rows).lower()
    assert len(rows) >= 8
    for term in ('plurium','obscuur','05392','73m2','skpt','rekonvensi','permohonan'):
        assert term in text
    assert all(x.get('source')=='DETERMINISTIC_EXTRACTIVE' for x in rows)


def test_executive_summary_suppresses_identity_identifiers():
    import app
    result={
        'source_text':'Replik gugatan perdata. Eksepsi obscuur libel. Sertipikat Hak Milik Nomor 05392.',
        'case_posture':'PERDATA_LITIGASI',
        'facts':['Ninik Sugiyanti, lahir di Malang, NIK: 3507226507700001, Phone 081234567890, email: user@example.com.',
                 'Sertipikat Hak Milik Nomor 05392 menjadi objek sengketa.'],
        'legal_issues':['Apakah gugatan memenuhi syarat formil?'],
        'legal_analysis':'Perkara harus diuji berdasarkan kompetensi forum dan bukti pertanahan.',
        'evidentiary_gaps':['Warkah dan SKPT perlu diverifikasi.'],
    }
    text=' '.join(app._compact_executive_summary(result))
    assert '3507226507700001' not in text
    assert '081234567890' not in text
    assert 'user@example.com' not in text
    assert '05392' in text


def test_local_mode_reports_complete_document_reading(client):
    narrative=('Pekerja mengajukan gugatan hubungan industrial terkait PHK dan pesangon. '
               'Dokumen pendukung hubungan kerja, surat PHK dan perhitungan hak perlu diuji. ' * 3)
    r=post_json(client,'/api/case-analysis',{'title':'Local complete','narrative':narrative,'verify_online':False})
    assert r.status_code==200
    x=r.get_json()['data']
    assert x['analysis_provenance']['mode']=='LOCAL_DETERMINISTIC'
    assert x['document_reading']['status']=='COMPLETE'
    assert x['document_reading']['segments_read']==x['document_reading']['segments_total']==1

# --------------------------------------------------------------------------
# v1.3.11 legal OCR post-processing + evidence grouping
# --------------------------------------------------------------------------

def test_legal_ocr_postprocess_is_conservative_and_auditable():
    from services.legal_ocr_postprocess import postprocess_legal_ocr
    raw=(
        'Tentang Exepsio Decklinatoir. Para ahliwaris kamijawab dalam REKONVENSI. '
        'Sertipikat Hak Milik Nomor 05392 seluas 73M2. '
        'Ninik Sugiyanti NIK: 3507226507700001. Riuwayat mendapatkan tanag tanpa atass dasar lain.'
    )
    text, diag=postprocess_legal_ocr(raw)
    low=text.lower()
    assert 'exceptio declinatoir' in low
    assert 'ahli waris' in low
    assert 'kami jawab' in low
    assert 'nomor 05392' in low
    assert '73 m2' in low
    assert 'tanah' in low
    assert 'atas' in low
    assert '3507226507700001' in text
    assert 'Riuwayat' in text
    assert diag['policy']=='CONSERVATIVE_LEGAL_OCR_NORMALIZATION'
    assert diag['corrections_count'] >= 5
    assert diag['raw_sha256'] != diag['normalized_sha256']


def test_evidence_grouping_is_extractive_and_traceable():
    from services.legal_ocr_postprocess import group_source_ledger
    ledger=[
        {'label':'DOCUMENT','statement':'REPLIK DALAM EKSEPSI','evidence':'REPLIK DALAM EKSEPSI'},
        {'label':'SOURCE FACT','statement':'Exceptio Declinatoir mempersoalkan kewenangan Pengadilan Agama.','evidence':'Exceptio Declinatoir mempersoalkan kewenangan Pengadilan Agama.'},
        {'label':'SOURCE FACT','statement':'Pembagian waris telah dilakukan kepada para ahli waris.','evidence':'Pembagian waris telah dilakukan kepada para ahli waris.'},
        {'label':'SOURCE FACT','statement':'Sertipikat Hak Milik Nomor 05392 dan proses PTSL menjadi objek sengketa.','evidence':'Sertipikat Hak Milik Nomor 05392 dan proses PTSL menjadi objek sengketa.'},
        {'label':'SOURCE FACT','statement':'Eksepsi Plurium Litis Consortium dan Obscuur Libel dimohonkan ditolak.','evidence':'Eksepsi Plurium Litis Consortium dan Obscuur Libel dimohonkan ditolak.'},
        {'label':'DOCUMENT','statement':'PERMOHONAN','evidence':'PERMOHONAN'},
        {'label':'ALLEGATION','statement':'Mengabulkan gugatan Penggugat dan menghukum Para Tergugat.','evidence':'Mengabulkan gugatan Penggugat dan menghukum Para Tergugat.'},
    ]
    groups=group_source_ledger(ledger)
    ids={g['id'] for g in groups}
    assert {'jurisdiction','inheritance','land_registration','formal_defenses','relief'} <= ids
    for g in groups:
        assert g['traceability']=='EXTRACTIVE_GROUPING_ONLY'
        assert all(isinstance(i,int) and 0 <= i < len(ledger) for i in g['source_item_indexes'])
        joined=' '.join(x['statement'] for x in g['items'])
        for fact in g['material_facts']:
            assert all(token in joined for token in fact.split()[:3])


def test_version_metadata_runtime_contract():
    import re
    import version
    assert version.PRODUCT_NAME == "LexiCore"
    assert version.PRODUCT_LABEL == "LexiCore Assistant"
    assert version.INITIATIVE == "Evidence-to-Action Legal Intelligence"
    assert version.FIRM_NAME == "ELF - Erfan's Law Firm"
    assert re.fullmatch(r"\d+\.\d+\.\d+-rc\d+", version.LEXICORE_VERSION)
    assert version.LEXICORE_VERSION.startswith(version.PUBLIC_VERSION + "-rc")
    assert version.release_metadata()["product_label"] == version.PRODUCT_LABEL
    assert version.release_metadata()["version"] == version.LEXICORE_VERSION


# --------------------------------------------------------------------------
# v1.3.13 — Lease drafting + deterministic contract review V2
# --------------------------------------------------------------------------

def test_lease_draft_is_comprehensive_and_keeps_placeholders():
    from legal_drafting import build_legal_draft
    content, clause_count = build_legal_draft(
        'Perjanjian Sewa Menyewa', 'Pemilik A', 'Penyewa B', '2026-09-06', 12,
        'Rumah tinggal di [ALAMAT LENGKAP], harga sewa [NILAI].'
    )
    assert clause_count >= 20
    for heading in (
        'OBJEK SEWA & STATUS PENGUASAAN', 'SERAH TERIMA, KONDISI AWAL & INVENTARIS',
        'UTILITAS, IURAN, PAJAK & BIAYA LAIN', 'WANPRESTASI / DEFAULT, PEMBERITAHUAN & CURE PERIOD',
        'PENGAKHIRAN SEBELUM JATUH TEMPO', 'PENYELESAIAN PERSELISIHAN'
    ):
        assert heading in content
    assert '[TANGGAL BERAKHIR]' in content


def test_contract_review_v2_classifies_lease_and_finds_structural_gaps():
    from contract_review import ContractTypeClassifier, EntityExtractor, ClauseAnalyzer, RiskDetector
    text = '''
SURAT PERJANJIAN SEWA-MENYEWA RUMAH
Pasal 1
Pihak Pertama menyewakan rumah kepada Pihak Kedua.
Pasal 2
Pihak Kedua menggunakan rumah untuk Tempat Tinggal.
Pasal 3
Pihak Kedua menyewa selama 1 tahun terhitung sejak tanggal 6 September 2026 sampai dengan tanggal 6 September 2027.
Harga sewa Rp. 8.000.000 dan dibayar sebagai tanda pelunasan seluruh jumlah uang sewa.
Pasal 4
Rumah bebas dari sengketa.
Pasal 5
Pihak Kedua berkewajiban membayar listrik dan air PDAM.
Pasal 6
Pihak Kedua berkewajiban merawat rumah.
Pasal 7
Pihak Kedua tidak dibenarkan mengalihkan hak sewa atau mengubah struktur tanpa izin.
Pasal 8
Force majeure meliputi bencana alam.
Pasal 9
Pihak Kedua dapat memutus sewa sebelum jangka waktu berakhir dengan pemberitahuan tertulis sekurang-kurangnya 1 bulan sebelum berakhirnya jangka waktu.
Pasal 10
Pihak Pertama dapat memutus sewa apabila Pihak Kedua lalai membayar harga sewa setelah jatuh tempo.
Pasal 11
Pihak Kedua wajib mengosongkan dan menyerahkan kembali rumah.
Pasal 12
Perpanjangan berdasarkan kesepakatan.
Pasal 14
Perjanjian dibuat rangkap dua.
'''
    ctype, confidence = ContractTypeClassifier.detect(text)
    assert ctype == 'SEWA_MENYEWA' and confidence >= 80
    dates = EntityExtractor.extract_dates(text)
    assert dates['effective_date'] == '6 September 2026'
    assert dates['termination_date'] == '6 September 2027'
    coverage, missing = ClauseAnalyzer.coverage(text, ctype)
    assert any(x['key'] == 'handover' for x in missing)
    assert any(x['key'] == 'dispute' for x in missing)
    inconsistencies = ClauseAnalyzer.inconsistencies(text, ctype)
    assert any(x['category'] == 'Penomoran Pasal' for x in inconsistencies)
    assert any(x['category'] == 'Konsistensi Pembayaran' for x in inconsistencies)
    risks = RiskDetector.analyze(text, ctype, missing, inconsistencies)
    assert any(r.category == 'Kelengkapan Kontrak' for r in risks)


def test_contract_risk_keyword_matching_does_not_match_nda_inside_indonesian_words():
    from contract_review import RiskDetector
    risks = RiskDetector.analyze('Perjanjian ini ditandatangani berdasarkan kesepakatan para pihak.')
    assert not any(r.category == 'Kerahasiaan' for r in risks)


# --------------------------------------------------------------------------
# v1.3.13.1 — Deep clause evaluation table
# --------------------------------------------------------------------------

def test_contract_review_deep_clause_table_preserves_existing_and_gives_copy_ready_redraft():
    from contract_review import ClauseDeepEvaluator, ClauseAnalyzer
    text = """
Pasal 3
PIHAK KEDUA akan menyewa rumah selama 1 tahun terhitung sejak tanggal 6 September 2026 sampai dengan tanggal 6 September 2027.
Harga sewa Rp. 8.000.000 dibayar sebagai pelunasan seluruh jumlah uang sewa.
Pasal 9
PIHAK KEDUA dapat memutuskan hubungan sewa-menyewa sebelum jangka waktu berakhir dengan pemberitahuan tertulis sekurang-kurangnya 1 bulan sebelum berakhirnya jangka waktu.
Pasal 10
PIHAK PERTAMA dapat memutuskan hubungan sewa-menyewa apabila PIHAK KEDUA lalai membayar harga sewa selama 1 bulan setelah jatuh tempo.
"""
    coverage, missing = ClauseAnalyzer.coverage(text, 'SEWA_MENYEWA')
    inconsistencies = ClauseAnalyzer.inconsistencies(text, 'SEWA_MENYEWA')
    rows = ClauseDeepEvaluator.evaluate(text, 'SEWA_MENYEWA', missing, inconsistencies)
    p3 = next(r for r in rows if r['article_number'] == '3')
    assert p3['existing_clause'].startswith('Pasal 3')
    assert 'Rp. 8.000.000' in p3['existing_clause']
    assert 'inkonsistensi' in p3['risk_loophole'].lower() or 'pembayaran' in p3['risk_loophole'].lower()
    assert 'Jangka waktu sewa berlaku' in p3['recommended_redraft']
    assert all(set(('existing_clause','risk_loophole','recommended_redraft')).issubset(r) for r in rows)

def test_contract_review_deep_table_adds_missing_dispute_as_explicit_proposed_clause():
    from contract_review import ClauseDeepEvaluator, ClauseAnalyzer
    text = 'Pasal 1\nPIHAK PERTAMA menyewakan rumah kepada PIHAK KEDUA.\nPasal 2\nPIHAK KEDUA menggunakan rumah untuk tempat tinggal.'
    coverage, missing = ClauseAnalyzer.coverage(text, 'SEWA_MENYEWA')
    rows = ClauseDeepEvaluator.evaluate(text, 'SEWA_MENYEWA', missing, [])
    dispute = [r for r in rows if 'penyelesaian perselisihan' in r['risk_loophole'].lower()]
    assert dispute
    assert dispute[0]['existing_clause'].startswith('— Tidak ada pasal eksisting')
    assert 'PASAL [●] — PENYELESAIAN PERSELISIHAN' in dispute[0]['recommended_redraft']

# v1.3.13.2 Regulatory Corpus hierarchy / provision / implication regression

def test_regulatory_corpus_hierarchy_contract():
    from services.legal_research_service import search_legal_authorities
    r=search_legal_authorities('peradilan agama Pasal 49',mode='offline',limit=12)
    groups=r.get('hierarchical_results') or []
    assert groups
    labels=[g.get('label') for g in groups]
    assert any('Undang-Undang' in (x or '') for x in labels)
    rows=[x for g in groups for x in g.get('items',[])]
    assert all('practical_implication' in x for x in rows)
    assert all('specific_provisions' in x for x in rows)
    assert any(any((p.get('pasal') or '').startswith('Pasal 49') for p in x.get('specific_provisions',[])) for x in rows)


def test_regulatory_hierarchy_order_is_workspace_order():
    from services.legal_research_service import REGULATORY_HIERARCHY
    assert [x[0] for x in REGULATORY_HIERARCHY[:8]] == ['UU','PP','PERPRES','PER_MENTERI_LEMBAGA','PERMA','PERMK','PERKAP','PERDA']


# --------------------------------------------------------------------------
# v1.3.13.4 Compliance Risk Matrix
# --------------------------------------------------------------------------

def test_compliance_categories_have_specific_question_sets(client):
    for category in ('General Corporate','Tech/PDP','Employment','Commercial','Financial/AML'):
        r=client.get('/api/compliance/questions',query_string={'category':category})
        assert r.status_code==200
        data=r.get_json()['data']
        assert data['category']==category
        assert len(data['questions'])>=3
        assert all(q.get('key') and q.get('question') for q in data['questions'])

def test_compliance_assessment_returns_required_matrix_columns(client):
    r=post_json(client,'/api/compliance/assess',{'entity':'PT Contoh','category':'Tech/PDP','answers':{'privacy':'no','security':'partial','processor':'yes'}})
    assert r.status_code==200
    x=r.get_json()['data']
    assert x['risk_level'] in {'HIGH','MEDIUM','LOW'}
    assert x['professional_verification']=='PENDING'
    assert x['risk_matrix']
    for row in x['risk_matrix']:
        assert row['risk_identification']
        assert row['risk_level'] in {'HIGH','MEDIUM','LOW'}
        assert row['legal_justification']
        assert row['sanction_basis']
        assert isinstance(row['mitigation_checklist'],list) and row['mitigation_checklist']

def test_compliance_high_risk_is_justified_by_sanction_exposure(client):
    r=post_json(client,'/api/compliance/assess',{'category':'Financial/AML','answers':{'aml':'no','reporting':'no','financial_controls':'partial'}})
    x=r.get_json()['data']
    assert x['risk_level']=='HIGH'
    aml=next(row for row in x['risk_matrix'] if row['control_key']=='aml')
    assert aml['risk_level']=='HIGH'
    assert 'pidana' in aml['legal_justification'].lower() or 'pencabutan' in aml['legal_justification'].lower()

def test_compliance_history_preserves_risk_matrix(client):
    post_json(client,'/api/compliance/assess',{'entity':'PT Persist','category':'Commercial','answers':{'contracts':'partial','consumer':'yes','third_party':'no'}})
    r=client.get('/api/compliance?limit=5')
    assert r.status_code==200
    rows=r.get_json()['data']
    hit=next(x for x in rows if x.get('entity')=='PT Persist')
    assert hit.get('risk_matrix')
    assert hit.get('assessment_method')=='CATEGORY_SPECIFIC_DETERMINISTIC_MATRIX'


# v1.3.13.5 — four-part Case Analysis working paper
def test_case_working_paper_has_four_sequential_components():
    from services.case_working_paper import build_case_working_paper
    x={
      'case_posture':'PERDATA_LITIGASI',
      'domain_classification':{'domain_contract':['civil_procedure','civil_contract']},
      'case_readiness':{'components':{'evidence_map':{'percentage':70,'required':4,'fulfilled':3},'legal_analysis':{'percentage':75}}},
      'source_ledger':[{'label':'DOCUMENT','statement':'Perjanjian sewa tanggal 1 Januari 2026','evidence':'Perjanjian tertulis'}],
      'facts':['Para pihak menandatangani perjanjian tertulis.'],
      'legal_issues':['Apakah kewajiban pembayaran telah dilanggar?'],
      'arguments_for':['Ada dokumen perjanjian.'],
      'arguments_against':['Bukti pembayaran belum lengkap.'],
      'evidentiary_gaps':['Bukti pembayaran asli belum tersedia.'],
      'applicable_law':[{'source':'KUHPerdata — verifikasi pasal spesifik'}],
      'action_plan':[{'priority':'P1','issue':'Bukti pembayaran','current_status':'BELUM TERVERIFIKASI','action':'Dapatkan bukti pembayaran asli.','why_it_matters':'Menentukan ada/tidaknya wanprestasi.'}],
      'case_regulatory_snapshot':{'official_results':[]}
    }
    wp=build_case_working_paper(x)
    assert list(k for k in ('working_paper_percentage','evidence_map','legal_construction','action_plan') if k in wp)==['working_paper_percentage','evidence_map','legal_construction','action_plan']
    assert 15 <= wp['working_paper_percentage']['percentage'] <= 85
    assert wp['evidence_map']['rows']
    assert wp['legal_construction']['chains']
    assert wp['action_plan'][0]['step']==1

def test_criminal_evidence_basis_does_not_hardcode_old_kuhap_article():
    from services.case_working_paper import build_case_working_paper
    x={'case_posture':'PIDANA_PENYIDIKAN','domain_classification':{'domain_contract':['criminal']},'case_readiness':{'components':{}},'source_ledger':[],'case_regulatory_snapshot':{}}
    wp=build_case_working_paper(x)
    basis=wp['evidence_map']['legal_basis']
    assert basis['system']=='CRIMINAL'
    assert 'tanggal proses' in basis['citation'].lower()
    assert 'tidak di-hard-code' in basis['note'].lower()


# v1.3.13.6 — Menu 7 strict norm-conflict resolver
def test_norm_conflict_rule_a_vs_multiple_and_lex_superior():
    from norm_conflict import analyze_conflicts
    x=analyze_conflicts([
        'UU Nomor 99 Tahun 2020 tentang Perlindungan Contoh',
        'Peraturan Menteri Nomor 5 Tahun 2024 tentang Perlindungan Contoh',
        'Perda Nomor 2 Tahun 2025 tentang Perlindungan Contoh',
    ], 'Aturan-aturan tersebut bertentangan dan tidak dapat diterapkan bersama pada kewajiban yang sama.')
    rows=x['rule_comparison_matrix']
    assert len(rows)==2
    assert rows[0]['pair_label']=='Aturan A vs Aturan B'
    assert rows[1]['pair_label']=='Aturan A vs Aturan C'
    assert rows[0]['principle_applied']=='LEX_SUPERIOR'
    assert rows[0]['applicable_law'].startswith('UU Nomor 99')

def test_norm_conflict_lex_posterior_same_rank():
    from norm_conflict import analyze_conflicts
    x=analyze_conflicts([
        'UU Nomor 1 Tahun 2020 tentang Perlindungan Data',
        'UU Nomor 2 Tahun 2022 tentang Perlindungan Data',
    ], 'Kedua aturan bertentangan dan tidak dapat diterapkan bersama pada perlindungan data.')
    row=x['rule_comparison_matrix'][0]
    assert row['principle_applied']=='LEX_POSTERIOR'
    assert '2022' in row['applicable_law']

def test_norm_conflict_no_antinomy_no_fake_applicable_law():
    from norm_conflict import analyze_conflicts
    x=analyze_conflicts(['UU Nomor 1 Tahun 2020 tentang Pajak','PP Nomor 2 Tahun 2021 tentang Pajak'], 'Keduanya mengatur administrasi pajak.')
    row=x['rule_comparison_matrix'][0]
    assert row['classification']=='RELATIONSHIP_ONLY'
    assert row['applicable_law']=='NOT_DETERMINED'


# --------------------------------------------------------------------------
# v1.3.13.7 — Menu 8 professional client communication
# --------------------------------------------------------------------------

def test_client_update_professional_structure(client):
    r = post_json(client, "/api/communication/generate", {
        "client_name": "Budi", "document_type": "client_update",
        "matter": "sengketa perjanjian",
        "progress": "Somasi pertama telah dikirim dan bukti penerimaan telah tersimpan",
        "next_step": "mengevaluasi respons pihak lawan dan menyiapkan opsi tindakan berikutnya",
    })
    assert r.status_code == 200
    data = r.get_json()["data"]
    msg = data["message"]
    assert "Status saat ini" in msg
    assert "Makna bagi posisi hukum Anda" in msg
    assert "Langkah berikutnya" in msg
    assert "ELF - Erfan's Law Firm" in msg
    assert data["professional_status"] == "DRAFT_FOR_LAWYER_REVIEW"


def test_document_request_has_order_deadline_and_legal_reason(client):
    r = post_json(client, "/api/communication/generate", {
        "client_name": "Ninik", "document_type": "document_request",
        "matter": "sengketa tanah dan sertipikat",
        "deadline": "10 September 2026 pukul 16.00 WIB",
        "next_step": "melakukan pemetaan bukti dan verifikasi dasar hukum",
    })
    assert r.status_code == 200
    data = r.get_json()["data"]
    assert len(data["document_items"]) >= 4
    assert data["document_items"][0]["priority"] == 1
    assert all(x["deadline"] == "10 September 2026 pukul 16.00 WIB" for x in data["document_items"])
    assert all(x["legal_reason"] for x in data["document_items"])
    assert "Alasan hukum/strategis" in data["message"]
    assert "Tenggat:" in data["message"]


# v1.3.13.7.1 — WhatsApp delivery launcher + norm conflict relationship corrective
def test_client_communication_whatsapp_link_uses_registered_contact(client):
    save = client.post('/api/communications', json={
        'client_id':'CL-0001','client_name':'Klien Uji','whatsapp_number':'081234567890',
        'subject':'Pembaruan Perkara','message':'Perkara sedang ditelaah.','document_type':'client_update','status':'draft'
    })
    assert save.status_code == 200
    r = client.post('/api/communication/whatsapp-link', json={
        'client_id':'CL-0001','subject':'Pembaruan Perkara','message':'Perkara sedang ditelaah.'
    })
    assert r.status_code == 200
    data=r.get_json()
    assert data['whatsapp_number'] == '6281234567890'
    assert data['whatsapp_url'].startswith('https://wa.me/6281234567890?text=')
    assert data['delivery_status'] == 'WHATSAPP_COMPOSER_READY'


def test_client_communication_schema_persists_client_id_and_whatsapp():
    cols={name for name,_ in EXPECTED_COLUMNS['client_communications']}
    assert {'client_id','whatsapp_number'} <= cols
    assert SCHEMA_VERSION >= 10


def test_norm_conflict_does_not_misclassify_kuhperdata_as_perda():
    from norm_conflict import _manual_norm
    row=_manual_norm('KUHPerdata','Aturan A')
    assert row['type'] == 'UNKNOWN'
    assert row['rank'] is None


# --------------------------------------------------------------------------
# v1.3.13.8 — Legal Drafting structural standard V2
# --------------------------------------------------------------------------

def test_contract_templates_have_mandatory_core_architecture():
    from legal_drafting import build_legal_draft, CONTRACT_TYPES
    required = (
        'DEFINISI', 'HAK & KEWAJIBAN',
        'PERNYATAAN & JAMINAN (REPRESENTATIONS & WARRANTIES)',
        'WANPRESTASI / DEFAULT', 'FORCE MAJEURE',
        'SENGKETA', 'DOMISILI HUKUM',
    )
    for doc_type in sorted(CONTRACT_TYPES):
        content, _ = build_legal_draft(doc_type, 'Pihak A', 'Pihak B', '2026-09-04', 12, 'Objek transaksi')
        for token in required:
            assert token in content, (doc_type, token)


def test_civil_and_tun_litigation_have_identity_posita_and_primary_subsidiary_relief():
    from legal_drafting import build_legal_draft
    types = ('Gugatan Perdata','Jawaban Tergugat','Eksepsi Perdata','Replik','Duplik',
             'Gugatan PTUN','Jawaban PTUN','Eksepsi PTUN','Replik PTUN','Duplik PTUN')
    for doc_type in types:
        content, _ = build_legal_draft(doc_type, 'Pihak A', 'Pihak B', '2026-09-04', 12, 'Kronologi perkara')
        assert 'IDENTITAS PARA PIHAK' in content
        assert 'POSITA / FUNDAMENTUM PETENDI' in content
        assert 'PETITUM PRIMAIR' in content
        assert 'PETITUM SUBSIDAIR' in content


def test_criminal_drafting_focuses_dakwaan_and_element_proof():
    from legal_drafting import build_legal_draft
    eksepsi, _ = build_legal_draft('Eksepsi Pidana','Terdakwa','Penuntut Umum','2026-09-04',12,'Dakwaan')
    assert 'PASAL 143 KUHAP' in eksepsi
    assert 'SYARAT FORMIL DAN MATERIIL DAKWAAN' in eksepsi
    assert 'verifikasi terhadap KUHAP yang berlaku menurut tempus perkara' in eksepsi
    pledoi, _ = build_legal_draft('Nota Pembelaan / Pledoi','Terdakwa','Penuntut Umum','2026-09-04',12,'Pembelaan')
    assert 'MATRIS PEMBUKTIAN UNSUR PASAL' in pledoi
    assert 'ANALISIS BEBAN & KEKUATAN PEMBUKTIAN' in pledoi


def test_advisory_uses_mandatory_five_part_structure():
    from legal_drafting import build_legal_draft
    for doc_type in ('Legal Opinion','Legal Memorandum'):
        content, count = build_legal_draft(doc_type,'Klien','Objek','2026-09-04',12,'Fakta utama')
        assert count == 5
        positions = [content.index(x) for x in (
            'KASUS POSISI (FAKTA)', 'ISU HUKUM (LEGAL ISSUES)',
            'DASAR HUKUM / REGULASI (LEGAL FRAMEWORK)',
            'ANALISIS HUKUM PER ISU (LEGAL ANALYSIS)',
            'KESIMPULAN & REKOMENDASI MITIGASI')]
        assert positions == sorted(positions)


# --------------------------------------------------------------------------
# v1.3.13.10 Compliance questionnaire expansion
# --------------------------------------------------------------------------

def test_compliance_expanded_question_sets_have_groups_and_custom_options():
    from services.compliance_risk import question_set
    minimums={
        'General Corporate':9, 'Tech/PDP':9, 'Employment':9,
        'Commercial':9, 'Financial/AML':10,
    }
    for category,minimum in minimums.items():
        qs=question_set(category)
        assert len(qs)>=minimum, (category,len(qs))
        assert all(q.get('group') for q in qs)
        assert all(isinstance(q.get('options'),list) and len(q['options'])>=3 for q in qs)

def test_compliance_source_framework_controls_are_present():
    from services.compliance_risk import question_set
    checks={
        'General Corporate': {'financial_tax','related_party','asset_ip','confidentiality','litigation_exposure'},
        'Tech/PDP': {'dpo','dpia_transfer','data_subject_rights','consent_notice','incident_drill','privacy_training'},
        'Employment': {'pkwt_registration','tka','discipline','termination_reserve','outsourcing'},
        'Commercial': {'liability_indemnity','warranty_sla','standard_clause','competition','credit_enforcement','dispute_forum'},
        'Financial/AML': {'sanctions_lists','sanctions_match','pep_edd','aml_audit','aml_training'},
    }
    for category,keys in checks.items():
        found={q['key'] for q in question_set(category)}
        assert keys <= found, (category, keys-found)

def test_compliance_additional_material_controls_are_present():
    from services.compliance_risk import question_set
    assert 'beneficial_owner' in {q['key'] for q in question_set('General Corporate')}
    assert 'k3' in {q['key'] for q in question_set('Employment')}
    assert 'transaction_monitoring' in {q['key'] for q in question_set('Financial/AML')}
    assert 'aml_scope' in {q['key'] for q in question_set('Financial/AML')}

def test_compliance_not_applicable_answer_does_not_inflate_score():
    from services.compliance_risk import build_risk_matrix
    out=build_risk_matrix('Employment', {'tka':'na','outsourcing':'na'})
    tka=next(r for r in out['matrix'] if r['control_key']=='tka')
    assert tka['control_status']=='NA'
    assert tka['risk_level']=='LOW'
    assert 'NOT APPLICABLE' in tka['legal_justification']
    assert out['applicable_question_count']==out['question_count']-2


# --------------------------------------------------------------------------
# v1.3.13.10.1 Progressive user-defined answer options
# --------------------------------------------------------------------------

def test_custom_compliance_question_uses_user_defined_answer_risk_mapping():
    from services.compliance_risk import build_risk_matrix
    custom=[{
        'key':'privacy_importance',
        'group':'Kontrol Tambahan',
        'question':'Seberapa penting pelindungan data pribadi menurut Anda?',
        'severity':'HIGH',
        'risk':'Perlindungan data belum diprioritaskan secara memadai.',
        'options':[
            {'value':'opt_1','label':'Sangat penting dan sudah diprioritaskan','risk_level':'LOW'},
            {'value':'opt_2','label':'Penting tetapi belum konsisten','risk_level':'MEDIUM'},
            {'value':'opt_3','label':'Belum menjadi prioritas','risk_level':'HIGH'},
        ]
    }]
    out=build_risk_matrix('Tech/PDP', {'privacy_importance':'opt_3'}, custom)
    row=next(r for r in out['matrix'] if r['control_key']=='privacy_importance')
    assert row['source']=='USER_DEFINED'
    assert row['answer_label']=='Belum menjadi prioritas'
    assert row['risk_level']=='HIGH'
    assert row['control_status']=='HIGH'
    assert out['custom_question_count']==1


def test_custom_compliance_na_option_excluded_from_score():
    from services.compliance_risk import build_risk_matrix
    custom=[{
        'key':'optional_control','question':'Apakah kontrol ini relevan?',
        'options':[
            {'value':'yes','label':'Relevan','risk_level':'LOW'},
            {'value':'na','label':'Tidak relevan','risk_level':'NA'},
        ]
    }]
    out=build_risk_matrix('General Corporate', {'optional_control':'na'}, custom)
    row=next(r for r in out['matrix'] if r['control_key']=='optional_control')
    assert row['control_status']=='NA'
    assert row['risk_level']=='LOW'
    assert out['applicable_question_count']==out['question_count']-1


# --------------------------------------------------------------------------
# v1.3.13.11.2 — Unified PDF/DOCX export across LexiCore workspaces
# --------------------------------------------------------------------------

def test_generic_workspace_docx_export(client):
    payload={
        "title":"Compliance & Risk Assessment — Risk Profile",
        "subtitle":"PT Contoh",
        "blocks":[
            {"type":"heading","level":2,"text":"Risk Profile"},
            {"type":"paragraph","text":"Overall Risk Score: 68%"},
            {"type":"table","rows":[["Risiko","Level"],["PDP","HIGH"]]},
        ],
    }
    r=post_json(client,"/api/export/document/docx",payload)
    assert r.status_code==200
    assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in r.content_type
    assert r.data[:2]==b"PK"

def test_generic_workspace_docx_export_requires_content(client):
    r=post_json(client,"/api/export/document/docx",{"title":"Kosong","blocks":[]})
    assert r.status_code==400

def test_workspace_export_controls_exist_in_ui():
    from pathlib import Path
    html=(Path(__file__).resolve().parents[1]/"static"/"index.html").read_text(encoding="utf-8")
    for kind in ("review","corpus","research","risk","norm","client"):
        assert f"exportWorkspacePdf('{kind}')" in html
        assert f"exportWorkspaceDocx('{kind}')" in html
    assert "/export/document/docx" in html


def test_case_analysis_performance_corrective_contract():
    from pathlib import Path
    legal=(Path(__file__).resolve().parents[1]/"legal_sources.py").read_text(encoding="utf-8")
    retrieval=(Path(__file__).resolve().parents[1]/"services"/"regulatory_retrieval.py").read_text(encoding="utf-8")
    route=(Path(__file__).resolve().parents[1]/"routes"/"case_analysis.py").read_text(encoding="utf-8")
    assert "HEALTH_CACHE_TTL" in legal and "FETCH_CACHE_TTL" in legal
    assert 'b"%%EOF" not in tail' in legal
    assert "strict=False" in legal
    assert "max_candidates=2" in retrieval
    assert "case_nexus_status') != 'NO_CASE_NEXUS'" in retrieval
    assert "performance_timing_ms" in route



# --------------------------------------------------------------------------
# v1.3.13.11.2 Regression stabilization guards
# --------------------------------------------------------------------------

def test_case_analysis_frontend_renderers_and_tabs_are_not_regressed():
    from pathlib import Path
    html=(Path(__file__).resolve().parents[1]/"static"/"index.html").read_text(encoding="utf-8")
    for fn in ("caseTab","checkOfficialSources","renderOfficialVerification","renderCaseReadiness","renderCaseWorkingPaper"):
        assert html.count(f"function {fn}(")==1, fn
    assert "jsonWithTimeout" in html
    assert "caseAnalysisInFlight" in html

def test_whatsapp_backend_route_is_not_regressed():
    from pathlib import Path
    app=(Path(__file__).resolve().parents[1]/"app.py").read_text(encoding="utf-8")
    assert "/api/communication/whatsapp-link" in app
    assert "WHATSAPP_COMPOSER_READY" in app

def test_case_analysis_official_search_has_no_nested_query_executor():
    from pathlib import Path
    legal=(Path(__file__).resolve().parents[1]/"legal_sources.py").read_text(encoding="utf-8")
    retrieval=(Path(__file__).resolve().parents[1]/"services"/"regulatory_retrieval.py").read_text(encoding="utf-8")
    route=(Path(__file__).resolve().parents[1]/"routes"/"case_analysis.py").read_text(encoding="utf-8")
    assert "def federated_search_many" in legal
    assert "SEARCH_TIME_BUDGET_EXCEEDED" in legal
    assert "federated_search_many(" in retrieval
    assert "ex.submit(federated_search,q" not in retrieval.replace(" ","")
    assert "case_run_lock" in route
    assert "health_cache_only=True" in route


def test_pytest_backup_dir_is_isolated_from_project():
    """Destructive smoke tests must never write snapshots into project backups/."""
    from pathlib import Path
    configured = Path(os.environ["LEXICORE_BACKUP_DIR"]).resolve()
    project_backups = (Path(__file__).resolve().parents[1] / "backups").resolve()
    assert configured != project_backups
    assert "lexicore_test" in str(configured).lower()


def test_case_regulatory_ui_uses_user_facing_indonesian_labels():
    from pathlib import Path
    html = (Path(__file__).resolve().parents[1] / 'static' / 'index.html').read_text(encoding='utf-8')
    assert 'PENELUSURAN REGULASI TERKAIT PERKARA' in html
    assert 'Alur penyaringan regulasi' in html
    assert 'CASE-SCOPED REGULATORY RETRIEVAL' not in html
    assert 'Retrieval funnel' not in html
    assert 'seed fallback' not in html

# v1.3.13.11.9 — Case Analysis Legal Reasoning Guard benchmark

def test_case_reasoning_guard_eksepsi_tipikor_benchmark():
    from services.case_reasoning_guard import build_reasoning_guard
    text = '''
    EKSEPSI Perkara 92/PID.SUS.TPK/2026/PN.SBY. Terdakwa diajukan dalam perkara Tindak Pidana Korupsi.
    Perbuatan dianggap melanggar Pasal 603 UU Republik Indonesia Nomor 1 Tahun 2003 jo Pasal 18
    UU Nomor 31 Tahun 2099 sebagaimana diubah dengan UU Nomor 20 Tahun 2021 jo UU Nomor 1 Tahun 2026.
    Terdakwa adalah Direktur Utama PD BPR Arta Praja. Pemberian kredit kepada Dewi Mufarida Rp255.000.000
    dijamin Honda Jazz dan disebut dilindungi sertifikat fidusia. Eksepsi menyatakan Pengadilan Tipikor tidak
    berwenang karena perkara ini seharusnya penggelapan atas barang fidusia. Eksepsi juga menyebut Error in Persona,
    salah orang, karena yang seharusnya bertanggung jawab adalah pihak yang dititipi barang jaminan.
    '''
    domain={'domain_contract':['corruption','financial_services','criminal']}
    g=build_reasoning_guard(text,domain)
    assert g['policy']=='FAIL_CLOSED_EVIDENCE_DRIVEN'
    assert g['flags']['corruption_nexus'] is True
    assert g['flags']['banking_nexus'] is True
    assert g['flags']['fiduciary_nexus'] is True
    assert g['tempus']['status']=='TEMPUS_INSUFFICIENT'
    assert len(g['citation_anomalies'])==3
    codes={x['code'] for x in g['boundary_checks']}
    assert 'FORUM_VS_MERITS' in codes
    assert 'ERROR_IN_PERSONA_VS_ATTRIBUTION' in codes
    assert 'FIDUCIARY_NON_DISPOSITIVE' in codes
    assert 'silent correction' not in g['guarded_synthesis'].lower()


def test_case_reasoning_guard_does_not_inject_tipikor_language_into_unrelated_case():
    from services.case_reasoning_guard import build_reasoning_guard
    text='Penggugat mengajukan gugatan wanprestasi atas perjanjian sewa rumah karena Tergugat belum membayar sewa selama tiga bulan.'
    g=build_reasoning_guard(text,{'domain_contract':['civil_contract','civil_procedure']})
    synthesis=g['guarded_synthesis'].lower()
    for forbidden in ('mens rea','actual loss','jaksa','kerugian keuangan negara','kredit/perbankan','tipikor terdeteksi'):
        assert forbidden not in synthesis
    assert g['flags']['corruption_nexus'] is False


def test_case_reasoning_guard_no_silent_citation_normalization():
    from services.case_reasoning_guard import apply_reasoning_guard
    result={'legal_analysis':'AI menyatakan UU 31 Tahun 1999 berlaku final.','legal_issues':[],'recommendations':[],
            'domain_contract':{'domain_contract':['corruption','criminal']}}
    text='Dakwaan Tipikor menyebut UU Nomor 31 Tahun 2099 dan Pasal 603 UU Nomor 1 Tahun 2003.'
    guarded=apply_reasoning_guard(result,text)
    assert guarded['unguarded_legal_analysis'].startswith('AI menyatakan')
    assert guarded['reasoning_guard']['citation_anomalies']
    assert 'berpotensi salah ketik/OCR' in guarded['legal_analysis']
    assert any('surat dakwaan asli' in x.lower() for x in guarded['recommendations'])


# v1.3.13.11.10 — integrated Evidence & Consistency Guard

def test_consistency_guard_rejects_contact_metadata_as_evidence():
    from services.case_consistency_guard import classify_source_item, issue_source_nexus
    item={'label':'SOURCE FACT','statement':'0341-454415, Hp.:081233219270, Email: kantor@example.com'}
    meta=classify_source_item(item)
    assert meta['classification']=='DOCUMENT_METADATA'
    nx=issue_source_nexus('Apakah terdapat hubungan kausal antara keputusan kredit dan kerugian?',item)
    assert nx['eligible'] is False


def test_legal_construction_does_not_pair_issue_by_array_index():
    from services.case_working_paper import build_case_working_paper
    result={
      'legal_issues':['Apakah terdapat kerugian keuangan negara yang nyata dengan memperhitungkan outstanding dan agunan?'],
      'source_ledger':[
        {'label':'SOURCE FACT','statement':'Sentot Yusuf Patrikha, lahir di Malang, NIK: 3507242101560002, Advokat NIA: 02.10801.'},
        {'label':'SOURCE FACT','statement':'Pemberian kredit Rp255.000.000 disebut dijamin kendaraan dan BPKB masih tersimpan di BPR.'},
      ],
      'facts':['Identitas penasihat hukum.'], 'applicable_law':[], 'case_readiness':{'components':{}},
      'evidentiary_gaps':[], 'arguments_for':[], 'arguments_against':[],
      'domain_classification':{'domain_contract':['corruption','financial_services','criminal']},
      'legal_analysis':'Analisis bersyarat.'
    }
    wp=build_case_working_paper(result)
    chain=wp['legal_construction']['chains'][0]
    assert 'Sentot Yusuf' not in chain['material_fact']
    assert '255.000.000' in chain['material_fact'] or 'BPKB' in chain['material_fact']
    assert chain['evidence_nexus_score'] > 0


def test_working_probability_no_provision_bonus_without_verified_applicable_instrument():
    from services.case_working_paper import build_case_working_paper
    result={
      'case_readiness':{'components':{'evidence_map':{'percentage':50},'legal_analysis':{'percentage':50}}},
      'case_regulatory_snapshot':{'official_results':[
        {'positive_law_verification':{
          'legal_status':'IN_FORCE','identity_confirmed':False,'case_nexus_status':'NO_CASE_NEXUS',
          'tempus_status':'TEMPUS_VERIFIED','tempus_applicable':True,'final_status':'VERIFIED_NOT_RELEVANT',
          'provision_verification':{'status':'PROVISION_PARTIALLY_VERIFIED','verified_count':1}
        }}
      ]},
      'evidentiary_gaps':[], 'arguments_for':[], 'arguments_against':[], 'source_ledger':[],
      'domain_classification':{'domain_contract':['criminal']}, 'legal_issues':[], 'applicable_law':[]
    }
    p=build_case_working_paper(result)['working_paper_percentage']
    assert not any(x['variable']=='Dasar hukum resmi yang telah lolos verifikasi' for x in p['variables_increasing'])


def test_instrument_before_article_verification_gate():
    from services.positive_law_verification import verify_document_candidate
    candidate={
      'title':'Perubahan Pasal 18 Undang-Undang Kewarganegaraan','query':'Pasal 18 Tipikor',
      'url':'https://example.invalid/uu','authoritative':True,'case_nexus_status':'NO_CASE_NEXUS',
      'requested_provisions':['Pasal 18']
    }
    source='UNDANG-UNDANG NOMOR 3 TAHUN 1976. Pasal 18 ketentuan kewarganegaraan. Tanggal Berlaku 05-04-1976.'
    v=verify_document_candidate(candidate,source_text=source,snapshot={'event_year_candidate':2022})
    assert (v['provision_verification']['verified_count'] or 0)==0
    assert v['provision_verification']['status']=='PROVISION_BLOCKED_BY_INSTRUMENT_GATE'


def test_case_report_suppresses_irrelevant_official_instruments():
    from services.case_consistency_guard import regulation_reportable
    bad={'document_classification':{'legal_instrument_candidate':True},'positive_law_verification':{'final_status':'VERIFIED_NOT_RELEVANT','case_nexus_status':'NO_CASE_NEXUS'}}
    good={'document_classification':{'legal_instrument_candidate':True},'positive_law_verification':{'final_status':'POTENTIALLY_APPLICABLE','case_nexus_status':'CASE_NEXUS_UNCERTAIN'}}
    assert regulation_reportable(bad) is False
    assert regulation_reportable(good) is True


def test_global_tempus_consistency_does_not_promote_year_from_sk_number():
    from services.case_reasoning_guard import build_reasoning_guard
    from services.case_consistency_guard import apply_global_case_consistency
    text=(
        'Terdakwa dalam perkara Tipikor adalah Direktur Utama PD BPR. '
        'SK Walikota Nomor 188/591/HK/410.010.2/2021 tanggal 05 September 2011. '
        'Pemberian kredit kepada Dewi Mufarida dijamin BPKB dan fidusia. '
        'Pasal 603 UU Nomor 1 Tahun 2003 jo UU Nomor 1 Tahun 2026.'
    )
    guard=build_reasoning_guard(text,{'domain_contract':['corruption','financial_services','criminal']})
    result={
        'source_text':text,
        'reasoning_guard':guard,
        'legal_issues':['Tempus delicti harus diverifikasi.'],
        'legal_analysis':guard['guarded_synthesis'],
        'evidentiary_gaps':[],
        'source_ledger':[{'label':'SOURCE FACT','statement':'Pemberian kredit kepada Dewi Mufarida dijamin BPKB dan fidusia.'}],
    }
    out=apply_global_case_consistency(result)
    joined=' '.join(out['executive_summary'])
    assert 'peristiwa kredit disebut berkaitan dengan tahun 2021' not in joined
    assert 'tahun/tanggal perbuatan material belum dapat ditetapkan' in joined


def test_case_evidence_map_excludes_counsel_address_and_procedural_metadata():
    from services.case_consistency_guard import evidence_rows, material_source_ledger, classify_source_item
    ledger=[
      {'label':'SOURCE FACT','statement':'Kertarejasa Gang XIII Nomor 119 Candirenggo, Singosari Malang, Tlp. 0341-454415'},
      {'label':'SOURCE FACT','statement':'Berdasarkan Surat Kuasa Khusus tanggal 27 Agustus 2026 terdaftar di Kepaniteraan Pengadilan Negeri Surabaya Nomor 1159/HK./VIII/2026'},
      {'label':'SOURCE FACT','statement':'Pemberian Kredit kepada Dewi Mufarida sejumlah Rp255.000.000 melebihi plafon 65% dari hasil penilaian jaminan.'},
      {'label':'SOURCE FACT','statement':'Sementara diketahui buku BPKB kendaraan yang dijaminkan masih tersimpan di BPR Arta Praja.'},
      {'label':'SOURCE FACT','statement':'Eksepsi tentang Error in Persona'},
      {'label':'SOURCE FACT','statement':'Memerintahkan Terdakwa untuk dikeluarkan dari Rumah Tahanan Negara.'},
    ]
    assert classify_source_item(ledger[0])['classification']=='DOCUMENT_METADATA'
    assert classify_source_item(ledger[1])['classification']=='PROCEDURAL_METADATA'
    rows=evidence_rows(ledger,'CRIMINAL')
    text=' '.join(r['fact_proved'] for r in rows)
    assert 'Kertarejasa' not in text
    assert 'Surat Kuasa Khusus' not in text
    assert 'Error in Persona' not in text
    assert 'Memerintahkan Terdakwa' not in text
    assert 'Pemberian Kredit' in text
    assert 'BPKB' in text
    material=' '.join(x['statement'] for x in material_source_ledger(ledger))
    assert 'Kertarejasa' not in material
    assert 'Surat Kuasa Khusus' not in material


def test_case_evidence_map_rejects_hypothetical_legal_argument_as_fact():
    from services.case_consistency_guard import classify_source_item, evidence_rows
    item={'label':'SOURCE FACT','statement':'Kalau barang yang dilindungi sertifikat fidusia ditarik karena kredit macet dan dijual maka persoalan ini selesai.'}
    assert classify_source_item(item)['classification']=='LEGAL_ARGUMENT'
    assert evidence_rows([item],'CRIMINAL')==[]


def test_case_evidence_map_rejects_bare_amount_fragment():
    from services.case_consistency_guard import classify_source_item, evidence_rows
    item={'label':'SOURCE FACT','statement':'255.000.000,- (Dua ratus lima puluh lima juta rupiah).'}
    assert classify_source_item(item)['classification']=='NON_MATERIAL_FRAGMENT'
    assert evidence_rows([item],'CRIMINAL')==[]


def test_case_issue_nexus_never_uses_legal_argument_as_material_fact():
    from services.case_consistency_guard import ranked_support_for_issue
    ledger=[
      {'label':'SOURCE FACT','statement':'Eksepsi Kewenangan Mengadili, adanya missing link atau uraian dakwaan yang terputus mengenai pemberian kredit.'},
      {'label':'SOURCE FACT','statement':'Ditemukan pemberian kredit melebihi plafon 65% dari hasil penilaian jaminan.'},
    ]
    r=ranked_support_for_issue('Apakah penyimpangan prosedur/tata kelola perkreditan terbukti?',ledger,3)
    assert r
    assert all('missing link' not in x['statement'].lower() for x in r)


def test_regulation_reportable_requires_verified_instrument_and_case_nexus():
    from services.case_consistency_guard import regulation_reportable
    noisy={'document_classification':{'legal_instrument_candidate':True},'positive_law_verification':{'identity_confirmed':False,'case_nexus_status':'CASE_NEXUS_UNCERTAIN','final_status':'STATUS_UNCERTAIN'}}
    good={'document_classification':{'legal_instrument_candidate':True},'positive_law_verification':{'identity_confirmed':True,'case_nexus_status':'CASE_NEXUS_VERIFIED','final_status':'VERIFIED_APPLICABLE'}}
    assert regulation_reportable(noisy) is False
    assert regulation_reportable(good) is True


def test_case_readiness_is_not_win_probability():
    from services.case_working_paper import build_case_working_paper
    x={
      'case_readiness':{'components':{'evidence_map':{'percentage':0},'legal_analysis':{'percentage':30}}},
      'evidentiary_gaps':['Bukti primer belum tersedia'],
      'source_ledger':[], 'legal_issues':[], 'case_regulatory_snapshot':{'official_results':[]}
    }
    p=build_case_working_paper(x)['working_paper_percentage']
    assert p['metric']=='CASE_ANALYSIS_READINESS'
    assert 'peluang menang' in p['disclaimer'].lower() or 'bukan peluang' in p['disclaimer'].lower()
    assert 'win_probability' not in str(p).lower()

def test_case_export_privacy_compacts_context_and_uses_legal_disclaimer():
    from exporters.common import _case_export_sections
    x={
      'title':'Uji Privasi',
      'filename':'eksepsi.docx',
      'source_text':'Kepada Yth. Advokat A, NIK: 3507242101560002, Hp: 081233219270, email test@example.com, kesemuanya beralamat kantor di Jl. Kertarejasa 119. KASUS POSISI: Terdakwa didakwa terkait pemberian kredit Rp255.000.000 dan agunan.',
      'case_working_paper':{},
      'coverage_note':'Professional Verification: PENDING.'
    }
    _,sections=_case_export_sections(x)
    text=' '.join(' '.join(v if isinstance(v,list) else [str(v)]) for _,v in sections)
    assert '3507242101560002' not in text
    assert '081233219270' not in text
    assert 'test@example.com' not in text
    assert 'pemberian kredit' in text.lower()
    assert 'medical advice' not in text.lower()

def test_case_export_has_executive_legal_review_first_when_available():
    from exporters.common import _case_export_sections
    x={'professional_review':{'review_status':'HOLD_FOR_VERIFICATION','document_structure':{'document_type':'EKSEPSI_OR_OBJECTION'},'findings':[{'type':'TEMPUS_GAP','severity':'CRITICAL','finding':'Tempus belum pasti','source_text':'Pasal 603 disebut','recommendation':'Verifikasi tanggal perbuatan'}],'strategic_recommendation':{'priorities':['Audit tempus']}}}
    _,sections=_case_export_sections(x)
    titles=[t for t,_ in sections]
    assert titles[0]=='Executive Legal Review'


def test_rc14_soft_edge_secondary_stack_freeze_contract():
    from pathlib import Path
    html=(Path(__file__).resolve().parents[1]/"static"/"index.html").read_text(encoding="utf-8")
    assert 'id="workspaceFrozenStack"' in html
    assert 'id="workspaceFixedHeader"' in html
    assert 'id="workspaceSubnavHost"' in html
    assert '--mobile-frozen-header-h' in html
    assert '--workspace-stack-h' in html
    assert 'syncFrozenHeaderHeight' in html
    assert 'syncWorkspaceSubnav' in html
    assert "nav.dataset.panelOwner=panel.id" in html
    assert "host.appendChild(nav)" in html
    assert '.workspace-frozen-stack' in html
    assert 'soft edge boundary' in html
    assert '.workspace-frozen-stack::after' in html
    assert 'linear-gradient(90deg,transparent,rgba(76,102,128,.20)' in html
    assert 'border-bottom:0' in html
    assert 'position:fixed!important;top:var(--mobile-frozen-header-h)!important' in html
    assert 'calc(var(--mobile-frozen-header-h) + var(--workspace-stack-h) + 12px)' in html
    assert '.workspace-subnav-host .panel-section-nav' in html
    assert 'position:static!important' in html
    assert '.case-workrail' in html
    assert 'overflow-x:auto!important' in html
    assert 'overflow-y:hidden!important' in html


def test_rc15_case_completion_notice_and_user_facing_source_status_contract():
    from pathlib import Path
    html=(Path(__file__).resolve().parents[1]/"static"/"index.html").read_text(encoding="utf-8")
    assert 'function showCaseAnalysisCompleteNotice()' in html
    assert 'Working Paper siap ditinjau.' in html
    assert 'Lihat Working Paper' in html
    assert 'function markCaseWorkingPaperReady()' in html
    assert 'analysis-ready' in html
    assert 'function userConnectivityStatus(row)' in html
    assert 'Akses otomatis dibatasi oleh situs' in html
    assert 'Sumber dapat diakses' in html
    assert 'Sumber belum merespons tepat waktu' in html
    assert "showCaseAnalysisCompleteNotice();loadAll();" in html


def test_rc16_client_communication_guided_lifecycle_contract():
    from pathlib import Path
    html=(Path(__file__).resolve().parents[1]/"static"/"index.html").read_text(encoding="utf-8")
    assert 'id="clientFlowStatus"' in html
    assert 'data-client-step="identity"' in html
    assert 'data-client-step="status"' in html
    assert 'Lanjut: Status & Instruksi' in html
    assert 'Buat Client Document' in html
    assert 'Simpan Draft' in html
    assert 'Lihat Riwayat' in html
    assert 'function clientGoToStatus()' in html
    assert 'function clientValidateIdentity()' in html
    assert 'function clientValidateInstructions()' in html
    assert 'function resetClientFlow()' in html
    assert "window.lexicoreSectionShow=function(panelId,index)" in html
    assert "panel.id==='client'" in html
    assert "if(!clientDraftSaved){const ok=await saveClient({silent:true});if(!ok)return;}" in html


def test_rc17_evidence_taxonomy_distinguishes_pleaded_role_and_argument():
    from services.case_consistency_guard import classify_source_item, material_source_ledger
    role={'label':'ALLEGATION','statement':'Dalam Dakwaan Jaksa Penuntut Umum menyatakan Terdakwa sebagai Direktur Utama PD BPR Arta Praja diangkat berdasarkan SK Walikota.'}
    heading={'label':'SOURCE FACT','statement':'Kewenangan Mengadili'}
    argument={'label':'SOURCE FACT','statement':'Pertanyaan berikutnya mestinya adalah apakah jaminan itu dilindungi Sertipikat Fiducia?'}
    assert classify_source_item(role)['classification']=='ALLEGED_ROLE'
    assert classify_source_item(heading)['classification']=='HEADING_OR_SECTION'
    assert classify_source_item(argument)['classification']=='LEGAL_ARGUMENT'
    material=material_source_ledger([role,heading,argument])
    assert len(material)==1 and material[0]['display_classification']=='ALLEGED_ROLE'


def test_rc17_probative_weight_separates_semantic_relevance_from_sufficiency():
    from services.case_consistency_guard import ranked_support_for_issue
    ledger=[{'label':'SOURCE FACT','statement':'Pemberian kredit disebut melebihi plafon 65% dari nilai jaminan.'}]
    issue='Apakah kerugian keuangan negara telah dihitung sebagai kerugian nyata (actual loss), dengan memperhitungkan pembayaran, outstanding, agunan, dan hasil pemulihan?'
    rows=ranked_support_for_issue(issue,ledger,3)
    assert rows and rows[0]['score'] > 0.4
    assert rows[0]['probative_sufficient'] is False
    assert rows[0]['probative_level']=='LOW'


def test_rc17_legal_construction_blocks_support_when_only_semantic_nexus_exists():
    from services.case_working_paper import build_case_working_paper
    x={'source_ledger':[{'label':'SOURCE FACT','statement':'Pemberian kredit disebut melebihi plafon 65% dari nilai jaminan.'}],
       'legal_issues':['Apakah kerugian keuangan negara telah dihitung sebagai kerugian nyata (actual loss), dengan memperhitungkan pembayaran, outstanding, agunan, dan hasil pemulihan?'],
       'applicable_law':[], 'case_regulatory_snapshot':{'official_results':[]}}
    chain=build_case_working_paper(x)['legal_construction']['chains'][0]
    assert chain['construction_status']=='SEMANTIC_NEXUS_ONLY'
    assert chain['probative_sufficient'] is False


def test_rc17_export_uses_lawyer_facing_language_for_new_guards():
    from exporters.common import _case_export_sections
    x={'title':'Uji','case_working_paper':{'working_paper_percentage':{},'evidence_map':{},'legal_construction':{'chains':[{'issue':'Kerugian','material_fact':'Kredit melebihi plafon','supporting_evidence':'PLEADED_FACT','evidence_reference':'Kredit melebihi plafon','evidence_nexus_score':0.9,'evidence_nexus_reason':'topik terkait','probative_score':0.2,'probative_level':'LOW','probative_reason':'belum membuktikan kerugian nyata','legal_rule':'Dasar hukum spesifik belum terverifikasi.','causal_logic':'Belum cukup','construction_status':'SEMANTIC_NEXUS_ONLY'}]},'action_plan':[]},
       'case_regulatory_snapshot':{'official_results':[],'retrieval_funnel':{'verified_applicable':0}}, 'applicable_law':['UU terkait — perlu verifikasi']}
    _,sections=_case_export_sections(x)
    txt=' '.join(' '.join(v) for _,v in sections if isinstance(v,list))
    assert 'SEMANTIC_NEXUS_ONLY' not in txt
    assert 'PLEADED_FACT' not in txt
    assert 'Keterkaitan topik ditemukan' in txt
    assert any(t=='Kandidat Dasar Hukum yang Perlu Diverifikasi' for t,_ in sections)


def test_case_readiness_does_not_rise_from_identity_or_status_only():
    from services.case_working_paper import build_case_working_paper
    base={
      'case_readiness':{'components':{'evidence_map':{'percentage':0},'legal_analysis':{'percentage':38}}},
      'evidentiary_gaps':[
        'Laporan hasil perhitungan kerugian',
        'Bukti aliran dana',
        'Status pembayaran kredit',
        'Status agunan',
        'SOP kredit',
      ],
      'arguments_for':['Fakta pendukung sementara'],
      'arguments_against':[],
      'source_ledger':[], 'legal_issues':[], 'applicable_law':[],
      'case_regulatory_snapshot':{'official_results':[]}
    }
    p0=build_case_working_paper(base)['working_paper_percentage']
    with_status=dict(base)
    with_status['case_regulatory_snapshot']={'official_results':[
      {'positive_law_verification':{
        'legal_status':'IN_FORCE',
        'identity_confirmed':True,
        'case_nexus_status':'CASE_NEXUS_UNCERTAIN',
        'tempus_status':'TEMPUS_UNVERIFIED',
        'tempus_applicable':False,
        'final_status':'POTENTIALLY_APPLICABLE',
        'provision_verification':{'status':'PROVISION_VERIFIED','verified_count':2},
      }}
    ]}
    p1=build_case_working_paper(with_status)['working_paper_percentage']
    assert p1['percentage'] == p0['percentage']
    assert any(x['variable']=='Dasar hukum positif belum terverifikasi memadai' for x in p1['variables_decreasing'])
    assert not any(x['variable']=='Dasar hukum resmi yang telah lolos verifikasi' for x in p1['variables_increasing'])
