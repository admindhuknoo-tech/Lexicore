from io import BytesIO


def _fake_payload():
    return (
        [("Judul", "Case Analysis")],
        [
            ("Ringkasan Analisis Hukum", [
                "Paragraf naratif yang menjelaskan posisi hukum dan memerlukan format baku untuk keterbacaan profesional.",
                "ISU: Apakah fakta material telah diverifikasi? | STATUS: Perlu verifikasi",
            ]),
            ("Rencana Tindakan", [
                "TINDAKAN: Dapatkan dan uji bukti primer yang diperlukan.",
            ]),
        ],
    )


def test_pdf_uses_f4_compact_working_paper_geometry_and_roman_section_heading(monkeypatch):
    import exporters.case_pdf as mod
    from pypdf import PdfReader

    monkeypatch.setattr(mod, "_case_export_sections", lambda x: _fake_payload())
    monkeypatch.setattr(mod, "prepare_case_export_for_reader", lambda m, s: (m, s))

    bio = mod.export_case_pdf({"title": "Case Analysis", "source_text": ""})
    reader = PdfReader(bio)
    page = reader.pages[0]
    width_mm = float(page.mediabox.width) * 25.4 / 72.0
    height_mm = float(page.mediabox.height) * 25.4 / 72.0
    assert abs(width_mm - 215.0) < 0.5
    assert abs(height_mm - 330.0) < 0.5
    text = " ".join((page.extract_text() or "").split())
    assert "I. RINGKASAN ANALISIS HUKUM" in text


def test_docx_uses_f4_compact_reference_profile(monkeypatch):
    import exporters.case_docx as mod
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    monkeypatch.setattr(mod, "_case_export_sections", lambda x: _fake_payload())
    monkeypatch.setattr(mod, "prepare_case_export_for_reader", lambda m, s: (m, s))

    bio = mod.export_case_docx({"title": "Case Analysis", "source_text": ""})
    doc = Document(bio)
    sec = doc.sections[0]

    assert abs(sec.page_width.mm - 215.0) < 0.5
    assert abs(sec.page_height.mm - 330.0) < 0.5
    assert abs(sec.top_margin.mm - 24.0) < 0.5
    assert abs(sec.left_margin.mm - 28.0) < 0.5
    assert abs(sec.bottom_margin.mm - 22.0) < 0.5
    assert abs(sec.right_margin.mm - 22.0) < 0.5

    normal = doc.styles["Normal"]
    assert normal.font.name == "Times New Roman"
    assert abs(normal.font.size.pt - 10.5) < 0.1
    assert normal.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert abs(float(normal.paragraph_format.line_spacing) - 1.15) < 0.01
    assert normal.paragraph_format.first_line_indent.mm == 0

    headings = [p for p in doc.paragraphs if p.text.startswith("I. ")]
    assert headings
    hp = headings[0]
    assert hp.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert hp.runs[0].bold is True
    assert hp.runs[0].font.name == "Times New Roman"
    assert abs(hp.runs[0].font.size.pt - 12.5) < 0.1


def test_structured_legal_lines_are_not_first_line_indented(monkeypatch):
    import exporters.case_docx as mod
    from docx import Document

    monkeypatch.setattr(mod, "_case_export_sections", lambda x: _fake_payload())
    monkeypatch.setattr(mod, "prepare_case_export_for_reader", lambda m, s: (m, s))

    doc = Document(mod.export_case_docx({"title": "Case Analysis", "source_text": ""}))
    issue = next(p for p in doc.paragraphs if p.text.startswith("ISU:"))
    narrative = next(p for p in doc.paragraphs if p.text.startswith("Paragraf naratif"))
    assert issue.paragraph_format.first_line_indent.mm == 0
    assert narrative.paragraph_format.first_line_indent.mm == 0
