"""LexiCore - Lawyer Operating System API."""
import os, re, shutil, tempfile
from datetime import datetime
from io import BytesIO
from flask import Flask, request, jsonify, send_file, send_from_directory
from flask_cors import CORS
from werkzeug.utils import secure_filename
from docx import Document

from env_loader import load_project_env
load_project_env()

from contract_review import ContractReviewEngine
from database import (init_database, DraftManager, AnalysisManager, ResearchManager,
                      RiskAssessmentManager, CommunicationManager, AuditLogger, RegulatoryCorpusManager,
                      get_db_connection, SCHEMA_VERSION)
from legal_sources import source_health, federated_search, public_source_registry, public_secondary_source_registry, supplementary_search_many
from ai_engine import status as ai_status
from regulatory_db import get_all_regulations, search_regulations, retrieve_for_case
from norm_conflict import analyze_conflicts
from regulatory_intelligence import normalized_catalog, graph_for_query, timeline_for_regulation, resolve_conflict_from_query, intelligence_for_case
from services.backup_service import create_database_backup
from version import LEXICORE_VERSION, PRODUCT_NAME, PRODUCT_LABEL, INITIATIVE, FIRM_NAME, release_metadata
from routes.case_analysis import create_case_analysis_blueprint
from legal_drafting import build_legal_draft, template_catalog, TEMPLATE_REGISTRY
from services.case_domain_classifier import classify_case
from services.legal_research_service import search_legal_authorities
from services.case_law_summary import summarize_case_source, summarize_doctrine_source, CASE_TYPES, DOCTRINE_TYPES
from services.compliance_risk import build_risk_matrix, question_set, normalize_category
from services.case_consistency_guard import executive_fact_candidates, classify_source_item
from client_communication import build_client_communication

BASE_DIR=os.path.dirname(os.path.abspath(__file__))
app=Flask(__name__, static_folder='static')
# Security fix (v1.3.5.4): `CORS(app)` with no restriction reflects
# Access-Control-Allow-Origin for ANY requesting origin. This API has zero
# authentication (see the app.run() note below), and its own frontend is
# served from this exact same Flask origin — so cross-origin access has no
# legitimate use here. Left wide open, any website the lawyer happens to
# have open in another tab could silently fetch() this locally running API
# and read or delete confidential client/case data (drafts, contract
# analyses, full case narratives with BAP/financial details) without the
# user ever visiting LexiCore itself. That's a realistic "malicious webpage
# talks to your local server" attack, not a theoretical one, and it
# defeats the network-exposure guard further down in this file, which only
# protects against LAN attackers, not against the browser itself.
# Cross-origin access is now opt-in only, via a comma-separated allowlist in
# LEXICORE_CORS_ORIGINS (e.g. for a separate frontend dev server). By
# default no origin is allowed, which matches how the app is actually used.
_cors_origins = [o.strip() for o in os.environ.get('LEXICORE_CORS_ORIGINS', '').split(',') if o.strip()]
if _cors_origins:
    CORS(app, origins=_cors_origins)
    print(f'LexiCore: CORS diaktifkan untuk origin: {_cors_origins}')
app.config.update(MAX_CONTENT_LENGTH=50*1024*1024, UPLOAD_FOLDER=os.path.join(BASE_DIR,'uploads'), ALLOWED_EXTENSIONS={'pdf','docx','png','jpg','jpeg','webp','tif','tiff'})
os.makedirs(app.config['UPLOAD_FOLDER'],exist_ok=True)
init_database()


def allowed_file(filename):
    return '.' in filename and filename.rsplit('.',1)[1].lower() in app.config['ALLOWED_EXTENSIONS']


@app.route('/api/ocr/status', methods=['GET'])
def ocr_status_endpoint():
    from services.document_ocr import ocr_runtime_status
    return jsonify(success=True, data=ocr_runtime_status(load_engine=False))

def risk_level(risks):
    high=sum(1 for r in risks if r.get('risk_level')=='HIGH'); med=sum(1 for r in risks if r.get('risk_level')=='MEDIUM')
    return 'HIGH' if high>=2 else ('MEDIUM' if high or med>=3 else 'LOW')

def sentences(text):
    return [s.strip() for s in re.split(r'(?<=[.!?])\s+', text or '') if len(s.strip())>20]


# v1.3.3.13 — review & cleanup helpers
HISTORY_TABLES = {
    'drafts': 'drafts',
    'contract_reviews': 'contract_analyses',
    'research_notes': 'legal_research',
    'risk_assessments': 'risk_assessments',
    'case_analyses': 'case_analyses',
    'client_documents': 'client_communications',
}

def _managed_upload_delete(file_path):
    """Delete only files physically located under LexiCore's managed uploads directory."""
    if not file_path:
        return False
    try:
        root=os.path.realpath(app.config['UPLOAD_FOLDER'])
        candidate=os.path.realpath(file_path if os.path.isabs(file_path) else os.path.join(BASE_DIR,file_path))
        if os.path.commonpath([root,candidate]) != root:
            return False
        if os.path.isfile(candidate):
            os.remove(candidate); return True
    except Exception:
        return False
    return False

def _history_row(kind, item_id):
    table=HISTORY_TABLES.get(kind)
    if not table:
        return None
    conn=get_db_connection(); row=conn.execute(f'SELECT * FROM {table} WHERE id=?',(item_id,)).fetchone(); conn.close()
    if not row:
        return None
    d=dict(row)
    json_fields={
        'contract_reviews': ('parties','risks'),
        'research_notes': ('keywords',),
        'risk_assessments': ('answers','findings','recommendations'),
        'case_analyses': ('facts','legal_issues','applicable_law','arguments_for','arguments_against','evidence_needed','risks','recommendations','domain_classification','analysis_provenance','case_readiness'),
    }.get(kind,())
    import json as _json
    dict_json_fields={'answers','domain_classification','analysis_provenance','case_readiness'}
    for key in json_fields:
        try: d[key]=_json.loads(d.get(key) or ('{}' if key in dict_json_fields else '[]'))
        except Exception: d[key] = {} if key in dict_json_fields else []
    if kind=='case_analyses':
        try:
            from database import CaseRegulatorySnapshotManager
            d['case_regulatory_snapshot']=CaseRegulatorySnapshotManager.latest(item_id)
        except Exception:
            d['case_regulatory_snapshot']=None
    return d

def _delete_history_item(kind, item_id):
    table=HISTORY_TABLES.get(kind)
    if not table:
        return False
    conn=get_db_connection(); cur=conn.cursor()
    if kind=='contract_reviews':
        row=cur.execute('SELECT file_path FROM contract_analyses WHERE id=?',(item_id,)).fetchone()
        if row: _managed_upload_delete(row['file_path'])
    if kind=='case_analyses':
        cur.execute('DELETE FROM legal_source_verifications WHERE case_analysis_id=?',(item_id,))
        cur.execute('DELETE FROM case_regulatory_snapshots WHERE case_analysis_id=?',(item_id,))
    cur.execute(f'DELETE FROM {table} WHERE id=?',(item_id,)); ok=cur.rowcount>0
    conn.commit(); conn.close()
    if ok: AuditLogger.log_action('delete_history_item',kind,item_id,'Deleted saved working-history record')
    return ok

def _delete_history_kind(kind):
    table=HISTORY_TABLES.get(kind)
    if not table:
        return 0
    create_database_backup(os.environ.get('LEXICORE_DB_PATH','lexicore.db'), reason=f'pre_delete_{kind}')
    conn=get_db_connection(); cur=conn.cursor()
    if kind=='contract_reviews':
        rows=cur.execute('SELECT file_path FROM contract_analyses').fetchall()
        for row in rows: _managed_upload_delete(row['file_path'])
    if kind=='case_analyses':
        cur.execute('DELETE FROM legal_source_verifications')
        cur.execute('DELETE FROM case_regulatory_snapshots')
    count=cur.execute(f'SELECT COUNT(*) AS n FROM {table}').fetchone()['n']
    cur.execute(f'DELETE FROM {table}')
    conn.commit(); conn.close()
    if count: AuditLogger.log_action('delete_history_all',kind,None,f'Deleted {count} saved records')
    return int(count or 0)

@app.route('/api/history/<kind>/<int:item_id>', methods=['GET','DELETE'])
def history_item(kind,item_id):
    if kind not in HISTORY_TABLES:
        return jsonify(success=False,error='Jenis riwayat tidak dikenal'),404
    if request.method=='GET':
        row=_history_row(kind,item_id)
        return jsonify(success=True,data=row) if row else (jsonify(success=False,error='Riwayat tidak ditemukan'),404)
    ok=_delete_history_item(kind,item_id)
    return jsonify(success=ok,deleted=1 if ok else 0)

@app.route('/api/history/<kind>', methods=['DELETE'])
def history_clear_kind(kind):
    if kind=='norm_conflicts':
        conn=get_db_connection(); cur=conn.cursor(); row=cur.execute("SELECT COUNT(*) AS n FROM audit_log WHERE action='norm_conflict_analysis'").fetchone(); n=int(row['n'] or 0); cur.execute("DELETE FROM audit_log WHERE action='norm_conflict_analysis'"); conn.commit(); conn.close();
        if n: AuditLogger.log_action('delete_history_all','norm_conflicts',None,f'Deleted {n} norm-conflict activity records')
        return jsonify(success=True,deleted=n)
    if kind not in HISTORY_TABLES:
        return jsonify(success=False,error='Jenis riwayat tidak dikenal'),404
    return jsonify(success=True,deleted=_delete_history_kind(kind))

@app.route('/api/history', methods=['DELETE'])
def history_clear_all():
    d=request.get_json(silent=True) or {}
    if d.get('confirm')!='HAPUS SEMUA RIWAYAT':
        return jsonify(success=False,error='Konfirmasi penghapusan tidak cocok'),400
    order=['client_documents','case_analyses','risk_assessments','research_notes','contract_reviews','drafts']
    deleted={k:_delete_history_kind(k) for k in order}
    conn=get_db_connection(); cur=conn.cursor(); row=cur.execute("SELECT COUNT(*) AS n FROM audit_log WHERE action='norm_conflict_analysis'").fetchone(); n=int(row['n'] or 0); cur.execute("DELETE FROM audit_log WHERE action='norm_conflict_analysis'"); conn.commit(); conn.close(); deleted['norm_conflicts']=n
    AuditLogger.log_action('purge_working_history','history',None,'User cleared all saved drafts and working-history records')
    return jsonify(success=True,deleted=deleted)

@app.route('/')
def index(): return send_from_directory(app.static_folder,'index.html')

@app.route('/api/drafts',methods=['GET','POST'])
def drafts():
    if request.method=='GET':
        data=DraftManager.get_all_drafts(request.args.get('limit',50,type=int)); return jsonify(success=True,data=data,count=len(data))
    data=request.get_json(silent=True) or {}
    if not data.get('content','').strip(): return jsonify(success=False,error='Konten draft wajib diisi'),400
    i=DraftManager.create_draft(data); return jsonify(success=True,draft_id=i,message='Draft berhasil disimpan')

@app.route('/api/drafts/<int:draft_id>',methods=['GET','PUT','DELETE'])
def draft_detail(draft_id):
    if request.method=='GET':
        d=DraftManager.get_draft(draft_id); return (jsonify(success=True,data=d) if d else (jsonify(success=False,error='Draft tidak ditemukan'),404))
    if request.method=='DELETE':
        return jsonify(success=DraftManager.delete_draft(draft_id))
    ok=DraftManager.update_draft(draft_id,request.get_json(silent=True) or {}); return jsonify(success=ok)

def _build_legal_draft(doc_type, p1, p2, date, duration, prompt):
    return build_legal_draft(doc_type,p1,p2,date,duration,prompt)


@app.route('/api/drafting/templates')
def drafting_templates():
    return jsonify(success=True,data=template_catalog(),count=len(TEMPLATE_REGISTRY))


@app.route('/api/generate/draft',methods=['POST'])
def generate_draft():
    data=request.get_json(silent=True) or {}
    doc_type=(data.get('doc_type') or 'Perjanjian Kerjasama').strip()
    p1=(data.get('party1') or '[PIHAK / KLIEN PERTAMA]').strip()
    p2=(data.get('party2') or '[PIHAK / SUBJEK KEDUA]').strip()
    date=data.get('effective_date') or datetime.now().strftime('%Y-%m-%d')
    try: duration=max(1,int(data.get('duration') or 12))
    except (TypeError,ValueError): duration=12
    prompt=(data.get('prompt') or '').strip()
    if doc_type == 'Surat Kuasa': doc_type='Surat Kuasa Khusus'
    if doc_type not in TEMPLATE_REGISTRY: return jsonify(success=False,error='Jenis dokumen tidak didukung'),400
    content,cc=_build_legal_draft(doc_type,p1,p2,date,duration,prompt)
    wc=len(content.split())
    i=DraftManager.create_draft(dict(title=f'{doc_type} — {p1} / {p2}',doc_type=doc_type,party1=p1,party2=p2,effective_date=date,duration=duration,content=content,word_count=wc,clause_count=cc,status='generated'))
    return jsonify(success=True,draft_id=i,content=content,word_count=wc,clause_count=cc,template_family=doc_type,message='Draft berhasil digenerate sesuai jenis dokumen')


@app.route('/api/drafts/<int:draft_id>/export/docx')
def export_docx(draft_id):
    d=DraftManager.get_draft(draft_id)
    if not d: return jsonify(success=False,error='Draft tidak ditemukan'),404
    doc=Document(); doc.add_heading(d['title'],0)
    for block in d['content'].split('\n\n'):
        if block.strip(): doc.add_paragraph(block.strip())
    bio=BytesIO(); doc.save(bio); bio.seek(0); safe=secure_filename(d['title']) or 'lexicore-draft'
    return send_file(bio,as_attachment=True,download_name=f'{safe}.docx',mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')


@app.route('/api/export/document/docx', methods=['POST'])
def export_workspace_docx():
    """Generic DOCX exporter for non-drafting LexiCore workspaces.

    The frontend sends already-rendered, user-visible content as structured blocks.
    This keeps export logic presentation-only: it does not rerun legal analysis or
    alter any substantive result before export.
    """
    data=request.get_json(silent=True) or {}
    title=(data.get('title') or 'LexiCore Working Paper').strip()
    subtitle=(data.get('subtitle') or '').strip()
    blocks=data.get('blocks') or []
    if not isinstance(blocks,list) or not blocks:
        return jsonify(success=False,error='Belum ada hasil yang dapat diekspor'),400

    doc=Document()
    doc.add_heading(title,0)
    if subtitle:
        p=doc.add_paragraph(subtitle)
        try: p.style='Subtitle'
        except Exception: pass

    for block in blocks[:500]:
        if not isinstance(block,dict):
            continue
        kind=(block.get('type') or 'paragraph').lower()
        if kind=='heading':
            text=str(block.get('text') or '').strip()
            if text:
                level=max(1,min(int(block.get('level') or 2),4))
                doc.add_heading(text,level=level)
        elif kind=='table':
            rows=block.get('rows') or []
            rows=[r for r in rows if isinstance(r,list)]
            if not rows: continue
            width=max((len(r) for r in rows), default=0)
            if not width: continue
            table=doc.add_table(rows=1, cols=width)
            table.style='Table Grid'
            for j,val in enumerate(rows[0]):
                table.rows[0].cells[j].text=str(val or '')
            for row in rows[1:]:
                cells=table.add_row().cells
                for j in range(width):
                    cells[j].text=str(row[j] if j < len(row) else '')
        elif kind=='list':
            items=block.get('items') or []
            for item in items:
                text=str(item or '').strip()
                if text: doc.add_paragraph(text,style='List Bullet')
        else:
            text=str(block.get('text') or '').strip()
            if text: doc.add_paragraph(text)

    doc.add_paragraph('')
    footer=doc.add_paragraph('LexiCore Assistant by ELF — Erfan’s Law Firm | Working document — professional verification required before legal use.')
    try: footer.style='Caption'
    except Exception: pass
    bio=BytesIO(); doc.save(bio); bio.seek(0)
    safe=secure_filename(title) or 'lexicore-working-paper'
    return send_file(bio,as_attachment=True,download_name=f'{safe}.docx',mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

@app.route('/api/review',methods=['POST'])
def review_contract():
    f=request.files.get('file')
    if not f or not f.filename: return jsonify(success=False,error='Pilih file PDF/DOCX'),400
    if not allowed_file(f.filename): return jsonify(success=False,error='Hanya PDF dan DOCX yang didukung'),400
    td=tempfile.mkdtemp()
    try:
        path=os.path.join(td,secure_filename(f.filename)); f.save(path)
        # Bug fix (v1.3.5.4): a malformed/corrupt/encrypted PDF or DOCX made
        # DocumentExtractor raise PdfReadError/PackageNotFoundError (or similar)
        # with nothing catching it here. Flask then returned an unhandled 500
        # HTML error page instead of JSON, which the frontend cannot parse —
        # the user just saw a broken request with no explanation. Extraction
        # and analysis failures are now caught and reported the same way as
        # every other validation error in this API.
        try:
            result=ContractReviewEngine.review(path)
        except ValueError as exc:
            return jsonify(success=False,error=str(exc)),400
        except Exception:
            return jsonify(success=False,error='Dokumen tidak dapat dibaca. File mungkin rusak, terenkripsi/berpassword, atau formatnya tidak valid. Coba unggah ulang atau gunakan salinan lain.'),422
        result['risk_score']=risk_level(result.get('risks',[])); result['analysis_id']=AnalysisManager.save_analysis(result); return jsonify(success=True,data=result)
    finally: shutil.rmtree(td,ignore_errors=True)

@app.route('/api/analyses')
def analyses():
    data=AnalysisManager.get_all_analyses(request.args.get('limit',50,type=int)); return jsonify(success=True,data=data,count=len(data))

@app.route('/api/research/search',methods=['POST'])
def research_search():
    d=request.get_json(silent=True) or {}
    q=(d.get('query') or d.get('citation') or d.get('title') or '').strip()
    if len(q)<3: return jsonify(success=False,error='Masukkan query dasar hukum minimal 3 karakter'),400
    mode=(d.get('mode') or 'hybrid').strip().lower()
    focus=(d.get('focus') or 'mixed').strip().lower()
    direct_ids=None
    if focus in {'jurisprudence','case_law'}:
        direct_ids=('ma','mk','jdih_mk')
    result=search_legal_authorities(q,mode=mode,limit=max(1,min(int(d.get('limit') or 12),30)),direct_ids=direct_ids)
    result['research_focus']=focus
    if focus in {'jurisprudence','case_law'}:
        result['research_note']='Fokus yurisprudensi: sumber resmi MA/MK diprioritaskan. Ringkasan putusan hanya dibentuk dari teks putusan yang dimasukkan/diambil, bukan dari judul hasil pencarian.'
    elif focus=='doctrine':
        result['research_note']='Fokus doktrin: masukkan teks buku/jurnal/pendapat ahli yang dapat dipertanggungjawabkan. LexiCore tidak memperlakukan doktrin sebagai sumber hukum primer.'
    return jsonify(success=True,data=result)


@app.route('/api/research/summarize',methods=['POST'])
def research_summarize():
    d=request.get_json(silent=True) or {}; text=(d.get('source_text') or '').strip()
    if len(text)<80: return jsonify(success=False,error='Masukkan teks sumber hukum minimal 80 karakter'),400

    source_type=d.get('source_type','putusan_pengadilan')
    source_labels={
      'kuhp_2023':'KUHP Nasional — UU No. 1 Tahun 2023 sebagaimana disesuaikan/diubah dengan UU No. 1 Tahun 2026',
      'penyesuaian_pidana_2026':'UU No. 1 Tahun 2026 tentang Penyesuaian Pidana',
      'kuhperdata_bw':'Kitab Undang-Undang Hukum Perdata (KUHPerdata / Burgerlijk Wetboek)',
      'uud1945':'UUD Negara Republik Indonesia Tahun 1945','tap_mpr':'Ketetapan MPR (TAP MPR)',
      'uu':'Undang-Undang (UU)','perppu':'Peraturan Pemerintah Pengganti Undang-Undang (Perppu)',
      'pp':'Peraturan Pemerintah (PP)','perpres':'Peraturan Presiden (Perpres)',
      'perda_prov':'Peraturan Daerah Provinsi','perda_kab_kota':'Peraturan Daerah Kabupaten/Kota',
      'per_mpr':'Peraturan MPR','per_dpr':'Peraturan DPR','per_dpd':'Peraturan DPD','per_bpk':'Peraturan BPK',
      'per_bi':'Peraturan Bank Indonesia','per_menteri':'Peraturan Menteri','per_lembaga':'Peraturan Badan / Lembaga / Komisi',
      'per_gubernur':'Peraturan Gubernur','per_bupati':'Peraturan Bupati','per_walikota':'Peraturan Wali Kota','per_desa':'Peraturan Desa / yang setingkat',
      'putusan_ma':'Putusan Mahkamah Agung','putusan_mk':'Putusan Mahkamah Konstitusi','putusan_pengadilan':'Putusan Pengadilan',
      'perma':'Peraturan Mahkamah Agung (PERMA)','sema':'Surat Edaran Mahkamah Agung (SEMA)',
      'per_mk':'Peraturan Mahkamah Konstitusi (PMK)','per_ky':'Peraturan Komisi Yudisial','reg_yudisial_lain':'Regulasi / Pedoman Instansi Yudikatif Lainnya',
      'legal_opinion':'Legal Opinion / Pendapat Hukum','doctrine':'Doktrin / Buku / Jurnal','memo':'Memo / Literatur',
      'case':'Putusan Pengadilan','statute':'Peraturan Perundang-undangan'
    }
    case_types={'putusan_ma','putusan_mk','putusan_pengadilan','case'}
    secondary_types={'legal_opinion','doctrine','memo'}
    ss=sentences(text)

    structured={}
    if source_type in case_types:
        structured=summarize_case_source(text)
        issue=structured['chronology']
        holding=structured['disposition']
        reasoning=structured['ratio_decidendi']
        issue_label='KRONOLOGI KASUS'; result_label='AMAR PUTUSAN'; reasoning_label='PERTIMBANGAN HUKUM HAKIM / RATIO DECIDENDI'
    elif source_type in secondary_types:
        structured=summarize_doctrine_source(text)
        issue=structured['doctrine_topic']
        holding=structured['doctrine_thesis']
        reasoning=structured['doctrine_analysis']
        issue_label='TOPIK / ISU DOKTRIN'; result_label='POKOK PENDAPAT / TESIS'; reasoning_label='ARGUMENTASI / ANALISIS DOKTRIN'
    else:
        issue=d.get('issue') or next((s for s in ss if any(k in s.lower() for k in ['menimbang','mengingat','berdasarkan','ketentuan','mengatur','dimaksud'])), ss[0] if ss else '')
        holding=next((s for s in ss if any(k in s.lower() for k in ['wajib','dilarang','berhak','berwenang','harus','dapat','ditetapkan','diatur'])), ss[1] if len(ss)>1 else (ss[0] if ss else ''))
        reasoning=' '.join(ss[1:5])[:1200] if len(ss)>1 else text[:1200]
        issue_label='RUANG LINGKUP / ISU PENGATURAN'; result_label='NORMA / POKOK PENGATURAN'; reasoning_label='KONTEKS / ANALISIS'

    words=re.findall(r'\b[a-zA-ZÀ-ÿ]{5,}\b',text.lower()); stop={'dengan','dalam','bahwa','untuk','adalah','karena','kepada','tersebut','sebagai','pihak','serta','yang','dari'}; freq={w:words.count(w) for w in set(words) if w not in stop}; keywords=[w for w,_ in sorted(freq.items(),key=lambda x:x[1],reverse=True)[:8]]
    record=dict(title=d.get('title') or 'Legal Research Note',jurisdiction=d.get('jurisdiction','Indonesia'),citation=d.get('citation',''),source_type=source_type,source_text=text,issue=issue,holding=holding,reasoning=reasoning,keywords=keywords,research_payload=structured)
    record['research_id']=ResearchManager.save(record)
    record.update(source_label=source_labels.get(source_type,source_type),issue_label=issue_label,result_label=result_label,reasoning_label=reasoning_label)
    return jsonify(success=True,data=record)

@app.route('/api/research')
def research_list(): return jsonify(success=True,data=ResearchManager.all(request.args.get('limit',30,type=int)))

@app.route('/api/compliance/questions')
def compliance_questions():
    category=normalize_category(request.args.get('category','General Corporate'))
    return jsonify(success=True,data={'category':category,'questions':question_set(category)})

@app.route('/api/compliance/assess',methods=['POST'])
def compliance_assess():
    d=request.get_json(silent=True) or {}
    answers=d.get('answers') or {}
    custom_controls=d.get('custom_controls') or []
    category=normalize_category(d.get('category','General Corporate'))
    assessed=build_risk_matrix(category,answers,custom_controls)
    findings=[r['risk_identification'] for r in assessed['matrix'] if r['risk_level'] in {'HIGH','MEDIUM'}]
    recommendations=[]
    for row in assessed['matrix']:
        if row['risk_level'] in {'HIGH','MEDIUM'}:
            recommendations.extend(row.get('mitigation_checklist') or [])
    # retain the legacy fields for compatibility while persisting the richer matrix.
    record=dict(
        title=d.get('title','Compliance & Risk Assessment'),
        entity=d.get('entity',''),
        category=category,
        answers=answers,
        score=assessed['score'],
        risk_level=assessed['risk_level'],
        findings=findings,
        recommendations=list(dict.fromkeys(recommendations)),
        risk_matrix=assessed['matrix'],
        professional_verification='PENDING',
        assessment_method='CATEGORY_SPECIFIC_DETERMINISTIC_MATRIX',
        custom_controls=custom_controls,
        system_question_count=assessed.get('system_question_count',0),
        custom_question_count=assessed.get('custom_question_count',0),
    )
    if not bool(d.get('preview')):
        record['assessment_id']=RiskAssessmentManager.save(record)
    else:
        record['preview']=True
    return jsonify(success=True,data=record)

@app.route('/api/compliance')
def compliance_list(): return jsonify(success=True,data=RiskAssessmentManager.all(request.args.get('limit',30,type=int)))

# Case Analysis uses a local issue-spotting/rule registry. It does NOT claim an exhaustive live corpus.
CASE_LAW_REGISTRY = [
    {'domain':'Konstitusi & Hak Dasar','keywords':['konstitusi','hak asasi','diskriminasi','kebebasan','kewenangan negara','uji materi'],'sources':['UUD Negara Republik Indonesia Tahun 1945','Putusan Mahkamah Konstitusi yang relevan']},
    {'domain':'Pidana Materiil','keywords':['pidana','tersangka','terdakwa','dakwaan','penipuan','penggelapan','pencurian','penganiayaan','ancaman','pemalsuan','korupsi'],'sources':['UU No. 1 Tahun 2023 tentang KUHP','UU No. 1 Tahun 2026 tentang Penyesuaian Pidana','Undang-undang pidana khusus yang relevan']},
    {'domain':'Acara Pidana','keywords':['penyidikan','penyelidikan','penangkapan','penahanan','penggeledahan','penyitaan','praperadilan','penuntutan','saksi','barang bukti'],'sources':['UU No. 20 Tahun 2025 tentang Kitab Undang-Undang Hukum Acara Pidana','PERMA/SEMA dan putusan pengadilan yang relevan']},
    {'domain':'Tindak Pidana Korupsi','keywords':['tindak pidana korupsi','tipikor','kerugian keuangan negara','kerugian keuangan daerah','penyalahgunaan kewenangan'],'sources':['UU No. 31 Tahun 1999 jo UU No. 20 Tahun 2001 tentang Pemberantasan Tindak Pidana Korupsi','KUHP Nasional/ketentuan transisi yang relevan menurut tempus']},
    {'domain':'Perbankan & Jasa Keuangan','keywords':['bpr','bank perekonomian rakyat','bank perkreditan rakyat','fasilitas kredit','pemberian kredit','komite kredit','analisa 5c','slik','ojk'],'sources':['UU Perbankan beserta perubahannya yang berlaku pada tempus','POJK/SEOJK yang berlaku pada tempus dan relevan dengan BPR/kredit']},
    {'domain':'BUMD & Pemerintahan Daerah','keywords':['perumda','bumd','pemerintah daerah','pemerintah kota','kpm','dewan pengawas'],'sources':['UU Pemerintahan Daerah beserta perubahannya','PP tentang BUMD dan peraturan daerah/pendirian Perumda yang relevan']},
    {'domain':'Perdata & Perikatan','keywords':['wanprestasi','perjanjian','utang','piutang','ganti rugi','jual beli','sewa','perbuatan melawan hukum','kontrak'],'sources':['KUHPerdata / Burgerlijk Wetboek (periksa status pasal secara spesifik)','Undang-undang sektoral yang menyimpangi atau menggantikan BW','Yurisprudensi yang relevan']},
    {'domain':'Perusahaan & Komersial','keywords':['perseroan','direksi','komisaris','pemegang saham','pt ','perusahaan','merger','akuisisi','usaha'],'sources':['Undang-undang perseroan/perusahaan yang berlaku','Peraturan pelaksana dan regulasi sektor usaha yang relevan']},
    {'domain':'Ketenagakerjaan','keywords':['pekerja','buruh','karyawan','phk','upah','hubungan kerja','pesangon'],'sources':['Peraturan ketenagakerjaan yang berlaku beserta perubahan dan peraturan pelaksananya','Putusan PHI/MA yang relevan']},
    {'domain':'Pertanahan & Properti','keywords':['tanah','sertifikat','sertipikat','shm','hgb','hak milik','waris tanah','bpn','kantor pertanahan','surat ukur','data fisik','data yuridis','agraria'],'sources':['UUPA dan peraturan pertanahan yang berlaku','Peraturan pendaftaran tanah/ATR-BPN dan putusan yang relevan']},
    {'domain':'Hukum Acara Perdata','keywords':['gugatan','penggugat','tergugat','eksepsi','obscuur libel','plurium litis consortium','rekonvensi','kompetensi absolut','kewenangan absolut','pasal 8 rv'],'sources':['HIR/RBg/Rv sepanjang relevan dan masih digunakan sebagai sumber hukum acara','PERMA/SEMA serta yurisprudensi hukum acara perdata yang relevan']},
    {'domain':'Keluarga & Peradilan Agama','keywords':['perkawinan','cerai','talak','nafkah','hak asuh','waris islam','wakaf','pengadilan agama','peradilan agama','pasal 49 undang-undang nomor 3 tahun 2006','pasal 49 uu'],'sources':['UU Peradilan Agama beserta perubahan yang relevan','Peraturan perkawinan/keluarga jika materi sengketa memerlukannya','PERMA/SEMA dan putusan peradilan agama yang relevan']},
    {'domain':'Konsumen','keywords':['konsumen','pelaku usaha','garansi','produk','jasa','kerugian konsumen'],'sources':['Undang-undang perlindungan konsumen dan peraturan pelaksananya']},
    {'domain':'Teknologi, ITE & Data','keywords':['data pribadi','privasi','elektronik','internet','digital','sistem elektronik','informasi elektronik'],'sources':['Peraturan perlindungan data pribadi yang berlaku','Peraturan informasi dan transaksi elektronik yang berlaku beserta perubahannya','Peraturan pelaksana sektor digital yang relevan']},
    {'domain':'Administrasi Pemerintahan','keywords':['keputusan tata usaha','pejabat','izin','perizinan','administrasi pemerintahan','ptun'],'sources':['Peraturan administrasi pemerintahan dan peradilan tata usaha negara yang berlaku','Peraturan sektoral pemberi kewenangan']},
]

def _case_issue_spot(text):
    low=text.lower(); matched=[]
    for item in CASE_LAW_REGISTRY:
        hits=[k for k in item['keywords'] if k in low]
        if hits: matched.append((len(hits),item,hits))
    matched.sort(key=lambda x:x[0],reverse=True)
    return matched


def _unique(items, limit=None):
    out=[]; seen=set()
    for item in items:
        key=re.sub(r'\s+',' ',str(item or '')).strip().lower()
        if not key or key in seen: continue
        seen.add(key); out.append(item)
        if limit and len(out)>=limit: break
    return out


def _money_values(text):
    vals=[]
    for m in re.finditer(r'Rp\.?\s*([0-9][0-9\.\,]*)', text, re.I):
        raw=m.group(0).replace('\n',' ')
        if raw not in vals: vals.append(raw)
    return vals[:12]


def _date_values(text):
    pats=[r'\b\d{1,2}[-/]\d{1,2}[-/]\d{4}\b',r'\b\d{1,2}\s+(?:Januari|Februari|Maret|April|Mei|Juni|Juli|Agustus|September|Oktober|November|Desember)\s+\d{4}\b',r'\b(?:Januari|Februari|Maret|April|Mei|Juni|Juli|Agustus|September|Oktober|November|Desember)\s+\d{4}\b']
    vals=[]
    for pat in pats:
        vals += re.findall(pat,text,re.I)
    return _unique(vals,12)


def _provision_refs(text):
    refs=[]
    patterns=[
      r'Pasal\s+\d+[A-Za-z]?(?:\s+ayat\s*\([^\)]*\))?(?:\s+huruf\s+[a-z])?',
      r'Undang-Undang(?:\s+Republik Indonesia)?\s+Nomor\s+\d+\s+Tahun\s+\d{4}[^\n\.;]{0,120}',
      r'UU\s+No\.?\s*\d+\s+Tahun\s+\d{4}[^\n\.;]{0,100}',
      r'Peraturan\s+Otoritas\s+Jasa\s+Keuangan[^\n\.;]{0,120}',
      r'PERMA\s+Nomor\s+\d+\s+Tahun\s+\d{4}',r'SEMA\s+Nomor\s+\d+\s+Tahun\s+\d{4}'
    ]
    for pat in patterns:
        refs += re.findall(pat,text,re.I)
    return _unique([re.sub(r'\s+',' ',x).strip() for x in refs],20)


def _select_sentences(text, markers, limit=8, exclude=None):
    exclude=exclude or []
    out=[]
    for s in sentences(text):
        low=s.lower()
        if any(k in low for k in markers) and not any(k in low for k in exclude):
            out.append(s[:900])
    return _unique(out,limit)


def _deep_case_profile(text):
    low=text.lower(); ss=sentences(text)
    domain_contract=classify_case(text)
    active=set(domain_contract.get('domain_contract') or [])
    primary=domain_contract.get('primary_domain')
    is_corruption=primary=='corruption'
    is_criminal=is_corruption or 'criminal' in active or str(domain_contract.get('posture','')).startswith('PIDANA')
    is_credit='financial_services' in active
    is_land='land_property' in active
    is_civil_procedure='civil_procedure' in active
    is_religious='religious_court' in active
    is_employment='employment' in active

    incr_markers=['mengakui','benar','saya setujui','saya menyetujui','pemutus kredit','menyimpang','menyimpangi','tidak membuat justifikasi','tanpa survey','tanpa survei','tidak memeriksa','tidak sesuai ketentuan','bertentangan dengan ketentuan','memiliki kewenangan untuk menolak','saya acc','saya memutuskan']
    mitig_markers=['tidak pernah memerintah','tidak pernah melarang','percaya kepada','berdasarkan pertimbangan','komite kredit','kredit sebelumnya','riwayat pembayaran lancar','bpkb','agunan','fidusia','dokumen fiat','dokumen analisa','tidak menerima','tidak mengetahui','tidak ikut langsung','bagian marketing','kewenangan bagian kredit']
    incr=_select_sentences(text,incr_markers,10)
    mitig=_select_sentences(text,mitig_markers,10)

    facts=[]
    for s in ss:
        l=s.lower()
        if any(k in l for k in ['rp.','rp ','tanggal','tahun 20','direktur','debitur','perjanjian kredit','pemutus kredit','agunan','survei','survey']): facts.append(s[:700])
    facts=_unique(facts,12) or ss[:8]

    issues=[]
    if is_corruption:
        issues += [
          'Apakah penyimpangan prosedur/tata kelola perkreditan dapat dibuktikan sebagai perbuatan koruptif, bukan semata keputusan bisnis/perbankan yang buruk?',
          'Apakah terdapat kesengajaan atau tujuan menguntungkan diri sendiri, orang lain, atau korporasi yang dapat diatribusikan secara personal kepada pihak yang dianalisis?',
          'Apakah kerugian keuangan negara/daerah telah dihitung sebagai kerugian nyata (actual loss), dengan memperhitungkan pembayaran, outstanding, agunan, dan hasil pemulihan?',
          'Apakah terdapat hubungan kausal langsung antara keputusan pihak yang dianalisis dengan kerugian, atau terdapat kontribusi pihak lain/intervening acts?',
          'Apakah pembagian tanggung jawab antara pemutus akhir, komite kredit, pejabat pemasaran, kepatuhan, SPI, dewan pengawas, dan debitur telah dipetakan secara individual?'
        ]
    if is_credit:
        issues += [
          'Apakah SOP/PKPB dan aturan internal yang dipakai penyidik benar-benar masih berlaku pada tempus perbuatan?',
          'Apakah deviasi dari batas agunan, survei, analisa 5C, provisi/administrasi, atau prosedur internal mempunyai mekanisme pengecualian/justifikasi?',
          'Siapa yang secara fungsional wajib melakukan survei, verifikasi pendapatan, SLIK, penilaian agunan, dan monitoring pascapencairan?'
        ]
    if is_criminal:
        issues += ['Apakah hak tersangka, pendampingan penasihat hukum, pemberitahuan sangkaan, dan proses pemeriksaan telah dipenuhi sesuai hukum acara yang berlaku?']
    if is_civil_procedure:
        issues += ['Apakah gugatan memenuhi syarat formil, para pihak lengkap, posita-petitum konsisten, dan pengadilan yang dipilih memiliki kompetensi absolut serta relatif?']
    if is_land:
        issues += ['Bagaimana riwayat alas hak, data fisik/yuridis, proses pendaftaran, status sertipikat, dan hubungan antara bukti keperdataan dengan administrasi pertanahan?']
    if is_religious:
        issues += ['Apakah materi sengketa termasuk kompetensi absolut Peradilan Agama berdasarkan status para pihak dan objek sengketa yang sebenarnya?']
    if is_employment:
        issues += ['Apakah hubungan kerja, jenis perjanjian, alasan tindakan perusahaan, prosedur bipartit/PHI, dan hak normatif pekerja telah dipetakan berdasarkan ketentuan yang berlaku?']

    years=sorted(set(int(x) for x in re.findall(r'\b(20\d{2})\b',text)))
    temporal=False
    if (is_criminal or is_corruption or is_credit) and years and min(years)<=2022 and any(y>=2023 for y in years):
        temporal=True
        issues.append('Isu tempus/ketentuan peralihan: perbuatan disebut terjadi sebelum berlakunya sebagian norma yang dirujuk dalam proses 2026; perlu comparative-law analysis untuk menentukan norma yang sah dan paling menguntungkan sesuai ketentuan peralihan.')

    elements=[]
    if is_corruption:
        elements=[
          {'element':'Perbuatan melawan hukum/penyalahgunaan kewenangan','prosecution_support':'Ada indikasi deviasi prosedur dan kewenangan pemutus akhir jika fakta/dokumen membuktikannya.','defense_focus':'Pisahkan pelanggaran SOP/administratif dari kualitas perbuatan pidana; uji status berlaku SOP dan adanya collective decision process.','status':'DISPUTED'},
          {'element':'Mens rea / tujuan menguntungkan','prosecution_support':'Dapat diinferensikan dari keputusan yang sadar menyimpang jika dibarengi bukti motif/benefit.','defense_focus':'Cari atau bantah bukti fee, kickback, afiliasi, aliran dana, atau keuntungan personal; motif peningkatan kinerja bank tidak identik dengan motif koruptif.','status':'EVIDENCE NEEDED'},
          {'element':'Kerugian keuangan negara/daerah nyata','prosecution_support':'Kredit bermasalah dapat menjadi basis perhitungan jika auditor membuktikan actual loss.','defense_focus':'Minta LHP/PKKN, outstanding, pembayaran, pencadangan, nilai dan realisasi agunan, serta metode audit; plafon kredit tidak otomatis sama dengan kerugian.','status':'MUST BE PROVEN'},
          {'element':'Kausalitas','prosecution_support':'Jaksa dapat menghubungkan persetujuan kredit dengan pencairan dan kerugian.','defense_focus':'Uji tindakan debitur, kualitas analisa staf, monitoring, kondisi usaha, dan tindakan setelah pencairan sebagai faktor kausal/intervening.','status':'DISPUTED'},
          {'element':'Personal responsibility','prosecution_support':'Pemutus akhir dapat diposisikan sebagai pihak yang memiliki final authority.','defense_focus':'Buat responsibility matrix per fungsi; tandatangan final tidak otomatis berarti seluruh fakta teknis diketahui atau dibuat oleh pemutus akhir.','status':'MUST BE INDIVIDUALIZED'}
        ]

    gaps=[]
    if is_corruption:
        gaps += ['Laporan hasil perhitungan kerugian keuangan negara/daerah lengkap beserta metode dan tanggal cut-off','Bukti aliran dana/keuntungan pribadi, afiliasi, gratifikasi, fee, kickback, atau hubungan khusus dengan debitur','Status pembayaran kredit, outstanding principal, bunga, kolektibilitas, CKPN/pencadangan, dan pemulihan','Status dan nilai aktual agunan serta tindakan eksekusi fidusia/penjualan agunan']
    if is_credit:
        gaps += ['PKPB/SOP yang berlaku tepat pada tanggal keputusan kredit','Seluruh SK/SE Direksi dan perubahan/revokasinya','POJK/aturan eksternal yang berlaku pada tempus kredit','SLIK, berita acara survei, analisa 5C, analisa pendapatan, appraisal agunan, opinion kepatuhan, notulen/lembar komite, fiat, dan bukti pencairan','Laporan SPI/internal audit dan tindak lanjut atas temuan']
    if is_criminal:
        gaps += ['BAP saksi lain, ahli, dokumen penyitaan, audit, dan alat bukti yang dipakai untuk menghubungkan masing-masing unsur pasal']
    if is_civil_procedure:
        gaps += ['Salinan lengkap gugatan, relaas/panggilan, surat kuasa khusus, identitas dan kapasitas seluruh pihak, serta dokumen yang menentukan kompetensi absolut/relatif']
    if is_land:
        gaps += ['Sertipikat/buku tanah, surat ukur, warkah/alasan hak, riwayat peralihan, SKPT terbaru, dan dokumen ukur/pengumuman data fisik-yuridis yang relevan']
    if is_religious:
        gaps += ['Dokumen status para pihak dan objek hubungan hukum yang menentukan kewenangan Peradilan Agama']
    if is_employment:
        gaps += ['Perjanjian kerja, peraturan perusahaan/PKB, slip upah, surat peringatan/pemberitahuan, risalah bipartit, dan dokumen dasar PHK/tindakan perusahaan']

    strategy=[]
    if is_corruption:
        strategy += [
          'Jangan membangun pembelaan utama dengan menyangkal seluruh penyimpangan jika dokumen/BAP justru mengakuinya; akui area tata kelola yang memang terbukti lalu pisahkan dari unsur pidana korupsi.',
          'Susun element-by-element defense matrix: fakta → bukti → unsur pasal → bukti jaksa → bantahan → bukti pembelaan → status pembuktian.',
          'Prioritaskan audit actual loss dan causal chain sebelum menyimpulkan terpenuhinya unsur kerugian negara.',
          'Audit hubungan pihak yang dianalisis dengan debitur dan lakukan financial tracing untuk menguji ada/tidaknya keuntungan personal.',
          'Ajukan saksi a de charge dan ahli yang relevan (perbankan/BPR, audit kerugian, tata kelola, dan bila perlu hukum pidana) berdasarkan gap bukti yang ditemukan.'
        ]
    if is_credit:
        strategy += ['Rekonstruksi credit approval chain dari permohonan sampai monitoring pascapencairan dan petakan siapa mengetahui apa, kapan, serta berdasarkan dokumen apa.']
    if temporal:
        strategy += ['Buat matriks tempus delicti vs norma lama/norma baru/ketentuan peralihan dan bandingkan unsur serta ancaman pidana untuk memastikan penerapan hukum yang sah dan lebih menguntungkan bila relevan.']
    if is_civil_procedure:
        strategy += ['Susun matriks eksepsi dan jawaban per posita; pisahkan isu kompetensi, pihak, kejelasan gugatan, pembuktian pokok perkara, dan rekonvensi bila relevan.']
    if is_land:
        strategy += ['Rekonstruksi chain of title/warkah dan cocokkan dengan buku tanah, surat ukur, SKPT, serta tindakan administrasi BPN sebelum menyimpulkan kekuatan hak.']
    if is_religious:
        strategy += ['Uji kompetensi absolut berdasarkan substansi hubungan hukum, bukan sekadar label gugatan, dan verifikasi Pasal 49 UU Peradilan Agama beserta perkembangan putusan yang relevan.']
    if is_employment:
        strategy += ['Petakan hubungan kerja, proses bipartit dan hak normatif; jangan mengandalkan klausul kontrak untuk meniadakan hak yang bersifat wajib.']

    # Fail-closed base synthesis.  Domain-specific narrative is applied later by
    # services.case_reasoning_guard only when the source itself supports the
    # required nexus/gates.  This prevents a classifier hit from becoming a
    # hard-coded legal conclusion.
    conclusion='Dokumen memerlukan pengujian fakta, bukti, identitas norma, status berlaku, tempus, keterkaitan perkara, dan atribusi pertanggungjawaban sebelum kesimpulan hukum spesifik dapat dibuat.'

    return {
      'facts':facts,'incriminating_facts':incr,'mitigating_facts':mitig,'legal_issues':_unique(issues,14),
      'element_matrix':elements,'evidentiary_gaps':_unique(gaps,18),'strategy':_unique(strategy,14),
      'conclusion':conclusion,'temporal_issue':temporal,'money_values':_money_values(text),'date_values':_date_values(text),'provision_refs':_provision_refs(text),
      'is_criminal':is_criminal,'is_corruption':is_corruption,'is_credit':is_credit,'domain_contract':domain_contract
    }



def _case_domain_allowed_registry_labels(domain_contract):
    active=set((domain_contract or {}).get('domain_contract') or [])
    mapping={
      'corruption':{'Tindak Pidana Korupsi','Pidana Materiil'},
      'financial_services':{'Perbankan & Jasa Keuangan'},
      'criminal':{'Pidana Materiil','Acara Pidana'},
      'regional_government':{'BUMD & Pemerintahan Daerah','Administrasi Pemerintahan'},
      'civil_contract':{'Perdata & Perikatan'},
      'civil_procedure':{'Hukum Acara Perdata'},
      'land_property':{'Pertanahan & Properti'},
      'religious_court':{'Keluarga & Peradilan Agama'},
      'employment':{'Ketenagakerjaan'},
      'corporate':{'Perusahaan & Komersial'},
      'consumer':{'Konsumen'},
      'data_privacy':{'Teknologi, ITE & Data'},
      'bankruptcy':{'Perdata & Perikatan'},
      'arbitration':{'Perdata & Perikatan'},
    }
    out=set()
    for did in active: out |= mapping.get(did,set())
    return out


def _qualified_document_refs(refs):
    """Only parent-qualified instruments enter Applicable Law; orphan Pasal stays evidence text."""
    out=[]
    for ref in refs or []:
        t=re.sub(r'\s+',' ',str(ref or '')).strip()
        low=t.lower()
        if not t: continue
        if re.fullmatch(r'pasal\s+\d+[a-z]?(?:\s+ayat\s*\([^)]*\))?(?:\s+huruf\s+[a-z])?',t,re.I):
            continue
        if any(k in low for k in ('undang-undang','uu no','peraturan otoritas','pojk','seojk','perma','sema','kuhp','kuhap','kuhperdata','uupa')):
            out.append(t)
    return _unique(out,12)

def _case_analysis_payload(text,title='Case Analysis',input_type='narrative',filename='',official_verification=None):
    profile=_deep_case_profile(text)
    domain_contract=profile.get('domain_contract') or classify_case(text)
    allowed_labels=_case_domain_allowed_registry_labels(domain_contract)
    matched=_case_issue_spot(text)
    applicable=[]
    # Domain contract is authoritative: keyword hits outside active domains are ignored.
    for _,item,hits in matched:
        if item.get('domain') not in allowed_labels:
            continue
        for src in item['sources']:
            if not any(x.get('source')==src for x in applicable):
                applicable.append({'domain':item['domain'],'source':src,'status':'PERLU VERIFIKASI PASAL, PERUBAHAN, STATUS BERLAKU, DAN RELEVANSI TERHADAP FAKTA'})
        if len(applicable)>=12: break
    for ref in _qualified_document_refs(profile.get('provision_refs')):
        if not any(ref.lower() in x['source'].lower() for x in applicable):
            applicable.insert(0,{'domain':'Norma disebut dalam dokumen','source':ref,'status':'WAJIB DICEK KE SUMBER RESMI + TEMPUS/STATUS BERLAKU'})
    if not applicable:
        applicable=[{'domain':'Umum','source':'Peraturan dan putusan resmi yang relevan setelah isu diklasifikasikan','status':'PERLU IDENTIFIKASI LANJUT'}]

    official_verification=official_verification or {}
    verification_status=official_verification.get('status','NOT_RUN')
    legal_status=('AKSES SUMBER RESMI OTORITATIF LENGKAP — ANALISIS SUBSTANTIF TETAP PROFESSIONAL REVIEW PENDING' if verification_status=='FULL_OFFICIAL_SOURCE_ACCESS' else
                  'AKSES SUMBER RESMI OTORITATIF PARSIAL — DASAR HUKUM BELUM BOLEH DINYATAKAN LENGKAP' if verification_status=='PARTIAL_OFFICIAL_SOURCE_ACCESS' else
                  'SUMBER RESMI OTORITATIF BELUM TERVERIFIKASI')

    evidence=['Kronologi bertanggal dan identitas/kapasitas seluruh pihak','Dokumen asli yang membentuk hubungan hukum dan kewenangan','Bukti elektronik/keuangan beserta metadata bila relevan']
    evidence += profile['evidentiary_gaps']
    risks=['Analisis berbasis dokumen yang diunggah; fakta di luar dokumen dapat mengubah kesimpulan.','Pengakuan dalam BAP/dokumen harus dibaca bersama dokumen primer, BAP pihak lain, audit, dan alat bukti lain; satu dokumen tidak boleh diperlakukan sebagai keseluruhan berkas perkara.','Status berlaku, perubahan, pencabutan, ketentuan peralihan, lex specialis/lex posterior, serta putusan pengujian wajib diverifikasi pada sumber resmi.']

    pros=[]
    if profile['mitigating_facts']:
        pros += ['Fakta yang berpotensi mendukung pembelaan/mitigasi harus diuji terhadap dokumen primer: '+x for x in profile['mitigating_facts'][:5]]
    cons=[]
    if profile['incriminating_facts']:
        cons += ['Fakta yang berpotensi digunakan penuntut/lawan: '+x for x in profile['incriminating_facts'][:5]]
    if not pros: pros=['Belum teridentifikasi fakta meringankan yang kuat hanya dari teks; diperlukan pembacaan bukti lain.']
    if not cons: cons=['Belum teridentifikasi fakta memberatkan yang kuat hanya dari teks; diperlukan pembacaan bukti lain.']

    # Keep deterministic payload neutral.  A case-specific guarded synthesis is
    # applied after deterministic/AI processing in routes/case_analysis.py.
    analysis=profile['conclusion']

    # v1.3.1.1: restore Regulatory Corpus + Norm Conflict integration from v1.3.0.
    regulatory_matches=retrieve_for_case(text, profile.get('provision_refs',[]), limit=10, allowed_domains=domain_contract.get('domain_contract') or [])
    regulatory_intelligence=intelligence_for_case(regulatory_matches)
    norm_conflicts=analyze_conflicts(profile.get('provision_refs',[]), text, regulatory_matches)

    return dict(
      title=title,input_type=input_type,filename=filename,source_text=text,
      case_posture=domain_contract.get('posture'),domain_classification=domain_contract,domain_contract=domain_contract,
      facts=profile['facts'],incriminating_facts=profile['incriminating_facts'],mitigating_facts=profile['mitigating_facts'],
      legal_issues=profile['legal_issues'],applicable_law=applicable,legal_analysis=analysis,
      element_matrix=profile['element_matrix'],arguments_for=pros,arguments_against=cons,
      evidence_needed=_unique(evidence,24),evidentiary_gaps=profile['evidentiary_gaps'],risks=risks,
      recommendations=profile['strategy'] or ['Susun matriks Fakta → Bukti → Norma → Unsur → Kesimpulan → Risiko.'],
      money_values=profile['money_values'],date_values=profile['date_values'],provision_refs=profile['provision_refs'],
      regulatory_matches=regulatory_matches,
      regulatory_corpus_status={'mode':'LOCAL_CURATED_CORPUS','matches':len(regulatory_matches),'official_verification_required':True,'professional_verification':'PENDING'},
      regulatory_intelligence=regulatory_intelligence,
      norm_conflicts=norm_conflicts,
      official_verification=official_verification,legal_status=legal_status,
      analytical_method='DEEP_CASE_ANALYSIS_V2',
      coverage_note=PRODUCT_LABEL + ' membaca dokumen dengan pendekatan source-grounded dan fail-closed: fakta sumber → isu → domain/nexus → identitas/status norma → tempus → unsur/atribusi → kesimpulan bersyarat. Istilah domain-spesifik hanya boleh muncul bila nexus sumbernya terpenuhi; kesimpulan final tetap Professional Verification: PENDING.'
    )



def _sanitize_executive_text(value):
    """Remove identifiers that do not belong in an executive summary.

    The full source remains available in the source ledger; this only keeps the
    high-level working-paper summary from unnecessarily repeating NIK, phone,
    email, and detailed address data.
    """
    t=re.sub(r'\s+',' ',str(value or '')).strip()
    t=re.sub(r'\bN\.?I\.?K\.?\s*[:\-]?\s*\d{12,18}\b','NIK [REDACTED]',t,flags=re.I)
    t=re.sub(r'\b(?:\+62|62|0)8\d{7,12}\b','[PHONE REDACTED]',t)
    t=re.sub(r'\b[A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,}\b','[EMAIL REDACTED]',t,flags=re.I)
    t=re.sub(r'\b(?:Jl\.?|Jalan)\s+[^.;]{3,100}(?=[.;]|$)','[ADDRESS REDACTED]',t,flags=re.I)
    return t.strip(' ;,-')


def _source_ledger_candidates(source_text, limit=48):
    """Create a broad extractive source ledger for audit/traceability.

    This function intentionally preserves document structure, pleading labels,
    prayers, objections and other source material because ``source_ledger`` is
    the audit trace of what was actually read.  It is *not* the lawyer-facing
    Evidence Map.  The stricter ``material_source_ledger`` / ``evidence_rows``
    projection is responsible for excluding counsel addresses, identity blocks,
    procedural metadata, prayers, legal argument and other non-merits material.

    Keeping these two layers separate prevents a recurring regression where
    cleaning the Evidence Map accidentally deleted source trace needed by other
    workflows/tests.
    """
    text=str(source_text or '').replace('\r','\n')
    text=re.sub(r'(?i)dipindai\s+dengan\s+camscanner',' ',text)
    text=re.sub(r'[ \t]+',' ',text)
    raw=[re.sub(r'\s+',' ',x).strip(' -\t') for x in re.split(r'\n{1,}|(?<=[.!?;:])\s+(?=[A-Z0-9IV])',text)]

    anchors=(
        'perbuatan','kejadian','terjadi','terdakwa','penggugat','tergugat','saksi','debitur','kreditur',
        'pemberian kredit','kredit macet','pencairan','pembayaran','plafon','appraisal','agunan','jaminan',
        'fidusia','bpkb','rekening','audit','kerugian','dana','keuntungan','gratifikasi','fee','kickback',
        'direktur','jabatan','kewenangan','keputusan','komite','persetujuan','sop','kepatuhan','pengawas',
        'perjanjian','wanprestasi','melawan hukum','hak','kewajiban','tanah','sertipikat','sertifikat','shm',
        'waris','ahli waris','penguasaan','penyerahan','jual beli','utang','piutang','tidak ditemukan',
        'dialihkan','dijual','disimpan','tersimpan','digantikan','diangkat','diperpanjang','eksepsi',
        'kompetensi','kewenangan mengadili','error in persona','obscuur','plurium','dakwaan','rekonvensi',
        'konvensi','permohonan','petitum','kesimpulan','replik','duplik','ptsl','skpt'
    )
    structural_re=re.compile(r'(?i)^(?:dalam\s+)?(?:eksepsi|konvensi|rekonvensi|kesimpulan|permohonan|petitum|replik|duplik)(?:\b.*)?$')
    scored=[]
    for idx,chunk in enumerate(raw):
        if not chunk or len(chunk)>900:
            continue
        low=chunk.lower()
        structural=bool(structural_re.match(chunk))
        if len(chunk)<12 and not structural:
            continue
        score=sum(2 for a in anchors if a in low)
        if re.search(r'(?i)\b(?:Rp\.?\s*)?[0-9][0-9.]{3,}',chunk): score+=1
        if re.search(r'(?i)\b(?:tidak|belum|melebihi|memberikan|menerima|menjual|mengalihkan|menyimpan|menjabat|diangkat|digantikan)\b',chunk): score+=2
        if structural: score=max(score,2)
        if score<=0:
            continue
        label='DOCUMENT' if structural else ('ALLEGATION' if any(x in low for x in ('mendalilkan','menyatakan','dakwaan','didakwa','menuntut')) else 'SOURCE FACT')
        scored.append((score,idx,chunk,label))

    strongest=sorted(scored,key=lambda x:(-x[0],x[1]))[:max(limit*3,limit)]
    strongest=sorted(strongest,key=lambda x:x[1])
    out=[]; seen=set()
    for score,idx,chunk,label in strongest:
        key=re.sub(r'\W+',' ',chunk.lower()).strip()[:180]
        if not key or key in seen:
            continue
        seen.add(key)
        out.append({'label':label,'statement':chunk[:780],'evidence':chunk[:780],'segment':None,'source':'DETERMINISTIC_EXTRACTIVE'})
        if len(out)>=limit:
            break
    return out

def _compact_executive_summary(result):
    source=re.sub(r'\s+',' ',str(result.get('source_text') or '')).strip()

    def fact_quality(value):
        t=re.sub(r'\s+',' ',str(value or '')).strip(' -;:')
        if not t or len(t) < 24 or len(t) > 360:
            return False
        low=t.lower()
        noise=(
            'demi keadilan dan kebenaran','paraf:','mengertikah tersangka',
            'anak dani','anak dan (alm)','jo.pasal','jo. pasal','----------',
            'nama lengkap drs.','tindak pidana khusus','nik:','n.i.k','no.handphone',
            'phone ','email :','alamat :','tempat tinggal','lahir di'
        )
        if any(x in low for x in noise):
            return False
        # Reject OCR fragments dominated by statute chaining or punctuation.
        if t.count('Pasal ') >= 2 or sum(t.count(ch) for ch in '_-|') >= 5:
            return False
        return True

    facts=[_sanitize_executive_text(x) for x in _unique(result.get('facts') or [], 14) if fact_quality(x)]
    semantic_facts=[_sanitize_executive_text(x) for x in executive_fact_candidates(result.get('source_ledger') or [],5) if fact_quality(x)]
    if semantic_facts:
        facts=_unique(semantic_facts+facts,10)

    # Case capsule uses the same global tempus guard as Legal Construction.
    # A year occurring in a case number, SK number, statute citation or object
    # model must never become the year of the legal event merely by proximity.
    case_bits=[]
    if source and re.search(r'Perumda\s+BPR|PD\.?\s*BPR', source, re.I):
        if re.search(r'Tindak\s+Pidana\s+Korupsi|Tipikor', source, re.I):
            bpr_label = ('Perumda BPR' if re.search(r'Perumda\s+BPR', source, re.I)
                         else ('PD BPR' if re.search(r'PD\.?\s*BPR', source, re.I) else 'BPR'))
            case_bits.append(f'dokumen berkaitan dengan perkara tindak pidana korupsi dan fasilitas kredit pada {bpr_label}')
        guard=result.get('reasoning_guard') or {}
        tempus=guard.get('tempus') or {}
        years=tempus.get('event_year_candidates') or []
        if len(years)==1 and tempus.get('status') not in ('TEMPUS_INSUFFICIENT','TEMPUS_REQUIRES_EVENT_DATE'):
            case_bits.append('kandidat tahun perbuatan yang perlu diverifikasi: '+str(years[0]))
        else:
            case_bits.append('tahun/tanggal perbuatan kredit belum dapat ditetapkan secara aman dari dokumen ini')
        if re.search(r'Direktur\s+Utama', source, re.I):
            case_bits.append('dokumen menyebut pihak yang dianalisis pernah menjabat Direktur Utama; periode dan atribusi tindakannya tetap harus dipetakan')
    if case_bits:
        facts=['; '.join(case_bits)+'.'] + facts
    else:
        posture=str(result.get('case_posture') or (result.get('domain_classification') or {}).get('posture') or '')
        civil_bits=[]
        if posture=='PERDATA_LITIGASI' or re.search(r'penggugat|tergugat|gugatan',source,re.I):
            if re.search(r'jawaban\s+(?:pertama\s+)?(?:para\s+)?tergugat|jawaban.*gugatan',source,re.I):
                civil_bits.append('dokumen merupakan jawaban/pleading dalam sengketa perdata')
            elif re.search(r'gugatan',source,re.I):
                civil_bits.append('dokumen berkaitan dengan gugatan perdata')
            if re.search(r'obscuur\s+libel',source,re.I): civil_bits.append('terdapat eksepsi obscuur libel')
            if re.search(r'plurium\s+litis\s+consortium',source,re.I): civil_bits.append('terdapat eksepsi kurang pihak/plurium litis consortium')
            if re.search(r'kompetensi\s+absolut|kewenangan\s+absolut|pengadilan\s+agama',source,re.I): civil_bits.append('kompetensi absolut forum dipersoalkan')
            if re.search(r'sertipikat|sertifikat|shm|skpt|surat ukur|kantor pertanahan',source,re.I): civil_bits.append('objek/fakta pertanahan dan riwayat pendaftaran menjadi bagian material sengketa')
        if civil_bits:
            facts=['; '.join(civil_bits)+'.'] + [f for f in facts if not re.search(r'phone|email|alamat\s*:',f,re.I)]

    issues=_unique(result.get('legal_issues') or [], 5)
    gaps=_unique(result.get('evidentiary_gaps') or result.get('evidence_needed') or [], 5)
    out=[]
    if facts:
        out.append('Fakta kunci: ' + _sanitize_executive_text('; '.join(str(x) for x in facts[:3])))
    elif source:
        out.append('Fakta kunci: dokumen sumber telah terbaca, tetapi fakta material belum cukup bersih untuk diringkas tanpa verifikasi terhadap teks sumber.')
    if issues:
        out.append('Isu utama: ' + '; '.join(str(x) for x in issues[:3]))
    analysis=str(result.get('legal_analysis') or '').strip()
    if analysis:
        out.append('Posisi analitis awal: ' + analysis[:900])
    if gaps:
        out.append('Gap pembuktian kritis: ' + '; '.join(str(x) for x in gaps[:4]))
    return out[:4]


def _clean_applicable_law(result):
    raw=result.get('applicable_law') or []
    snapshot=result.get('case_regulatory_snapshot') or {}
    domain_ids={str(d.get('id')) for d in (snapshot.get('domains') or [])[:4] if isinstance(d,dict)}
    domain_to_labels={
        'financial_services': {'Perbankan & Jasa Keuangan','Perusahaan & Komersial'},
        'corruption': {'Pidana Materiil'},
        'criminal': {'Pidana Materiil','Acara Pidana'},
        'civil_contract': {'Perdata & Perikatan'},
        'corporate': {'Perusahaan & Komersial'},
        'employment': {'Ketenagakerjaan'},
        'land_property': {'Pertanahan & Properti'},
        'religious_court': {'Keluarga & Peradilan Agama'},
        'civil_procedure': {'Hukum Acara Perdata'},
        'constitutional': {'Konstitusi & Hak Dasar'},
        'regional_government': {'Administrasi Pemerintahan','Perusahaan & Komersial'},
        'administrative': {'Administrasi Pemerintahan','Hukum Acara TUN'},
        'public_information': {'Keterbukaan Informasi Publik','Administrasi Pemerintahan'},
        'investment': {'Penanaman Modal','Perusahaan & Komersial'},
    }
    allowed_labels={'Norma disebut dalam dokumen'}
    for did in domain_ids:
        allowed_labels |= domain_to_labels.get(did,set())
    out=[]; seen=set()
    orphan=re.compile(r'^Pasal\s+\d+[A-Za-z]?(?:\s+ayat\s*\([^)]*\))?(?:\s+huruf\s+[a-z])?$',re.I)
    noisy=re.compile(r'^Pasal\s+\d+[A-Za-z]$',re.I)
    for item in raw:
        if not isinstance(item,dict): continue
        domain=str(item.get('domain') or '')
        source=str(item.get('source') or '').strip()
        if allowed_labels and domain not in allowed_labels: continue
        if not source or orphan.fullmatch(source) or noisy.fullmatch(source): continue
        key=(domain.lower(),source.lower())
        if key in seen: continue
        seen.add(key); out.append(item)
    return out[:14]

def _ensure_evidence_to_action(result):
    """Normalize v1.3.1 Evidence-to-Action fields for AI and deterministic fallback."""
    result = result or {}
    # Build a compact source ledger from AI segment extraction when available.
    ledger=[]
    label_map=(
        ('facts','SOURCE FACT'),('admissions','ADMISSION'),('denials_or_limits','DENIAL / LIMITATION'),
        ('allegations','ALLEGATION')
    )
    for seg in result.get('ai_evidence_map') or []:
        segno=seg.get('_segment')
        explicit=seg.get('source_items') or []
        for item in explicit:
            if not isinstance(item,dict): continue
            ledger.append({
                'label':item.get('label') or 'SOURCE FACT',
                'statement':item.get('statement') or '',
                'evidence':item.get('evidence') or '',
                'segment':item.get('segment') or segno
            })
        if not explicit:
            for key,label in label_map:
                for item in (seg.get(key) or []):
                    if isinstance(item,dict):
                        ledger.append({'label':label,'statement':item.get('statement',''),'evidence':item.get('evidence',''),'segment':segno})
            for doc in seg.get('documents') or []:
                ledger.append({'label':'DOCUMENT','statement':str(doc),'evidence':'','segment':segno})
    if not ledger:
        # Prefer a deeper extractive ledger from the actual source.  Facts from
        # the deterministic analyzer remain useful, but OCR documents commonly
        # contain far more material than that small list captures.
        ledger.extend(_source_ledger_candidates(result.get('source_text') or '', limit=36))
        existing={re.sub(r'\W+',' ',str(x.get('statement') or '').lower()).strip()[:180] for x in ledger}
        for fact in result.get('facts') or []:
            statement=re.sub(r'\s+',' ',str(fact)).strip()[:650]
            key=re.sub(r'\W+',' ',statement.lower()).strip()[:180]
            if statement and key not in existing:
                ledger.append({'label':'SOURCE FACT','statement':statement,'evidence':statement,'segment':None,'source':'DETERMINISTIC_FACT'})
                existing.add(key)
    result['source_ledger']=ledger[:160]
    # Keep the raw ledger for internal traceability, but create a separate
    # user-facing projection that contains only case-material source items.
    # Identity/contact/procedural metadata must never masquerade as evidence.
    from services.case_consistency_guard import material_source_ledger
    result['material_source_ledger']=material_source_ledger(result['source_ledger'], limit=80)
    from services.case_consistency_guard import classify_source_item
    hygiene_counts={}
    for _item in result['source_ledger']:
        if not isinstance(_item,dict):
            continue
        _cls=classify_source_item(_item).get('classification') or 'UNKNOWN'
        hygiene_counts[_cls]=hygiene_counts.get(_cls,0)+1
    result['evidence_hygiene']={
        'policy':'MATERIAL_EVIDENCE_ONLY',
        'raw_source_items':len(result['source_ledger']),
        'user_visible_material_items':len(result['material_source_ledger']),
        'classification_counts':hygiene_counts,
        'excluded_from_evidence_map':[
            'DOCUMENT_METADATA','PARTY_IDENTITY','PROCEDURAL_METADATA','PETITUM_OR_PRAYER',
            'LEGAL_REFERENCE','LEGAL_ARGUMENT','PLEADING_ASSERTION','NON_MATERIAL_FRAGMENT','DOCUMENT_STRUCTURE'
        ],
    }
    # v1.3.11: extractive evidence grouping, traceable to source_ledger indexes.
    from services.legal_ocr_postprocess import group_source_ledger
    result['evidence_groups']=group_source_ledger(result['source_ledger'])

    gaps=_unique(result.get('evidentiary_gaps') or [],18)
    established=_unique(result.get('facts') or [],12)
    if result.get('incriminating_facts'):
        established += _unique(result.get('incriminating_facts') or [],5)
    not_established=[]
    gap_text=' '.join(gaps).lower()
    for label,keys in [
        ('Mens rea / pengetahuan dan kehendak belum terbukti secara memadai',('mens rea','niat','pengetahuan','kehendak')),
        ('Kerugian aktual belum terverifikasi secara memadai',('kerugian','actual loss','audit')),
        ('Causal nexus antara tindakan dan akibat belum terpetakan secara memadai',('kausal','causal','sebab')),
        ('Tanggung jawab personal belum dapat dipisahkan dari fungsi organisasi',('tanggung jawab','personal','peran','kewenangan'))
    ]:
        if any(k in gap_text for k in keys): not_established.append(label)
    if not not_established:
        not_established=_unique(gaps,8) or ['Tidak ada kesimpulan unsur yang boleh dianggap final hanya dari satu dokumen sumber.']
    must_verify=_unique((result.get('evidence_needed') or []) + [
        'Verifikasi bunyi, status berlaku, perubahan, dan tempus seluruh norma pada sumber resmi.',
        'Cocokkan setiap inference LexiCore dengan bagian dokumen sumber dan alat bukti primer.'
    ],12)
    ds=result.get('decision_summary') if isinstance(result.get('decision_summary'),dict) else {}
    result['executive_summary']=_compact_executive_summary(result)
    result['applicable_law']=_clean_applicable_law(result)
    result['decision_summary']={
        'already_established':_unique(ds.get('already_established') or established,14),
        'not_yet_established':_unique(ds.get('not_yet_established') or not_established,14),
        'must_be_verified':_unique(ds.get('must_be_verified') or must_verify,14)
    }

    action=result.get('action_plan') if isinstance(result.get('action_plan'),list) else []
    if not action:
        # Deterministic priorities: evidentiary gaps first, then evidence collection, then legal verification/strategy.
        for i,g in enumerate(gaps[:5]):
            action.append({'priority':'P1' if i<3 else 'P2','issue':g,'current_status':'BELUM TERBUKTI / BELUM TERVERIFIKASI','action':'Dapatkan dan uji bukti primer yang secara langsung menjawab gap ini.','why_it_matters':'Kesimpulan hukum tidak boleh dinaikkan menjadi established fact sebelum gap ini ditutup.','source_segments':[]})
        for ev in (result.get('evidence_needed') or [])[:4]:
            action.append({'priority':'P2','issue':'Kelengkapan bukti','current_status':'PERLU DIKUMPULKAN','action':ev,'why_it_matters':'Diperlukan untuk menguji fakta, unsur, kausalitas, atau tanggung jawab personal.','source_segments':[]})
        action.append({'priority':'P3','issue':'Verifikasi hukum positif','current_status':result.get('legal_status','BELUM TERVERIFIKASI'),'action':'Verifikasi pasal, perubahan/pencabutan, ketentuan peralihan, dan putusan relevan pada sumber resmi.','why_it_matters':'Mencegah penggunaan norma yang salah tempus atau sudah berubah.','source_segments':[]})
    # Normalize priorities and cap.
    normalized=[]; seen_actions=set(); seen_issues=set()
    for a in action[:24]:
        if not isinstance(a,dict): continue
        pr=str(a.get('priority') or 'P2').upper(); pr=pr if pr in ('P1','P2','P3') else 'P2'
        issue=a.get('issue') or 'Isu'; act=a.get('action') or ''
        if str(issue).strip().lower()=='kelengkapan bukti' and act:
            issue=act
        issue_key=re.sub(r'[^a-z0-9]+',' ',str(issue).lower()).strip()
        # One actionable row per substantive issue. This prevents the same
        # evidentiary gap from reappearing at a lower priority with a generic
        # collection instruction.
        if issue_key and issue_key in seen_issues:
            continue
        sig=(issue_key, re.sub(r'[^a-z0-9]+',' ',str(act).lower()).strip())
        if sig in seen_actions: continue
        seen_actions.add(sig)
        if issue_key: seen_issues.add(issue_key)
        normalized.append({'priority':pr,'issue':issue,'current_status':a.get('current_status') or 'PERLU DIUJI','action':act, 'why_it_matters':a.get('why_it_matters') or '', 'source_segments':a.get('source_segments') or []})
    result['action_plan']=normalized[:14]
    result['professional_verification']='PENDING'
    result['evidence_to_action_version']='1.3.1'
    return result



def _readiness_tokens(text):
    """Conservative legal-workflow matching tokens; not a merits predictor."""
    stop={
      'yang','dan','atau','untuk','dari','dalam','pada','dengan','sebagai','terhadap','oleh','agar','ini','itu','belum','perlu',
      'wajib','harus','dapat','bukti','dokumen','data','hasil','secara','lebih','serta','status','verifikasi','resmi','sumber'
    }
    toks=re.findall(r'[A-Za-zÀ-ÿ0-9]{4,}', str(text or '').lower())
    return {t for t in toks if t not in stop}


def _case_readiness_profile(result):
    """Measure preparation completeness, never probability of winning a case.

    The score is deliberately evidence-led and reproducible.  It rises when
    requested evidence becomes traceable in the source ledger and when the
    legal-analysis/action layers are actually populated.  Missing evidence is
    converted into projected score gains so Action Plan can explain the value
    of obtaining a specific document or proof item.
    """
    ledger=result.get('source_ledger') or []
    # Evidence readiness is satisfied only by actual document/evidence
    # candidates, not by contact metadata, party identity or a pleading that
    # merely says a document exists.  Assertions remain traceable but do not
    # close evidentiary gaps by themselves.
    readiness_items=[]
    for v in ledger:
        if not isinstance(v,dict): continue
        meta=classify_source_item(v)
        if meta.get('classification')=='ACTUAL_EVIDENTIARY_ITEM':
            readiness_items.append(v)
    corpus=' '.join(str(v.get('statement',''))+' '+str(v.get('evidence','')) for v in readiness_items)
    corpus_tokens=_readiness_tokens(corpus)
    needs=_unique((result.get('evidence_needed') or []) + (result.get('evidentiary_gaps') or []), 16)
    evidence_items=[]
    fulfilled=0
    for i,need in enumerate(needs,1):
        nt=_readiness_tokens(need)
        overlap=len(nt & corpus_tokens)
        # Conservative: require at least two meaningful shared tokens, or all
        # tokens where the requirement is exceptionally short.
        threshold=1 if len(nt)<=2 else 2
        ok=bool(nt) and overlap>=threshold
        fulfilled += int(ok)
        evidence_items.append({'id':i,'requirement':need,'status':'TERPENUHI' if ok else 'BELUM TERPENUHI'})
    if needs:
        evidence_pct=round(100*fulfilled/len(needs))
    else:
        evidence_pct=100 if ledger else 0

    legal_fields=[
      result.get('legal_analysis'), result.get('legal_issues'), result.get('element_matrix'),
      result.get('mens_rea_analysis'), result.get('actual_loss_analysis'), result.get('causation_analysis'),
      result.get('personal_responsibility_analysis'), result.get('temporal_law_analysis')
    ]
    legal_done=sum(1 for v in legal_fields if v and (not isinstance(v,(list,dict)) or len(v)>0))
    legal_pct=round(100*legal_done/len(legal_fields))

    actions=result.get('action_plan') or []
    # Action closure is tied to whether an evidence requirement is already
    # satisfied. It does not claim that a legal recommendation itself is done.
    action_total=max(len(actions),1)
    unresolved=sum(1 for e in evidence_items if e['status']!='TERPENUHI')
    action_closed=max(0, len(actions)-min(len(actions),unresolved)) if actions else 0
    action_pct=round(100*action_closed/action_total) if actions else (100 if not needs else evidence_pct)

    weights={'evidence':50,'legal_analysis':30,'action_plan':20}
    overall=round(evidence_pct*.50 + legal_pct*.30 + action_pct*.20)
    overall=max(0,min(100,overall))

    missing=[e for e in evidence_items if e['status']!='TERPENUHI']
    gain_unit=(weights['evidence']/max(len(evidence_items),1)) if evidence_items else 0
    projections=[]
    running=overall
    for e in missing:
        projected=min(100, round(running+gain_unit))
        projections.append({**e,'estimated_gain':round(projected-running,1),'projected_readiness':projected})
    return {
      'metric':'CASE_PREPARATION_COMPLETENESS',
      'label':'Case Readiness / Kelengkapan Persiapan',
      'overall_percentage':overall,
      'components':{
        'evidence_map':{'percentage':evidence_pct,'weight':50,'fulfilled':fulfilled,'required':len(evidence_items)},
        'legal_analysis':{'percentage':legal_pct,'weight':30,'completed_fields':legal_done,'required_fields':len(legal_fields)},
        'action_plan':{'percentage':action_pct,'weight':20,'closed':action_closed,'total':len(actions)}
      },
      'evidence_requirements':evidence_items,
      'missing_evidence':missing,
      'projections':projections,
      'disclaimer':'Persentase ini mengukur kelengkapan persiapan perkara berdasarkan evidence, analisis, dan action plan; bukan probabilitas menang/kalah atau prediksi putusan.'
    }


def _instrument_mentions(text):
    """Return named legal instruments with text offsets for article qualification.

    Verification must never assume that an orphan ``Pasal N`` belongs to a
    particular statute.  We therefore bind an article only to an instrument
    that is actually mentioned close to it in the uploaded document.
    """
    specs=[
      ('UU', r'(?:Undang-Undang(?:\s+Republik Indonesia)?|UU)\s+(?:Nomor|No\.?)\s*\d+[A-Za-z]?\s+Tahun\s+\d{4}'),
      ('POJK', r'(?:Peraturan\s+Otoritas\s+Jasa\s+Keuangan|POJK)\s+(?:Nomor|No\.?)\s*[0-9A-Za-z./-]+(?:\s+Tahun\s+\d{4}|/POJK\.\d{2}/\d{4})?'),
      ('SEOJK', r'(?:Surat\s+Edaran\s+Otoritas\s+Jasa\s+Keuangan|SEOJK)\s+(?:Nomor|No\.?)\s*[0-9A-Za-z./-]+(?:\s+Tahun\s+\d{4}|/SEOJK\.\d{2}/\d{4})?'),
      ('PP', r'(?:Peraturan\s+Pemerintah|PP)\s+(?:Nomor|No\.?)\s*\d+[A-Za-z]?\s+Tahun\s+\d{4}'),
      ('PERPRES', r'(?:Peraturan\s+Presiden|Perpres)\s+(?:Nomor|No\.?)\s*\d+[A-Za-z]?\s+Tahun\s+\d{4}'),
      ('PERMA', r'(?:Peraturan\s+Mahkamah\s+Agung|PERMA)\s+(?:Nomor|No\.?)\s*\d+[A-Za-z]?\s+Tahun\s+\d{4}'),
      ('SEMA', r'(?:Surat\s+Edaran\s+Mahkamah\s+Agung|SEMA)\s+(?:Nomor|No\.?)\s*\d+[A-Za-z]?\s+Tahun\s+\d{4}'),
      ('PERMEN', r'(?:Peraturan\s+Menteri(?:\s+[A-Za-zÀ-ÿ .()/-]+?)?|Permen[A-Za-z]*)\s+(?:Nomor|No\.?)\s*[0-9A-Za-z./-]+\s+Tahun\s+\d{4}'),
    ]
    out=[]
    for kind,pat in specs:
        for m in re.finditer(pat,text,re.I):
            label=re.sub(r'\s+',' ',m.group(0)).strip(' ,.;:-')
            out.append({'kind':kind,'label':label,'start':m.start(),'end':m.end()})
    return sorted(out,key=lambda x:x['start'])


def _qualified_legal_refs(text, max_items=12):
    """Create certainty-preserving verification queries.

    A query such as ``Pasal 603`` is legally ambiguous.  It is emitted only
    when the source document lets us associate the article with a named
    instrument.  Unqualified articles are intentionally withheld from JDIH
    verification rather than guessed.
    """
    instruments=_instrument_mentions(text)
    qualified=[]; unresolved=[]
    article_pat=re.compile(r'Pasal\s+\d+[A-Za-z]?(?:\s+ayat\s*\([^\)]*\))?(?:\s+huruf\s+[a-z])?',re.I)
    for am in article_pat.finditer(text):
        article=re.sub(r'\s+',' ',am.group(0)).strip()
        candidates=[]
        for inst in instruments:
            # Legal citations frequently name the statute immediately before OR
            # after the article. Keep the window deliberately narrow.
            if inst['end'] <= am.start():
                distance=am.start()-inst['end']; direction_penalty=0
            elif inst['start'] >= am.end():
                distance=inst['start']-am.end(); direction_penalty=120
            else:
                distance=0; direction_penalty=0
            if distance <= 900:
                candidates.append((distance+direction_penalty,inst))
        if candidates:
            _,inst=min(candidates,key=lambda x:x[0])
            qualified.append(f"{inst['label']} {article}")
        else:
            unresolved.append(article)
    # Named instruments are still useful as exact instrument-level checks.
    queries=[x['label'] for x in instruments]
    queries=qualified+queries
    return {
      'queries':_unique(queries,max_items),
      'unresolved_articles':_unique(unresolved,12),
      'qualification_policy':'NO_ORPHAN_ARTICLE_GUESSING'
    }

def _case_verification_queries(text,title):
    matched=_case_issue_spot(text)
    qualified=_qualified_legal_refs(text,max_items=12)
    queries=list(qualified['queries'])

    # A case title can improve discovery, but never use it to override an exact
    # regulation + article reference extracted from the source document.
    if title and title.lower()!='case analysis' and not re.fullmatch(r'Pasal\s+\d+[A-Za-z]?',title.strip(),re.I):
        queries.append(title.strip())

    # Add domain discovery only after exact legal references.  Never add bare
    # 'Pasal N' from keyword issue spotting because that can point JDIH search
    # at the wrong statute and create a false impression of legal certainty.
    for _,item,hits in matched[:4]:
        q=' '.join(hits[:3]).strip()
        if q and not re.fullmatch(r'Pasal\s+\d+[A-Za-z]?',q,re.I):
            queries.append((item['domain']+' '+q).strip())
        for src in item['sources'][:2]:
            clean=src.replace('yang relevan','').replace('yang berlaku','').strip()
            if any(k in clean.lower() for k in ('uu no.','undang-undang','kuhp','kuhperdata','perma','sema','pojk','otoritas jasa keuangan')):
                if not re.fullmatch(r'Pasal\s+\d+[A-Za-z]?',clean,re.I):
                    queries.append(clean)

    if not queries:
        # Discovery fallback is descriptive, not an invented article citation.
        queries=[' '.join(re.findall(r'\b[\w-]{4,}\b',text[:1000],flags=re.UNICODE)[:8])]
    return _unique([q for q in queries if len(q)>=4],8)

@app.route('/api/regulations/catalog')
def regulations_catalog():
    data=RegulatoryCorpusManager.all(500)
    intel=normalized_catalog()
    return jsonify(success=True,data=data,count=len(data),mode='LOCAL_DATABASE',storage='SQLITE',schema_version=SCHEMA_VERSION,
                   intelligence_model=intel.get('model'),counts={'regulations':len(data),'articles':sum(len(x.get('articles',[])) for x in data)},
                   official_verification_required=True,professional_verification='PENDING')

@app.route('/api/regulations/search')
def regulations_search():
    q=(request.args.get('q') or '').strip(); mode=(request.args.get('mode') or 'offline').strip().lower()
    limit=request.args.get('limit',12,type=int)
    if mode not in ('offline','online','hybrid'): mode='offline'
    data=search_legal_authorities(q,mode=mode,limit=limit)
    return jsonify(success=True,query=q,mode=mode,data=data,
                   results=data.get('local_results',[]) if mode=='offline' else data.get('local_results',[])+data.get('online_results',[]),
                   total_matched=data.get('counts',{}).get('total',0),official_verification_required=True,professional_verification='PENDING')


@app.route('/api/regulatory-intelligence/catalog')
def regulatory_intelligence_catalog():
    return jsonify(success=True,data=normalized_catalog())

@app.route('/api/regulatory-intelligence/graph')
def regulatory_intelligence_graph():
    q=(request.args.get('q') or '').strip()
    return jsonify(success=True,data=graph_for_query(q,limit=request.args.get('limit',10,type=int)))

@app.route('/api/regulatory-intelligence/timeline')
def regulatory_intelligence_timeline():
    q=(request.args.get('q') or '').strip()
    rid=(request.args.get('regulation_id') or '').strip() or None
    return jsonify(success=True,data=timeline_for_regulation(regulation_id=rid,query=q))

@app.route('/api/regulatory-intelligence/compare',methods=['POST'])
def regulatory_intelligence_compare():
    d=request.get_json(silent=True) or {}
    qa=(d.get('norm_a') or '').strip(); qb=(d.get('norm_b') or '').strip(); context=(d.get('context') or '').strip()
    if not qa or not qb:
        return jsonify(success=False,error='norm_a dan norm_b wajib diisi dengan identitas peraturan yang cukup lengkap'),400
    result=resolve_conflict_from_query(qa,qb,context)
    code=200 if result.get('success') else 422
    return jsonify(success=bool(result.get('success')),data=result),code

@app.route('/api/norm-conflicts',methods=['POST'])
def norm_conflicts_endpoint():
    d=request.get_json(silent=True) or {}
    provisions=d.get('provisions') or []
    context=(d.get('context') or '').strip()
    matches=retrieve_for_case(context,provisions,limit=8)
    result=analyze_conflicts(provisions,context,matches)
    try:
        conflict_count=len((result or {}).get('conflicts_detected',[])) if isinstance(result,dict) else 0
        AuditLogger.log_action('norm_conflict_analysis','norm_conflict',None,f'provisions={len(provisions)} | conflicts={conflict_count}')
    except Exception:
        pass
    return jsonify(success=True,data=result,regulatory_matches=matches)

@app.route('/api/dashboard/metrics')
def dashboard_metrics():
    """Operational metrics for all eight LexiCore modules.

    Counts come from persisted SQLite records where applicable. Regulatory
    Corpus is counted from the curated local corpus. Norm Conflicts uses the
    audit log so no schema migration is required.
    """
    tables={
        'drafts':'drafts',
        'contract_reviews':'contract_analyses',
        'research_notes':'legal_research',
        'risk_assessments':'risk_assessments',
        'case_analyses':'case_analyses',
        'client_documents':'client_communications',
    }
    counts={}
    conn=None
    try:
        conn=get_db_connection()
        for key,table in tables.items():
            try:
                row=conn.execute(f'SELECT COUNT(*) AS n FROM {table}').fetchone()
                counts[key]=int(row['n'] if row is not None else 0)
            except Exception:
                counts[key]=0
        try:
            row=conn.execute("SELECT COUNT(*) AS n FROM audit_log WHERE action='norm_conflict_analysis'").fetchone()
            counts['norm_conflict_analyses']=int(row['n'] if row is not None else 0)
        except Exception:
            counts['norm_conflict_analyses']=0
    finally:
        if conn is not None:
            conn.close()
    counts['regulatory_corpus']=len(RegulatoryCorpusManager.all(1000))
    return jsonify(success=True,version=LEXICORE_VERSION,metrics=counts,metric_semantics={
        'regulatory_corpus':'core regulatory seed entries; case-specific regulations are retrieved dynamically',
        'norm_conflict_analyses':'persisted audit events',
        'client_documents':'client communication/document records'
    })

@app.route('/api/legal-sources')
def legal_sources_registry():
    data=public_source_registry()
    secondary=public_secondary_source_registry()
    return jsonify(success=True,data=data,count=len(data),secondary_research=secondary,secondary_count=len(secondary),
                   note='Secondary research providers such as Hukumonline are discovery/corroboration only and never replace official-source verification.')

@app.route('/api/legal-sources/health')
def legal_sources_health():
    data=source_health(); auth=[x for x in data if x.get('authoritative')]; directory=[x for x in data if not x.get('authoritative')]; return jsonify(success=True,data=data,reachable=sum(1 for x in auth if x.get('reachable')),count=len(auth),directory={'reachable':sum(1 for x in directory if x.get('reachable')),'count':len(directory)},checked_at=datetime.now().isoformat())

@app.route('/api/legal-sources/search')
def legal_sources_search():
    q=(request.args.get('q') or '').strip()
    if len(q)<3: return jsonify(success=False,error='Query minimal 3 karakter'),400
    official=federated_search(q)
    secondary=[]
    if (request.args.get('secondary') or '').strip().lower() in ('1','true','yes','on'):
        sec_map=supplementary_search_many([q],per_source_limit=6,max_workers=2,time_budget_seconds=10.0)
        secondary=sec_map.get(q,[])
    return jsonify(success=True,query=q,data=official,secondary_research=secondary,
                   secondary_disclaimer='Hukumonline/secondary results are research aids only; legal status and applicability require official sources.')

app.register_blueprint(create_case_analysis_blueprint(
    allowed_file=allowed_file,
    verification_queries=_case_verification_queries,
    case_payload=_case_analysis_payload,
    ensure_evidence_to_action=_ensure_evidence_to_action,
    readiness_profile=_case_readiness_profile,
))

@app.route('/api/communications',methods=['GET','POST'])
def communications():
    if request.method=='GET': return jsonify(success=True,data=CommunicationManager.all(request.args.get('limit',30,type=int)))
    d=request.get_json(silent=True) or {}
    if not d.get('message','').strip(): return jsonify(success=False,error='Isi komunikasi wajib diisi'),400
    i=CommunicationManager.save(d); return jsonify(success=True,communication_id=i,message='Komunikasi tersimpan')

@app.route('/api/communication/whatsapp-link', methods=['POST'])
def communication_whatsapp_link():
    d=request.get_json(silent=True) or {}
    client_id=(d.get('client_id') or '').strip()
    number=(d.get('whatsapp_number') or '').strip()
    if not number and client_id:
        contact=CommunicationManager.contact_by_client_id(client_id) or {}
        number=(contact.get('whatsapp_number') or '').strip()
    digits=''.join(ch for ch in number if ch.isdigit())
    if digits.startswith('0'):
        digits='62'+digits[1:]
    elif digits.startswith('8'):
        digits='62'+digits
    if not digits.startswith('62') or len(digits)<10:
        return jsonify(success=False,error='Nomor WhatsApp terdaftar belum valid. Gunakan format 08xxxxxxxxxx atau 62xxxxxxxxxx.'),400
    subject=(d.get('subject') or '').strip()
    message=(d.get('message') or '').strip()
    if not message:
        return jsonify(success=False,error='Isi pesan belum tersedia.'),400
    from urllib.parse import quote
    body=(subject+'\n\n' if subject else '')+message
    return jsonify(success=True,status='WHATSAPP_COMPOSER_READY',delivery_status='WHATSAPP_COMPOSER_READY',
                   client_id=client_id or None,whatsapp_number=digits,
                   whatsapp_url=f'https://wa.me/{digits}?text={quote(body)}')


@app.route('/api/communication/generate',methods=['POST'])
def generate_communication():
    d=request.get_json(silent=True) or {}
    result=build_client_communication(d)
    return jsonify(success=True,data=result)

@app.route('/api/audit/logs')
def audit(): return jsonify(success=True,data=AuditLogger.get_recent_logs(request.args.get('limit',50,type=int)))

@app.route('/api/ai/status')
def ai_status_endpoint(): return jsonify(success=True,**ai_status())

@app.route('/api/health')
def health(): return jsonify(status='online',timestamp=datetime.now().isoformat(),version=LEXICORE_VERSION,schema_version=SCHEMA_VERSION,product={'name':PRODUCT_NAME,'label':PRODUCT_LABEL,'initiative':INITIATIVE,'firm':FIRM_NAME,'release':release_metadata()},modules=['legal_drafting','contract_review','legal_research','compliance_risk','case_analysis','full_document_ai_reasoning','official_jdih_federation','dynamic_case_regulatory_retrieval','offline_regulatory_database','hybrid_legal_research','regulatory_corpus','regulatory_intelligence','legal_relationship_graph','legal_timeline','norm_conflict_detector','evidence_to_action','client_communication'],ai=ai_status())

if __name__=='__main__':
    # Security fix (v1.3.3.14): this app has zero authentication and stores
    # confidential client/case data. The previous hardcoded
    # app.run(debug=True, host='0.0.0.0', port=5000) had two compounding risks:
    #   1) debug=True enables the Werkzeug interactive debugger — if an
    #      unhandled exception is ever hit and that debugger page is reachable,
    #      it allows arbitrary Python code execution in the browser.
    #   2) host='0.0.0.0' binds to every network interface, so anyone on the
    #      same WiFi/LAN (office, coworking space) could reach the API and,
    #      combined with (1), potentially get code execution — not just read
    #      client data via the unauthenticated endpoints.
    # Defaults now are safe for local single-user use (127.0.0.1, debug off).
    # Set LEXICORE_DEBUG=1 only on a trusted machine during development, and
    # LEXICORE_HOST/LEXICORE_PORT explicitly if this is ever deployed for real
    # multi-user access — which also requires adding authentication first.
    debug_mode = os.environ.get('LEXICORE_DEBUG', '0') == '1'
    host = os.environ.get('LEXICORE_HOST', '127.0.0.1')
    port = int(os.environ.get('LEXICORE_PORT', '5000'))
    if host == '0.0.0.0' and not os.environ.get('LEXICORE_ALLOW_NETWORK_EXPOSURE'):
        print('LexiCore: LEXICORE_HOST=0.0.0.0 diminta tetapi belum ada autentikasi di API ini.')
        print('  Set LEXICORE_ALLOW_NETWORK_EXPOSURE=1 juga jika ini benar-benar disengaja (jaringan tepercaya saja).')
        host = '127.0.0.1'
    print(f'{PRODUCT_LABEL} [{LEXICORE_VERSION}] — Comprehensive Legal Intelligence + Hybrid Research + Expanded Drafting — http://{host}:{port}')
    if debug_mode:
        print('LexiCore: berjalan dengan LEXICORE_DEBUG=1 (Werkzeug debugger aktif). Jangan gunakan di jaringan bersama.')
    app.run(debug=debug_mode, host=host, port=port, threaded=True)
